from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from pyproj import Transformer

from sekolah.models import Sekolah, KategoriSekolah


class Command(BaseCommand):

    help = 'Import data sekolah dari Excel'

    def add_arguments(self, parser):
        parser.add_argument(
            'file_excel',
            type=str
        )

    def handle(self, *args, **kwargs):

        file_excel = kwargs['file_excel']

        workbook = load_workbook(file_excel)
        sheet = workbook.active

        total_baru = 0
        total_update = 0

        # UTM Zone 50N → WGS84
        transformer = Transformer.from_crs(
            "EPSG:32650",
            "EPSG:4326",
            always_xy=True
        )

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            npsn = row[0]
            nama = row[1]
            kategori_nama = row[2]
            status = row[3]
            kecamatan = row[4]
            desa = row[5]
            alamat = row[6]
            x_utm = row[7]
            y_utm = row[8]

            if not nama:
                continue

            # KONVERSI UTM → LATLONG
            longitude, latitude = transformer.transform(
                x_utm,
                y_utm
            )

            kategori, created = KategoriSekolah.objects.get_or_create(
                nama=kategori_nama
            )

            sekolah, created = Sekolah.objects.update_or_create(
                npsn=npsn,

                defaults={

                    'nama' : nama,
                    'kategori': kategori,
                    'status': status,
                    'kecamatan': kecamatan,
                    'desa': desa,
                    'alamat': alamat,

                    'x_utm': x_utm,
                    'y_utm': y_utm,

                    'latitude': latitude,
                    'longitude': longitude,
                }
            )

            if created:
                total_baru += 1
            else:
                total_update += 1

        self.stdout.write(
            self.style.SUCCESS(
                f'Import selesai. Baru: {total_baru}, Update: {total_update}'
            )
        )