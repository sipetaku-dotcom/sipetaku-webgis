from django.contrib.auth.models import User
from django.db import transaction
from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Sum, Count, Q, F
from django.http import HttpResponse
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.template.loader import render_to_string, get_template

from sekolah.models import Sekolah, KategoriSekolah
from guru.models import (
    Pegawai,
    RiwayatPegawai,
)
from murid.models import (
    Murid,
    TahunAjaran,
    RiwayatMurid,
)
from akun.models import ProfilUser
from prestasi.models import Prestasi
from aset.models import Aset, JenisAset
from akun.views import (
    ambil_profil_user,
    user_admin_kabupaten,
    hitung_kesiapan_tahun_ajaran,
)
from aset.penilaian import hitung_penilaian_sekolah

from datetime import date, timedelta
from openpyxl import Workbook
from xhtml2pdf import pisa
from io import BytesIO

from pyproj import Transformer

import json


def user_admin_kabupaten(request):

    profil = ambil_profil_user(request)

    if request.user.is_superuser:
        return True

    if profil and profil.role == 'admin_kabupaten':
        return True

    return False


def index(request):

    profil = ambil_profil_user(request)

    if request.user.is_authenticated:

        if user_admin_kabupaten(request):
            sekolah = Sekolah.objects.annotate(
                total_prestasi=Count('prestasi', distinct=True),
                
                total_murid=Count('murid', distinct=True)
            )

        elif profil and profil.role == 'admin_kecamatan':
            sekolah = Sekolah.objects.filter(
                kecamatan=profil.kecamatan
            ).annotate(
                total_prestasi=Count('prestasi', distinct=True),
                
                total_murid=Count('murid', distinct=True)
            )

        else:
            sekolah = Sekolah.objects.filter(
                user=request.user
            ).annotate(
                total_prestasi=Count('prestasi', distinct=True),
                
                total_murid=Count('murid', distinct=True)
            )

    else:
        sekolah = Sekolah.objects.annotate(
            total_prestasi=Count('prestasi', distinct=True),
            
            total_murid=Count('murid', distinct=True)
        )

    kecamatan = request.GET.get('kecamatan')
    kategori = request.GET.get('kategori')
    status_prestasi = request.GET.get('status_prestasi')

    if kecamatan:
        sekolah = sekolah.filter(kecamatan=kecamatan)

    if kategori:
        sekolah = sekolah.filter(kategori_id=kategori)

    if status_prestasi == 'punya':
        sekolah = sekolah.filter(
            total_prestasi__gt=0
        )

    if status_prestasi == 'tidak':
        sekolah = sekolah.filter(
            total_prestasi=0
        )

    if user_admin_kabupaten(request):

        daftar_kecamatan = Sekolah.objects.values_list(
            'kecamatan',
            flat=True
        ).distinct().order_by('kecamatan')

    elif profil and profil.role == 'admin_kecamatan':

        daftar_kecamatan = Sekolah.objects.filter(
            kecamatan=profil.kecamatan
        ).values_list(
            'kecamatan',
            flat=True
        ).distinct().order_by('kecamatan')

    else:

        daftar_kecamatan = sekolah.values_list(
            'kecamatan',
            flat=True
        ).distinct().order_by('kecamatan')

    daftar_kategori = KategoriSekolah.objects.all().order_by('nama')

    total_sekolah_tampil = sekolah.count()

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif:

        total_guru = (
            RiwayatPegawai.objects
            .filter(
                sekolah__in=sekolah,
                tahun_ajaran=tahun_aktif,
                jenis_pegawai='PENDIDIK',
                status_akhir='AKTIF'
            )
            .count()
        )

    else:

        total_guru = 0

    total_murid = sekolah.aggregate(
        total=Sum('total_murid')
    )['total'] or 0

    total_prestasi = sekolah.aggregate(
        total=Sum('total_prestasi')
    )['total'] or 0

    mode_peta = request.GET.get(
        'mode',
        'sekolah'
    )

    boleh_melihat_prioritas = (
        request.user.is_authenticated
        and (
            user_admin_kabupaten(request)
            or (
                profil
                and profil.role
                == 'admin_kecamatan'
            )
        )
    )

    if (
        mode_peta == 'prioritas'
        and not boleh_melihat_prioritas
    ):

        mode_peta = 'sekolah'

    if mode_peta == 'prioritas':

        daftar_sekolah_peta = list(
            sekolah
        )

        for data_sekolah in (
            daftar_sekolah_peta
        ):

            penilaian = (
                hitung_penilaian_sekolah(
                    data_sekolah
                )
            )

            skor_prioritas = (
                penilaian[
                    'skor_prioritas'
                ]
            )

            if not penilaian[
                'penilaian_final'
            ]:

                warna_prioritas = 'grey'

                status_prioritas = (
                    'Data Belum Lengkap'
                )

            elif skor_prioritas >= 80:

                warna_prioritas = 'red'

                status_prioritas = (
                    'Sangat Mendesak'
                )

            elif skor_prioritas >= 60:

                warna_prioritas = 'orange'

                status_prioritas = (
                    'Mendesak'
                )

            elif skor_prioritas >= 40:

                warna_prioritas = 'yellow'

                status_prioritas = (
                    'Perlu Perhatian'
                )

            elif skor_prioritas >= 20:

                warna_prioritas = 'blue'

                status_prioritas = 'Cukup'

            else:

                warna_prioritas = 'green'

                status_prioritas = 'Baik'

            data_sekolah.warna_prioritas = (
                warna_prioritas
            )

            data_sekolah.status_prioritas = (
                status_prioritas
            )

            data_sekolah.skor_prioritas_peta = (
                skor_prioritas
            )

            data_sekolah.progres_aset_peta = (
                penilaian['progres']
            )

            data_sekolah.penilaian_final_peta = (
                penilaian[
                    'penilaian_final'
                ]
            )

            data_sekolah.legalitas_aman_peta = (
                penilaian[
                    'legalitas_aman'
                ]
            )

            data_sekolah.alasan_prioritas_peta = (
                ' | '.join(
                    penilaian['alasan'][:3]
                )
            )

        sekolah = daftar_sekolah_peta

    context = {
        'sekolah': sekolah,
        'daftar_kecamatan': daftar_kecamatan,
        'daftar_kategori': daftar_kategori,
        'status_prestasi': status_prestasi,
        'total_sekolah_tampil': total_sekolah_tampil,
        'total_guru': total_guru,
        'total_murid': total_murid,
        'total_prestasi': total_prestasi,
        'mode_peta':
            mode_peta,

        'boleh_melihat_prioritas':
            boleh_melihat_prioritas,
    }

    return render(request, 'peta/index.html', context)


def detail_sekolah(request, id):

    sekolah = get_object_or_404(
        Sekolah,
        id=id
    )

    profil = ambil_profil_user(request)

    if request.user.is_authenticated:

        if user_admin_kabupaten(request):
            pass

        elif profil and profil.role == 'admin_kecamatan':

            if sekolah.kecamatan != profil.kecamatan:
                return redirect('/')

        else:

            if sekolah.user != request.user:
                return redirect('/')

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif:

        daftar_guru = (
            RiwayatPegawai.objects
            .filter(
                sekolah=sekolah,
                tahun_ajaran=tahun_aktif,
                jenis_pegawai='PENDIDIK',
                status_akhir='AKTIF'
            )
            .select_related(
                'pegawai'
            )
            .order_by(
                'pegawai__nama'
            )
        )

    else:

        daftar_guru = (
            RiwayatPegawai.objects.none()
        )
    daftar_murid = sekolah.murid.all()
    daftar_aset = sekolah.aset.all()
    daftar_prestasi = sekolah.prestasi.all()

    total_guru = daftar_guru.count()
    total_murid = daftar_murid.count()
    total_aset = daftar_aset.count()
    total_prestasi = daftar_prestasi.count()

    prestasi_per_tahun = daftar_prestasi.values(
        'tahun'
    ).annotate(
        jumlah=Count('id')
    ).order_by('tahun')

    label_tahun_prestasi = []
    data_tahun_prestasi = []

    for p in prestasi_per_tahun:
        label_tahun_prestasi.append(p['tahun'])
        data_tahun_prestasi.append(p['jumlah'])

    aset_per_kondisi = daftar_aset.values(
        'kondisi'
    ).annotate(
        jumlah=Count('id')
    ).order_by('kondisi')

    label_kondisi_aset = []
    data_kondisi_aset = []

    for a in aset_per_kondisi:
        label_kondisi_aset.append(a['kondisi'])
        data_kondisi_aset.append(a['jumlah'])

    jumlah_laki = daftar_murid.filter(
        jenis_kelamin='L'
    ).count()

    jumlah_perempuan = daftar_murid.filter(
        jenis_kelamin='P'
    ).count()

    context = {
        'sekolah': sekolah,
        'daftar_guru': daftar_guru,
        'daftar_murid': daftar_murid,
        'daftar_aset': daftar_aset,
        'daftar_prestasi': daftar_prestasi,
        'total_guru': total_guru,
        'total_murid': total_murid,
        'total_aset': total_aset,
        'total_prestasi': total_prestasi,
        'jumlah_laki': jumlah_laki,
        'jumlah_perempuan': jumlah_perempuan,
        'label_tahun_prestasi': json.dumps(label_tahun_prestasi),
        'data_tahun_prestasi': json.dumps(data_tahun_prestasi),
        'label_kondisi_aset': json.dumps(label_kondisi_aset),
        'data_kondisi_aset': json.dumps(data_kondisi_aset),
    }

    return render(
        request,
        'peta/detail_sekolah.html',
        context
    )


@login_required
def dashboard(request):

    # =====================================================
    # PROFIL DAN CAKUPAN AKSES
    # =====================================================

    profil = ambil_profil_user(request)

    if request.user.is_superuser:

        sekolah_qs = (
            Sekolah.objects
            .all()
        )

        judul_dashboard = (
            'Dashboard Superadmin'
        )

        nama_role = (
            'Superadmin'
        )

        cakupan_wilayah = (
            'Kabupaten Kutai Timur'
        )


    elif (
        profil
        and
        profil.role == 'admin_kabupaten'
    ):

        sekolah_qs = (
            Sekolah.objects
            .all()
        )

        judul_dashboard = (
            'Dashboard Admin Kabupaten'
        )

        nama_role = (
            'Admin Kabupaten'
        )

        cakupan_wilayah = (
            'Kabupaten Kutai Timur'
        )


    elif (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        sekolah_qs = (
            Sekolah.objects
            .filter(
                kecamatan=profil.kecamatan
            )
        )

        judul_dashboard = (
            'Dashboard Admin Kecamatan'
        )

        nama_role = (
            'Admin Kecamatan'
        )

        cakupan_wilayah = (
            f'Kecamatan {profil.kecamatan}'
        )


    else:

        messages.error(
            request,
            'Anda tidak memiliki akses ke dashboard ini.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = TahunAjaran.objects.filter(aktif=True).first()

    # Harus dihitung sebelum blok status kesiapan
    total_sekolah = sekolah_qs.count()

    # =========================================================
    # STATUS KESIAPAN SEKOLAH
    # =========================================================
    tahun_tujuan = None

    total_siap = 0
    total_belum_siap = 0
    total_rombel_belum_siap = 0
    total_belum_ada_data = 0
    persentase_siap = 0

    if tahun_aktif:
        tahun_tujuan = TahunAjaran.objects.filter(
            aktif=False,
            tahun_sebelumnya=tahun_aktif,
        ).first()

    if tahun_aktif and tahun_tujuan:
        for sekolah in sekolah_qs:
            hasil_kesiapan = hitung_kesiapan_tahun_ajaran(
                sekolah,
                tahun_aktif,
                tahun_tujuan,
            )

            status_kesiapan = hasil_kesiapan.get("status")

            if (
                hasil_kesiapan.get("siap")
                or hasil_kesiapan.get("siap_tanpa_siswa")
            ):
                total_siap += 1

            elif status_kesiapan == "ROMBEL BELUM SIAP":
                total_rombel_belum_siap += 1

            elif status_kesiapan == "BELUM ADA DATA":
                total_belum_ada_data += 1

            else:
                total_belum_siap += 1

        if total_sekolah > 0:
            persentase_siap = round(
                (total_siap / total_sekolah) * 100,
                1,
            )



    # =====================================================
    # TOTAL GURU AKTIF
    # =====================================================

    if tahun_aktif:

        total_guru = (
            RiwayatPegawai.objects
            .filter(
                sekolah__in=sekolah_qs,
                tahun_ajaran=tahun_aktif,
                jenis_pegawai='PENDIDIK',
                status_akhir='AKTIF'
            )
            .count()
        )

    else:

        total_guru = 0


    # =====================================================
    # TOTAL MURID AKTIF
    # =====================================================

    if tahun_aktif:

        total_murid = (
            RiwayatMurid.objects
            .filter(
                murid__sekolah__in=sekolah_qs,
                tahun_ajaran=tahun_aktif,
                status_akhir='AKTIF'
            )
            .values(
                'murid_id'
            )
            .distinct()
            .count()
        )

    else:

        total_murid = 0


    # =====================================================
    # TOTAL OPERATOR AKTIF
    # =====================================================

    total_operator = (
        ProfilUser.objects
        .filter(
            role='operator_sekolah',
            user__is_active=True,
            sekolah__in=sekolah_qs
        )
        .count()
    )


    # =====================================================
    # SEKOLAH TANPA OPERATOR
    # =====================================================

    sekolah_tanpa_operator = (
        sekolah_qs
        .filter(
            user__isnull=True
        )
        .count()
    )


    # =====================================================
    # DATA LAMA
    # SEMENTARA DIPERTAHANKAN AGAR TEMPLATE LAMA
    # TIDAK LANGSUNG ERROR
    # =====================================================

    total_aset = (
        Aset.objects
        .filter(
            sekolah__in=sekolah_qs
        )
        .count()
    )

    total_prestasi = (
        Prestasi.objects
        .filter(
            sekolah__in=sekolah_qs
        )
        .count()
    )

    total_kategori = (
        KategoriSekolah.objects
        .count()
    )


    # =====================================================
    # SEKOLAH PER KATEGORI
    # WAJIB MENGIKUTI CAKUPAN ROLE
    # =====================================================

    sekolah_per_kategori = (
        KategoriSekolah.objects
        .annotate(
            jumlah=Count(
                'sekolah',
                filter=Q(
                    sekolah__in=sekolah_qs
                ),
                distinct=True
            )
        )
        .order_by(
            'nama'
        )
    )


    # =====================================================
    # GURU AKAN PENSIUN SATU TAHUN KE DEPAN
    # =====================================================

    hari_ini = date.today()

    batas_pensiun = (
        hari_ini
        + timedelta(days=365)
    )


    if tahun_aktif:

        guru_akan_pensiun_qs = (
            RiwayatPegawai.objects
            .filter(
                sekolah__in=sekolah_qs,
                tahun_ajaran=tahun_aktif,
                jenis_pegawai='PENDIDIK',
                status_akhir='AKTIF',
                tanggal_pensiun__isnull=False,
                tanggal_pensiun__gte=hari_ini,
                tanggal_pensiun__lte=batas_pensiun
            )
            .select_related(
                'pegawai',
                'sekolah'
            )
            .order_by(
                'tanggal_pensiun',
                'pegawai__nama'
            )
        )

    else:

        guru_akan_pensiun_qs = (
            RiwayatPegawai.objects
            .none()
        )


    total_guru_akan_pensiun = (
        guru_akan_pensiun_qs
        .count()
    )


    # HANYA 10 GURU TERDEKAT UNTUK DASHBOARD

    guru_akan_pensiun = (
        guru_akan_pensiun_qs[:10]
    )


    # =====================================================
    # GURU PENSIUN PER KECAMATAN
    # =====================================================

    guru_pensiun_per_kecamatan = (
        guru_akan_pensiun_qs
        .values(
            'sekolah__kecamatan'
        )
        .annotate(
            jumlah=Count('id')
        )
        .order_by(
            'sekolah__kecamatan'
        )
    )


    # =====================================================
    # PRESTASI PER TINGKAT
    # WAJIB MENGIKUTI CAKUPAN ROLE
    # =====================================================

    prestasi_per_tingkat = (
        Prestasi.objects
        .filter(
            sekolah__in=sekolah_qs
        )
        .values(
            'tingkat'
        )
        .annotate(
            jumlah=Count('id')
        )
        .order_by(
            'tingkat'
        )
    )


    # =====================================================
    # DATA GRAFIK
    # =====================================================

    label_prestasi = []
    data_prestasi = []

    for prestasi in prestasi_per_tingkat:

        label_prestasi.append(
            prestasi['tingkat']
        )

        data_prestasi.append(
            prestasi['jumlah']
        )


    label_kategori = []
    data_kategori = []

    for kategori_sekolah in sekolah_per_kategori:

        label_kategori.append(
            kategori_sekolah.nama
        )

        data_kategori.append(
            kategori_sekolah.jumlah
        )


    label_pensiun = []
    data_pensiun = []

    for pensiun in guru_pensiun_per_kecamatan:

        label_pensiun.append(
            pensiun['sekolah__kecamatan']
        )

        data_pensiun.append(
            pensiun['jumlah']
        )


    # =====================================================
    # FILTER PRESTASI
    # =====================================================

    tingkat = (
        request.GET.get(
            'tingkat',
            ''
        )
        .strip()
    )

    tahun = (
        request.GET.get(
            'tahun',
            ''
        )
        .strip()
    )

    kecamatan = (
        request.GET.get(
            'kecamatan',
            ''
        )
        .strip()
    )

    kategori = (
        request.GET.get(
            'kategori',
            ''
        )
        .strip()
    )


    daftar_prestasi = (
        Prestasi.objects
        .filter(
            sekolah__in=sekolah_qs
        )
        .select_related(
            'sekolah',
            'murid',
            'sekolah__kategori'
        )
    )


    if tingkat:

        daftar_prestasi = (
            daftar_prestasi
            .filter(
                tingkat=tingkat
            )
        )


    if tahun:

        daftar_prestasi = (
            daftar_prestasi
            .filter(
                tahun=tahun
            )
        )


    if kecamatan:

        daftar_prestasi = (
            daftar_prestasi
            .filter(
                sekolah__kecamatan=kecamatan
            )
        )


    if kategori:

        daftar_prestasi = (
            daftar_prestasi
            .filter(
                sekolah__kategori_id=kategori
            )
        )


    # =====================================================
    # PILIHAN FILTER
    # =====================================================

    daftar_tahun = (
        Prestasi.objects
        .filter(
            sekolah__in=sekolah_qs
        )
        .values_list(
            'tahun',
            flat=True
        )
        .distinct()
        .order_by(
            '-tahun'
        )
    )


    daftar_kecamatan = (
        sekolah_qs
        .values_list(
            'kecamatan',
            flat=True
        )
        .distinct()
        .order_by(
            'kecamatan'
        )
    )


    daftar_kategori = (
        KategoriSekolah.objects
        .all()
        .order_by(
            'nama'
        )
    )


    # =====================================================
    # REKAPITULASI SEKOLAH PER KECAMATAN
    # =====================================================

    sekolah_per_kecamatan = []

    daftar_kecamatan_dashboard = (
        sekolah_qs
        .values_list(
            'kecamatan',
            flat=True
        )
        .distinct()
        .order_by(
            'kecamatan'
        )
    )


    for nama_kecamatan in daftar_kecamatan_dashboard:

        daftar_sekolah = (
            sekolah_qs
            .filter(
                kecamatan=nama_kecamatan
            )
            .order_by(
                'nama'
            )
        )


        jumlah_per_kategori = (
            daftar_sekolah
            .values(
                'kategori__nama'
            )
            .annotate(
                jumlah=Count('id')
            )
            .order_by(
                'kategori__nama'
            )
        )


        sekolah_per_kecamatan.append({
            'nama_kecamatan':
                nama_kecamatan,

            'jumlah_sekolah':
                daftar_sekolah.count(),

            'jumlah_per_kategori':
                jumlah_per_kategori,
        })


    # =====================================================
    # TOP SEKOLAH BERPRESTASI
    # SEMENTARA DIPERTAHANKAN UNTUK TEMPLATE LAMA
    # =====================================================

    ranking_sekolah = (
        sekolah_qs
        .annotate(
            jumlah_prestasi=Count(
                'prestasi'
            )
        )
        .order_by(
            '-jumlah_prestasi',
            'nama'
        )[:10]
    )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        # IDENTITAS DASHBOARD

        'judul_dashboard':
            judul_dashboard,

        'nama_role':
            nama_role,

        'cakupan_wilayah':
            cakupan_wilayah,

        'tahun_aktif':
            tahun_aktif,


        "tahun_tujuan": tahun_tujuan,
        "total_siap": total_siap,
        "total_belum_siap": total_belum_siap,
        "total_rombel_belum_siap": total_rombel_belum_siap,
        "total_belum_ada_data": total_belum_ada_data,
        "persentase_siap": persentase_siap,
        "kesiapan_tersedia": bool(tahun_aktif and tahun_tujuan),



        # RINGKASAN UTAMA

        'total_sekolah':
            total_sekolah,

        'total_murid':
            total_murid,

        'total_guru':
            total_guru,

        'total_operator':
            total_operator,

        'sekolah_tanpa_operator':
            sekolah_tanpa_operator,

        'total_guru_akan_pensiun':
            total_guru_akan_pensiun,


        # DATA LAMA UNTUK TEMPLATE SAAT INI

        'total_aset':
            total_aset,

        'total_kategori':
            total_kategori,

        'total_prestasi':
            total_prestasi,


        # REKAPITULASI

        'sekolah_per_kategori':
            sekolah_per_kategori,

        'sekolah_per_kecamatan':
            sekolah_per_kecamatan,


        # GURU PENSIUN

        'guru_akan_pensiun':
            guru_akan_pensiun,


        # PRESTASI

        'prestasi_per_tingkat':
            prestasi_per_tingkat,

        'daftar_prestasi':
            daftar_prestasi,

        'daftar_tahun':
            daftar_tahun,

        'tingkat_terpilih':
            tingkat,

        'tahun_terpilih':
            tahun,


        # FILTER

        'daftar_kecamatan':
            daftar_kecamatan,

        'daftar_kategori':
            daftar_kategori,

        'kecamatan_terpilih':
            kecamatan,

        'kategori_terpilih':
            kategori,


        # RANKING

        'ranking_sekolah':
            ranking_sekolah,


        # DATA GRAFIK

        'label_prestasi':
            json.dumps(
                label_prestasi
            ),

        'data_prestasi':
            json.dumps(
                data_prestasi
            ),

        'label_kategori':
            json.dumps(
                label_kategori
            ),

        'data_kategori':
            json.dumps(
                data_kategori
            ),

        'label_pensiun':
            json.dumps(
                label_pensiun
            ),

        'data_pensiun':
            json.dumps(
                data_pensiun
            ),
    }


    return render(
        request,
        'peta/dashboard.html',
        context
    )


@login_required
def export_sekolah_excel(request):

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sekolah"

    ws.append([
        'No',
        'Nama Sekolah',
        'NPSN',
        'Kategori',
        'Status',
        'Kecamatan',
        'Desa',
        'Alamat',
        'Latitude',
        'Longitude',
    ])

    profil = ambil_profil_user(request)

    if user_admin_kabupaten(request):
        sekolah = Sekolah.objects.all()

    elif profil and profil.role == 'admin_kecamatan':
        sekolah = Sekolah.objects.filter(
            kecamatan=profil.kecamatan
        )

    else:
        return redirect('/')

    for no, s in enumerate(sekolah, start=1):
        ws.append([
            no,
            s.nama,
            s.npsn,
            s.kategori.nama,
            s.status,
            s.kecamatan,
            s.desa,
            s.alamat,
            s.latitude,
            s.longitude,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=data_sekolah.xlsx'

    wb.save(response)

    return response


# @login_required
# def export_guru_excel(request):

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Data Guru"

#     ws.append([
#         'No',
#         'Nama Guru',
#         'NIP',
#         'NIK',
#         'Jenis Kelamin',
#         'Sekolah',
#         'Kecamatan',
#         'Jabatan',
#         'Status Pegawai',
#         'Mata Pelajaran',
#         'Tanggal Lahir',
#         'Tanggal Pensiun',
#         'No HP',
#     ])

#     profil = ambil_profil_user(request)

#     if user_admin_kabupaten(request):

#         guru = Guru.objects.select_related(
#             'sekolah'
#         ).all()

#     elif profil and profil.role == 'admin_kecamatan':

#         guru = Guru.objects.select_related(
#             'sekolah'
#         ).filter(
#             sekolah__kecamatan=profil.kecamatan
#         )

#     else:

#         return redirect('/')

#     for no, g in enumerate(guru, start=1):
#         ws.append([
#             no,
#             g.nama,
#             g.nip,
#             g.nik,
#             g.jenis_kelamin,
#             g.sekolah.nama,
#             g.sekolah.kecamatan,
#             g.jabatan,
#             g.status_pegawai,
#             g.mata_pelajaran,
#             g.tanggal_lahir,
#             g.tanggal_pensiun,
#             g.no_hp,
#         ])

#     response = HttpResponse(
#         content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
#     )

#     response['Content-Disposition'] = 'attachment; filename=data_guru.xlsx'

#     wb.save(response)

#     return response


@login_required
def export_murid_excel(request):

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Murid"

    ws.append([
        'No',
        'Nama Murid',
        'NISN',
        'NIK',
        'Jenis Kelamin',
        'Sekolah',
        'Kecamatan',
        'Kelas',
        'Tanggal Lahir',
        'Alamat',
        'Prestasi',
    ])

    profil = ambil_profil_user(request)

    if user_admin_kabupaten(request):

        murid = Murid.objects.select_related(
            'sekolah'
        ).all()

    elif profil and profil.role == 'admin_kecamatan':

        murid = Murid.objects.select_related(
            'sekolah'
        ).filter(
            sekolah__kecamatan=profil.kecamatan
        )

    else:

        return redirect('/')

    for no, m in enumerate(murid, start=1):
        ws.append([
            no,
            m.nama,
            m.nisn,
            m.nik,
            m.jenis_kelamin,
            m.sekolah.nama,
            m.sekolah.kecamatan,
            m.kelas,
            m.tanggal_lahir,
            m.alamat,
            m.prestasi,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=data_murid.xlsx'

    wb.save(response)

    return response


@login_required
def export_prestasi_excel(request):

    tingkat = request.GET.get('tingkat')
    tahun = request.GET.get('tahun')
    kecamatan = request.GET.get('kecamatan')
    kategori = request.GET.get('kategori')

    profil = ambil_profil_user(request)

    if user_admin_kabupaten(request):

        prestasi = Prestasi.objects.select_related(
            'sekolah',
            'murid'
        ).all()

    elif profil and profil.role == 'admin_kecamatan':

        prestasi = Prestasi.objects.select_related(
            'sekolah',
            'murid'
        ).filter(
            sekolah__kecamatan=profil.kecamatan
        )

    else:

        return redirect('/')

    if tingkat:
        prestasi = prestasi.filter(tingkat=tingkat)

    if tahun:
        prestasi = prestasi.filter(tahun=tahun)

    if kecamatan:
        prestasi = prestasi.filter(sekolah__kecamatan=kecamatan)

    if kategori:
        prestasi = prestasi.filter(sekolah__kategori_id=kategori)

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Prestasi"

    ws.append([
        'No',
        'Nama Murid',
        'Sekolah',
        'Kategori Sekolah',
        'Kecamatan',
        'Nama Lomba',
        'Bidang Lomba',
        'Tingkat',
        'Juara',
        'Tahun',
        'Keterangan',
    ])

    for no, p in enumerate(prestasi, start=1):
        ws.append([
            no,
            p.murid.nama,
            p.sekolah.nama,
            p.sekolah.kategori.nama,
            p.sekolah.kecamatan,
            p.nama_lomba,
            p.bidang_lomba,
            p.tingkat,
            p.juara,
            p.tahun,
            p.keterangan,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=data_prestasi_filter.xlsx'

    wb.save(response)

    return response


@login_required
def export_aset_excel(request):

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Aset"

    ws.append([
        'No',
        'Nama Aset',
        'Kategori',
        'Jumlah',
        'Kondisi',
        'Sekolah',
        'Kecamatan',
        'Keterangan',
    ])

    profil = ambil_profil_user(request)

    if user_admin_kabupaten(request):

        aset = Aset.objects.select_related(
            'sekolah'
        ).all()

    elif profil and profil.role == 'admin_kecamatan':

        aset = Aset.objects.select_related(
            'sekolah'
        ).filter(
            sekolah__kecamatan=profil.kecamatan
        )

    else:

        return redirect('/')

    for no, a in enumerate(aset, start=1):
        ws.append([
            no,
            a.nama,
            a.kategori,
            a.jumlah,
            a.kondisi,
            a.sekolah.nama,
            a.sekolah.kecamatan,
            a.keterangan,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=data_aset.xlsx'

    wb.save(response)

    return response


# @login_required
# def export_guru_pensiun_excel(request):

#     hari_ini = date.today()
#     batas_pensiun = hari_ini + timedelta(days=365)

#     profil = ambil_profil_user(request)

#     if user_admin_kabupaten(request):

#         guru = Guru.objects.select_related('sekolah').filter(
#             tanggal_pensiun__isnull=False,
#             tanggal_pensiun__gte=hari_ini,
#             tanggal_pensiun__lte=batas_pensiun
#         )

#     elif profil and profil.role == 'admin_kecamatan':

#         guru = Guru.objects.select_related('sekolah').filter(
#             sekolah__kecamatan=profil.kecamatan,
#             tanggal_pensiun__isnull=False,
#             tanggal_pensiun__gte=hari_ini,
#             tanggal_pensiun__lte=batas_pensiun
#         )

#     else:

#         return redirect('/')

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Guru Akan Pensiun"

#     ws.append([
#         'No',
#         'Nama Guru',
#         'Sekolah',
#         'Kecamatan',
#         'Jabatan',
#         'Status Pegawai',
#         'Tanggal Pensiun',
#         'No HP',
#     ])

#     for no, g in enumerate(guru, start=1):
#         ws.append([
#             no,
#             g.nama,
#             g.sekolah.nama,
#             g.sekolah.kecamatan,
#             g.jabatan,
#             g.status_pegawai,
#             g.tanggal_pensiun,
#             g.no_hp,
#         ])

#     response = HttpResponse(
#         content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
#     )

#     response['Content-Disposition'] = 'attachment; filename=guru_akan_pensiun.xlsx'

#     wb.save(response)

#     return response


def export_detail_sekolah_pdf(request, id):

    sekolah = Sekolah.objects.get(id=id)

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif:

        daftar_guru = (
            RiwayatPegawai.objects
            .filter(
                sekolah=sekolah,
                tahun_ajaran=tahun_aktif,
                jenis_pegawai='PENDIDIK',
                status_akhir='AKTIF'
            )
        )

    else:

        daftar_guru = (
            RiwayatPegawai.objects.none()
        )

    daftar_murid = Murid.objects.filter(
        sekolah=sekolah
    )

    daftar_prestasi = Prestasi.objects.filter(
        murid__sekolah=sekolah
    )

    daftar_aset = Aset.objects.filter(
        sekolah=sekolah
    )

    total_guru = daftar_guru.count()
    total_murid = daftar_murid.count()
    total_prestasi = daftar_prestasi.count()
    total_aset = daftar_aset.count()

    template = get_template(
        'peta/pdf_detail_sekolah.html'
    )

    html = template.render({
        'sekolah': sekolah,
        'total_guru': total_guru,
        'total_murid': total_murid,
        'total_prestasi': total_prestasi,
        'total_aset': total_aset,
        'daftar_prestasi': daftar_prestasi[:10],
    })

    response = HttpResponse(
        content_type='application/pdf'
    )

    response[
        'Content-Disposition'
    ] = f'filename="profil-{sekolah.nama}.pdf"'

    pisa_status = pisa.CreatePDF(
        html,
        dest=response
    )

    if pisa_status.err:
        return HttpResponse(
            'Terjadi kesalahan saat membuat PDF'
        )

    return response


@login_required
def import_sekolah_geojson(request):

    if not request.user.is_superuser:
        messages.error(
            request,
            "Anda tidak memiliki akses untuk import data sekolah."
        )
        return redirect("/")

    def normalisasi_npsn(nilai):
        if nilai is None:
            return ""

        if isinstance(nilai, float) and nilai.is_integer():
            nilai = int(nilai)

        return str(nilai).strip().upper()

    def siapkan_akun_operator(sekolah, npsn):
        """
        Membuat atau menyinkronkan akun operator sekolah.

        Password admin123 hanya diberikan kepada akun baru.
        Password akun lama tidak akan direset saat GeoJSON diunggah ulang.
        """

        akun_baru = False

        profil_sekolah = (
            ProfilUser.objects
            .filter(
                role="operator_sekolah",
                sekolah=sekolah,
            )
            .select_related("user")
            .first()
        )

        user_operator = sekolah.user

        if user_operator and profil_sekolah:
            if user_operator.id != profil_sekolah.user_id:
                raise ValueError(
                    "sekolah memiliki lebih dari satu akun operator"
                )

        if not user_operator and profil_sekolah:
            user_operator = profil_sekolah.user

        user_dengan_npsn = User.objects.filter(
            username=npsn
        ).first()

        if user_operator:
            profil_user = ProfilUser.objects.filter(
                user=user_operator
            ).first()

            if profil_user and profil_user.role != "operator_sekolah":
                raise ValueError(
                    f'akun "{user_operator.username}" bukan akun operator sekolah'
                )

            if user_dengan_npsn and user_dengan_npsn.id != user_operator.id:
                raise ValueError(
                    f'username NPSN "{npsn}" sudah digunakan akun lain'
                )

            field_diperbarui = []

            if user_operator.username != npsn:
                user_operator.username = npsn
                field_diperbarui.append("username")

            if not user_operator.first_name:
                user_operator.first_name = f"Operator {sekolah.nama}"
                field_diperbarui.append("first_name")

            if field_diperbarui:
                user_operator.save(
                    update_fields=field_diperbarui
                )

        else:
            if user_dengan_npsn:
                profil_user_npsn = ProfilUser.objects.filter(
                    user=user_dengan_npsn
                ).first()

                if (
                    profil_user_npsn
                    and profil_user_npsn.role == "operator_sekolah"
                    and profil_user_npsn.sekolah_id == sekolah.id
                ):
                    user_operator = user_dengan_npsn

                else:
                    raise ValueError(
                        f'username NPSN "{npsn}" sudah digunakan akun lain'
                    )

            else:
                user_operator = User.objects.create_user(
                    username=npsn,
                    password="admin123",
                    first_name=f"Operator {sekolah.nama}",
                )

                akun_baru = True

        profil_operator = ProfilUser.objects.filter(
            user=user_operator
        ).first()

        if profil_operator:
            if profil_operator.role != "operator_sekolah":
                raise ValueError(
                    f'akun "{npsn}" sudah memiliki peran lain'
                )

            profil_operator.kecamatan = sekolah.kecamatan
            profil_operator.sekolah = sekolah

            profil_operator.save(
                update_fields=[
                    "kecamatan",
                    "sekolah",
                ]
            )

        else:
            ProfilUser.objects.create(
                user=user_operator,
                role="operator_sekolah",
                kecamatan=sekolah.kecamatan,
                sekolah=sekolah,
            )

        if sekolah.user_id != user_operator.id:
            sekolah.user = user_operator
            sekolah.save(
                update_fields=["user"]
            )

        return akun_baru

    if request.method == "POST":

        file_geojson = request.FILES.get("file_geojson")

        if not file_geojson:
            messages.error(
                request,
                "Silakan pilih file GeoJSON terlebih dahulu."
            )
            return redirect("/import-sekolah-geojson/")

        try:
            data = json.load(file_geojson)

        except Exception:
            messages.error(
                request,
                "File yang diunggah bukan GeoJSON yang valid."
            )
            return redirect("/import-sekolah-geojson/")

        if data.get("type") != "FeatureCollection":
            messages.error(
                request,
                "GeoJSON harus bertipe FeatureCollection."
            )
            return redirect("/import-sekolah-geojson/")

        wgs84_ke_utm = Transformer.from_crs(
            "EPSG:4326",
            "EPSG:32650",
            always_xy=True,
        )

        total_baru = 0
        total_update = 0
        total_gagal = 0
        total_akun_baru = 0
        total_wgs84 = 0
        total_utm = 0

        daftar_gagal = []
        npsn_diproses = set()

        for nomor, feature in enumerate(
            data.get("features", []),
            start=1,
        ):
            try:
                geometry = feature.get("geometry") or {}
                properties = feature.get("properties") or {}

                if geometry.get("type") != "Point":
                    raise ValueError("geometry bukan Point")

                coordinates = geometry.get("coordinates") or []

                if len(coordinates) < 2:
                    raise ValueError("koordinat tidak lengkap")

                npsn = normalisasi_npsn(
                    properties.get("npsn")
                    or properties.get("NPSN")
                    or properties.get("Npsn")
                )

                if not npsn:
                    raise ValueError("NPSN tidak ditemukan")

                if len(npsn) > 50:
                    raise ValueError(
                        "NPSN melebihi 50 karakter"
                    )

                if npsn in npsn_diproses:
                    raise ValueError(
                        f'NPSN "{npsn}" muncul lebih dari satu kali'
                    )

                nama = (
                    properties.get("nama")
                    or properties.get("NAMA")
                    or properties.get("Nama")
                    or properties.get("nama_sekolah")
                    or properties.get("NAMA_SEKOLAH")
                )

                kategori_nama = (
                    properties.get("kategori")
                    or properties.get("KATEGORI")
                    or properties.get("jenjang")
                    or properties.get("JENJANG")
                )

                status = (
                    properties.get("status")
                    or properties.get("STATUS")
                    or ""
                )

                kecamatan = (
                    properties.get("kecamatan")
                    or properties.get("KECAMATAN")
                    or ""
                )

                desa = (
                    properties.get("desa")
                    or properties.get("DESA")
                    or ""
                )

                alamat = (
                    properties.get("alamat")
                    or properties.get("ALAMAT")
                    or ""
                )

                if not nama:
                    raise ValueError(
                        "nama sekolah tidak ditemukan"
                    )

                if not kategori_nama:
                    raise ValueError(
                        "kategori sekolah tidak ditemukan"
                    )

                nama = str(nama).strip()
                kategori_nama = str(kategori_nama).strip()
                kecamatan = str(kecamatan).strip()
                desa = str(desa).strip()
                alamat = str(alamat).strip()

                if not kecamatan:
                    raise ValueError(
                        "kecamatan tidak ditemukan"
                    )

                koordinat_x = float(coordinates[0])
                koordinat_y = float(coordinates[1])

                if (
                    -180 <= koordinat_x <= 180
                    and
                    -90 <= koordinat_y <= 90
                ):
                    longitude = koordinat_x
                    latitude = koordinat_y

                    x_utm, y_utm = wgs84_ke_utm.transform(
                        longitude,
                        latitude,
                    )

                    sistem_koordinat = "WGS84"

                else:
                    x_utm = koordinat_x
                    y_utm = koordinat_y
                    sistem_koordinat = "UTM 50N"

                status_teks = str(status).strip().lower()

                if status_teks == "negeri":
                    status = "Negeri"

                elif status_teks == "swasta":
                    status = "Swasta"

                else:
                    raise ValueError(
                        "status harus Negeri atau Swasta"
                    )

                kategori = KategoriSekolah.objects.filter(
                    nama__iexact=kategori_nama
                ).first()

                if not kategori:
                    raise ValueError(
                        f'kategori "{kategori_nama}" belum terdaftar'
                    )

                with transaction.atomic():

                    sekolah = Sekolah.objects.filter(
                        npsn=npsn
                    ).first()

                    sekolah_baru = False

                    if not sekolah:
                        sekolah_nama_sama = (
                            Sekolah.objects
                            .filter(
                                nama__iexact=nama,
                                kecamatan__iexact=kecamatan,
                            )
                            .first()
                        )

                        if sekolah_nama_sama:
                            if (
                                sekolah_nama_sama.npsn
                                and sekolah_nama_sama.npsn != npsn
                            ):
                                raise ValueError(
                                    "nama sekolah sudah terdaftar "
                                    f"dengan NPSN {sekolah_nama_sama.npsn}"
                                )

                            sekolah = sekolah_nama_sama

                        else:
                            sekolah = Sekolah()
                            sekolah_baru = True

                    sekolah.nama = nama
                    sekolah.npsn = npsn
                    sekolah.kategori = kategori
                    sekolah.status = status
                    sekolah.kecamatan = kecamatan
                    sekolah.desa = desa
                    sekolah.alamat = alamat
                    sekolah.x_utm = x_utm
                    sekolah.y_utm = y_utm

                    sekolah.save()

                    akun_baru = siapkan_akun_operator(
                        sekolah,
                        npsn,
                    )

                npsn_diproses.add(npsn)

                if sekolah_baru:
                    total_baru += 1
                else:
                    total_update += 1

                if akun_baru:
                    total_akun_baru += 1

                if sistem_koordinat == "WGS84":
                    total_wgs84 += 1
                else:
                    total_utm += 1

            except Exception as error:
                total_gagal += 1

                properties_error = (
                    feature.get("properties") or {}
                )

                identitas_error = (
                    properties_error.get("npsn")
                    or properties_error.get("NPSN")
                    or properties_error.get("nama")
                    or properties_error.get("NAMA")
                    or f"Feature {nomor}"
                )

                daftar_gagal.append(
                    f"{identitas_error}: {error}"
                )

        messages.success(
            request,
            (
                f"Import selesai. "
                f"Sekolah baru: {total_baru}, "
                f"sekolah diperbarui: {total_update}, "
                f"akun operator baru: {total_akun_baru}, "
                f"gagal: {total_gagal}, "
                f"WGS84: {total_wgs84}, "
                f"UTM 50N: {total_utm}."
            )
        )

        if daftar_gagal:
            messages.warning(
                request,
                "Detail gagal: "
                + " | ".join(daftar_gagal[:10])
            )

        return redirect("/import-sekolah-geojson/")

    return render(
        request,
        "peta/import_sekolah_geojson.html",
    )