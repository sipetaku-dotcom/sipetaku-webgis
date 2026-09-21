from openpyxl import load_workbook
from datetime import datetime, date
from django.utils import timezone
from openpyxl import Workbook
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Q, Max
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from sekolah.forms import SekolahForm
from guru.models import (
    Pegawai,
    RiwayatPegawai,
)
from guru.forms import (
    PegawaiForm,
    RiwayatPegawaiForm,
    MutasiPegawaiForm,
    AkhiriPegawaiForm,
)
from murid.models import (
    Murid,
    TahunAjaran,
    RiwayatMurid,
)
from murid.forms import (
    MuridForm,
    TahunAjaranForm,
)
from aset.models import (
    Aset,
    AsetSekolah,
    DataRombel,
    BidangTanahSekolah,
    LegalitasTanah,
    DokumenTanah,
    FotoAset,
    JenisAset,
)
from aset.forms import (
    AsetForm,
    AsetSekolahForm,
    DataRombelForm,
    BidangTanahSekolahForm,
    LegalitasTanahForm,
    DokumenTanahForm,
)
from aset.penilaian import (
    hitung_penilaian_sekolah
)
from prestasi.models import Prestasi
from prestasi.forms import PrestasiForm
from django.contrib import messages
from .models import ProfilUser
import json
from sekolah.models import (
    Sekolah,
    KategoriSekolah,
)
from sekolah.wilayah import DATA_WILAYAH
from django.contrib.auth.models import User
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.urls import reverse


def ambil_profil_user(request):

    if not request.user.is_authenticated:
        return None

    try:
        return request.user.profil
    except:
        return None


def user_admin_kabupaten(request):

    profil = ambil_profil_user(request)

    if request.user.is_superuser:
        return True

    if profil and profil.role == 'admin_kabupaten':
        return True

    return False


def user_admin_kecamatan(request):
    profil = ambil_profil_user(request)

    if profil and profil.role == 'admin_kecamatan':
        return True

    return False


def user_operator_sekolah(request):
    profil = ambil_profil_user(request)

    if profil and profil.role == 'operator_sekolah':
        return True

    return False


def login_view(request):

    if request.method == 'POST':

        username = request.POST.get('username')
        password = request.POST.get('password')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:

            login(request, user)

            return redirect('/')

        else:

            messages.error(
                request,
                'Username atau password anda salah'
            )

    return render(
        request,
        'akun/login.html'
    )


def logout_view(request):

    logout(request)

    return redirect('/login/')


@login_required
def tambah_sekolah(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    if request.method == 'POST':
        form = SekolahForm(request.POST, request.FILES)

        if form.is_valid():
            form.save()
            messages.success(request, 'Data sekolah berhasil ditambahkan')
            return redirect('/tambah-sekolah/')

    else:
        form = SekolahForm()

    context = {
        'form': form,
        'data_wilayah': json.dumps(DATA_WILAYAH),
    }

    return render(request, 'akun/tambah_sekolah.html', context)


@login_required
def data_sekolah(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    sekolah = Sekolah.objects.all().order_by('nama')

    context = {
        'sekolah': sekolah,
    }

    return render(request, 'akun/data_sekolah.html', context)


@login_required
def edit_sekolah(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    sekolah = get_object_or_404(
        Sekolah,
        id=id
    )

    if request.method == 'POST':

        form = SekolahForm(
            request.POST,
            request.FILES,
            instance=sekolah
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                'Data sekolah berhasil diubah'
            )

            return redirect('/data-sekolah/')

    else:

        form = SekolahForm(instance=sekolah)

    context = {
        'form': form,
        'sekolah': sekolah,
        'data_wilayah': json.dumps(DATA_WILAYAH),
    }

    return render(
        request,
        'akun/edit_sekolah.html',
        context
    )


@login_required
def hapus_sekolah(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    sekolah = get_object_or_404(
        Sekolah,
        id=id
    )

    sekolah.delete()

    messages.success(
        request,
        'Data sekolah berhasil dihapus'
    )

    return redirect('/data-sekolah/')


@login_required
def tambah_user_operator(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    sekolah_list = Sekolah.objects.filter(
        user__isnull=True
    ).order_by('kecamatan', 'desa', 'nama')

    data_sekolah = []

    for sekolah in sekolah_list:
        data_sekolah.append({
            'id': sekolah.id,
            'nama': sekolah.nama,
            'npsn': sekolah.npsn,
            'kecamatan': sekolah.kecamatan,
            'desa': sekolah.desa,
        })

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')
        sekolah_id = request.POST.get('sekolah')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect('/tambah-user-operator/')

        sekolah = Sekolah.objects.get(id=sekolah_id)

        user = User.objects.create_user(
            username=username,
            password=password,
            first_name=nama_lengkap
        )

        ProfilUser.objects.create(
            user=user,
            role='operator_sekolah',
            kecamatan=sekolah.kecamatan,
            sekolah=sekolah
        )

        sekolah.user = user
        sekolah.save()

        messages.success(request, 'User operator sekolah berhasil dibuat')
        return redirect('/tambah-user-operator/')

    context = {
        'data_wilayah': json.dumps(DATA_WILAYAH),
        'data_sekolah': json.dumps(data_sekolah),
    }

    return render(request, 'akun/tambah_user_operator.html', context)


@login_required
def data_operator(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    operator = ProfilUser.objects.filter(
        role='operator_sekolah'
    ).select_related('user', 'sekolah').order_by('sekolah__nama')

    context = {
        'operator': operator,
    }

    return render(request, 'akun/data_operator.html', context)


@login_required
def tambah_user_admin_kecamatan(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')
        kecamatan = request.POST.get('kecamatan')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect('/tambah-user-admin-kecamatan/')

        user = User.objects.create_user(
            username=username,
            password=password,
            first_name=nama_lengkap
        )

        ProfilUser.objects.create(
            user=user,
            role='admin_kecamatan',
            kecamatan=kecamatan
        )

        messages.success(request, 'User admin kecamatan berhasil dibuat')
        return redirect('/tambah-user-admin-kecamatan/')

    context = {
        'data_wilayah': DATA_WILAYAH,
    }

    return render(request, 'akun/tambah_user_admin_kecamatan.html', context)


@login_required
def data_admin_kecamatan(request):

    if not user_admin_kabupaten(request):
        return redirect('/')

    admin_kecamatan = ProfilUser.objects.filter(
        role='admin_kecamatan'
    ).select_related('user').order_by('kecamatan')

    context = {
        'admin_kecamatan': admin_kecamatan,
    }

    return render(request, 'akun/data_admin_kecamatan.html', context)


@login_required
def tambah_user_admin_kabupaten(request):

    if not request.user.is_superuser:
        return redirect('/')

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect('/tambah-user-admin-kabupaten/')

        user = User.objects.create_user(
            username=username,
            password=password,
            first_name=nama_lengkap
        )

        ProfilUser.objects.create(
            user=user,
            role='admin_kabupaten'
        )

        messages.success(request, 'User admin kabupaten berhasil dibuat')
        return redirect('/tambah-user-admin-kabupaten/')

    return render(request, 'akun/tambah_user_admin_kabupaten.html')


@login_required
def edit_operator(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    profil_operator = ProfilUser.objects.get(
        id=id,
        role='operator_sekolah'
    )

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')

        user = profil_operator.user

        if User.objects.filter(username=username).exclude(id=user.id).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect(f'/edit-operator/{id}/')

        user.first_name = nama_lengkap
        user.username = username

        if password:
            user.set_password(password)

        user.save()

        messages.success(request, 'Data operator berhasil diubah')
        return redirect('/data-operator/')

    context = {
        'operator': profil_operator,
    }

    return render(request, 'akun/edit_operator.html', context)


@login_required
def hapus_operator(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    profil_operator = ProfilUser.objects.get(
        id=id,
        role='operator_sekolah'
    )

    sekolah = profil_operator.sekolah
    user = profil_operator.user

    if sekolah:
        sekolah.user = None
        sekolah.save()

    user.delete()

    messages.success(request, 'User operator berhasil dihapus')
    return redirect('/data-operator/')


@login_required
def edit_admin_kecamatan(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    profil_admin = ProfilUser.objects.get(
        id=id,
        role='admin_kecamatan'
    )

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')

        user = profil_admin.user

        if User.objects.filter(username=username).exclude(id=user.id).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect(f'/edit-admin-kecamatan/{id}/')

        user.first_name = nama_lengkap
        user.username = username

        if password:
            user.set_password(password)

        user.save()

        messages.success(request, 'Data admin kecamatan berhasil diubah')
        return redirect('/data-admin-kecamatan/')

    context = {
        'admin_kecamatan': profil_admin,
    }

    return render(request, 'akun/edit_admin_kecamatan.html', context)


@login_required
def hapus_admin_kecamatan(request, id):

    if not user_admin_kabupaten(request):
        return redirect('/')

    profil_admin = ProfilUser.objects.get(
        id=id,
        role='admin_kecamatan'
    )

    user = profil_admin.user
    user.delete()

    messages.success(request, 'User admin kecamatan berhasil dihapus')
    return redirect('/data-admin-kecamatan/')


@login_required
def data_admin_kabupaten(request):

    if not request.user.is_superuser:
        return redirect('/')

    admin_kabupaten = ProfilUser.objects.filter(
        role='admin_kabupaten',
        user__is_superuser=False
    ).select_related('user').order_by('user__first_name')

    context = {
        'admin_kabupaten': admin_kabupaten,
    }

    return render(request, 'akun/data_admin_kabupaten.html', context)


@login_required
def edit_admin_kabupaten(request, id):

    if not request.user.is_superuser:
        return redirect('/')

    profil_admin = ProfilUser.objects.get(
        id=id,
        role='admin_kabupaten'
    )

    if request.method == 'POST':

        nama_lengkap = request.POST.get('nama_lengkap')
        username = request.POST.get('username')
        password = request.POST.get('password')

        user = profil_admin.user

        if User.objects.filter(username=username).exclude(id=user.id).exists():
            messages.error(request, 'Username sudah digunakan')
            return redirect(f'/edit-admin-kabupaten/{id}/')

        user.first_name = nama_lengkap
        user.username = username

        if password:
            user.set_password(password)

        user.save()

        messages.success(request, 'Data admin kabupaten berhasil diubah')
        return redirect('/data-admin-kabupaten/')

    context = {
        'admin_kabupaten': profil_admin,
    }

    return render(request, 'akun/edit_admin_kabupaten.html', context)


@login_required
def hapus_admin_kabupaten(request, id):

    if not request.user.is_superuser:
        return redirect('/')

    profil_admin = ProfilUser.objects.get(
        id=id,
        role='admin_kabupaten'
    )

    user = profil_admin.user
    user.delete()

    messages.success(request, 'User admin kabupaten berhasil dihapus')
    return redirect('/data-admin-kabupaten/')


@login_required
def dashboard_operator(request):

    # Pastikan hanya operator sekolah
    if not user_operator_sekolah(request):

        messages.error(
            request,
            'Anda tidak memiliki akses ke dashboard sekolah.'
        )

        return redirect('/')

    # Ambil sekolah milik operator
    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:

        return render(
            request,
            'akun/belum_terhubung.html'
        )

    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(aktif=True)
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')

    # =====================================================
    # KESIAPAN TAHUN AJARAN BERIKUTNYA
    # =====================================================

    kesiapan_tahun_ajaran = None
    tahun_asal_kesiapan = tahun_ajaran
    tahun_tujuan_kesiapan = None

    # Cari Tahun Ajaran yang menjadikan
    # tahun aktif sebagai tahun sebelumnya.
    tahun_tujuan_kesiapan = (
        TahunAjaran.objects
        .filter(
            tahun_sebelumnya=tahun_ajaran
        )
        .order_by('nama')
        .first()
    )

    if tahun_tujuan_kesiapan:

        kesiapan_tahun_ajaran = (
            hitung_kesiapan_tahun_ajaran(
                sekolah,
                tahun_asal_kesiapan,
                tahun_tujuan_kesiapan
            )
        )

    # =====================================================
    # PENILAIAN SEKOLAH
    # =====================================================

    penilaian = hitung_penilaian_sekolah(
        sekolah
    )

    # =====================================================
    # RINGKASAN SEKOLAH
    # =====================================================

    pegawai_sekolah = (
        RiwayatPegawai.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran,
            status_akhir='AKTIF',
            jenis_pegawai='PENDIDIK'
        )
        .select_related(
            'pegawai'
        )
    )

    total_guru = (
        pegawai_sekolah.count()
    )

    total_guru_laki_laki = (
        pegawai_sekolah
        .filter(
            pegawai__jenis_kelamin='L'
        )
        .count()
    )

    total_guru_perempuan = (
        pegawai_sekolah
        .filter(
            pegawai__jenis_kelamin='P'
        )
        .count()
    )

    murid_aktif = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            tahun_ajaran=tahun_ajaran,
            status_akhir='AKTIF'
        )
    )

    total_murid = murid_aktif.count()

    total_laki_laki = murid_aktif.filter(
        murid__jenis_kelamin='L'
    ).count()

    total_perempuan = murid_aktif.filter(
        murid__jenis_kelamin='P'
    ).count()

    total_siswa_baru = murid_aktif.filter(
        status_masuk='SISWA_BARU'
    ).count()

    total_pindahan = murid_aktif.filter(
        status_masuk='PINDAHAN'
    ).count()

    total_prestasi = Prestasi.objects.filter(
        sekolah=sekolah
    ).count()

    prestasi_terbaru = (
        Prestasi.objects
        .filter(
            sekolah=sekolah
        )
        .select_related(
            'murid'
        )
        .order_by(
            '-tahun',
            '-id'
        )[:5]
    )

    # =====================================================
    # DATA ASET SEKOLAH
    # =====================================================

    jenis_aset = (
        JenisAset.objects
        .filter(aktif=True)
    )

    data_aset_sekolah = list(
        AsetSekolah.objects
        .filter(
            sekolah=sekolah,
            sudah_diisi=True
        )
        .select_related(
            'jenis_aset'
        )
    )

    # =====================================================
    # KELENGKAPAN DATA
    # =====================================================

    jumlah_master = jenis_aset.count()

    jumlah_terisi = len(
        data_aset_sekolah
    )

    if jumlah_master > 0:

        persentase_pengisian = round(
            (
                jumlah_terisi
                / jumlah_master
            ) * 100
        )

    else:

        persentase_pengisian = 0

    # =====================================================
    # KONDISI ASET
    # =====================================================

    jumlah_ada = sum(
        1
        for data in data_aset_sekolah
        if data.ketersediaan == 'ADA'
    )

    jumlah_tidak_ada = sum(
        1
        for data in data_aset_sekolah
        if data.ketersediaan == 'TIDAK_ADA'
    )

    jumlah_rusak_ringan = sum(
        data.jumlah_rusak_ringan
        for data in data_aset_sekolah
    )

    jumlah_rusak_berat = sum(
        data.jumlah_rusak_berat
        for data in data_aset_sekolah
    )

    # =====================================================
    # DATA TANAH
    # =====================================================

    daftar_bidang_tanah = list(
        BidangTanahSekolah.objects
        .filter(
            sekolah=sekolah
        )
        .prefetch_related(
            'legalitas'
        )
        .order_by(
            'nomor_bidang'
        )
    )


    # =====================================================
    # JUMLAH BIDANG
    # =====================================================

    jumlah_bidang_tanah = (
        len(
            daftar_bidang_tanah
        )
    )


    # =====================================================
    # TOTAL LUAS TANAH
    #
    # Luas dihitung berdasarkan bidang,
    # bukan berdasarkan jumlah dokumen legalitas.
    # =====================================================

    total_luas_tanah = sum(
        (
            bidang.luas_bidang_m2
            or 0
        )
        for bidang in daftar_bidang_tanah
    )


    # =====================================================
    # BIDANG YANG MEMILIKI LEGALITAS
    # =====================================================

    jumlah_bidang_berlegalitas = sum(
        1
        for bidang in daftar_bidang_tanah
        if bidang.legalitas.all()
    )


    # =====================================================
    # BIDANG BELUM MEMILIKI LEGALITAS
    # =====================================================

    jumlah_bidang_belum_berlegalitas = (
        jumlah_bidang_tanah
        - jumlah_bidang_berlegalitas
    )


    # =====================================================
    # STATUS LEGALITAS
    # =====================================================

    if jumlah_bidang_tanah == 0:

        status_legalitas_tanah = (
            'Data Tanah Belum Diisi'
        )

        warna_legalitas_tanah = (
            'secondary'
        )


    elif jumlah_bidang_berlegalitas == 0:

        status_legalitas_tanah = (
            'Belum Ada Legalitas'
        )

        warna_legalitas_tanah = (
            'danger'
        )


    elif (
        jumlah_bidang_berlegalitas
        < jumlah_bidang_tanah
    ):

        status_legalitas_tanah = (
            'Legalitas Belum Lengkap'
        )

        warna_legalitas_tanah = (
            'warning'
        )


    else:

        status_legalitas_tanah = (
            'Legalitas Lengkap'
        )

        warna_legalitas_tanah = (
            'success'
        )

    # =====================================================
    # DATA ROMBEL
    # =====================================================

    daftar_rombel = list(
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )
        .order_by(
            'nama_kelas'
        )
    )

    data_siswa_per_rombel = []

    for data in daftar_rombel:

        data_siswa_per_rombel.append({
            'nama_kelas': data.nama_kelas,
            'jumlah_siswa': data.jumlah_siswa_aktif,
            'kapasitas': data.kapasitas_total,
            'kekurangan': data.kekurangan_kapasitas,
        })

    total_rombel = sum(
        data.jumlah_rombel
        for data in daftar_rombel
    )

    total_siswa_rombel = sum(
        data.jumlah_siswa_aktif
        for data in daftar_rombel
    )

    total_kapasitas = sum(
        data.kapasitas_total
        for data in daftar_rombel
    )

    total_kekurangan = sum(
        data.kekurangan_kapasitas
        for data in daftar_rombel
    )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'sekolah': sekolah,
        'tahun_ajaran': tahun_ajaran,
        'kesiapan_tahun_ajaran':
            kesiapan_tahun_ajaran,

        'tahun_asal_kesiapan':
            tahun_asal_kesiapan,

        'tahun_tujuan_kesiapan':
            tahun_tujuan_kesiapan,

        # =========================
        # RINGKASAN SEKOLAH
        # =========================

        'total_guru': total_guru,
        'total_guru_laki_laki': total_guru_laki_laki,
        'total_guru_perempuan': total_guru_perempuan,
        'total_murid': total_murid,
        'total_prestasi': total_prestasi,
        'prestasi_terbaru': prestasi_terbaru,
        'total_laki_laki': total_laki_laki,
        'total_perempuan': total_perempuan,
        'total_siswa_baru': total_siswa_baru,
        'total_pindahan': total_pindahan,

        # =========================
        # PENILAIAN
        # =========================

        'penilaian': penilaian,

        # =========================
        # KELENGKAPAN ASET
        # =========================

        'jumlah_master': jumlah_master,
        'jumlah_terisi': jumlah_terisi,
        'persentase_pengisian': persentase_pengisian,

        # =========================
        # KONDISI ASET
        # =========================

        'jumlah_ada': jumlah_ada,
        'jumlah_tidak_ada': jumlah_tidak_ada,
        'jumlah_rusak_ringan': jumlah_rusak_ringan,
        'jumlah_rusak_berat': jumlah_rusak_berat,

        # =========================
        # TANAH & LEGALITAS
        # =========================

        'daftar_bidang_tanah':
            daftar_bidang_tanah,

        'jumlah_bidang_tanah':
            jumlah_bidang_tanah,

        'total_luas_tanah':
            total_luas_tanah,

        'jumlah_bidang_berlegalitas':
            jumlah_bidang_berlegalitas,

        'jumlah_bidang_belum_berlegalitas':
            jumlah_bidang_belum_berlegalitas,

        'status_legalitas_tanah':
            status_legalitas_tanah,

        'warna_legalitas_tanah':
            warna_legalitas_tanah,

        # =========================
        # ROMBEL
        # =========================

        'total_rombel': total_rombel,
        'total_siswa_rombel': total_siswa_rombel,
        'total_kapasitas': total_kapasitas,
        'total_kekurangan': total_kekurangan,
        'data_siswa_per_rombel': data_siswa_per_rombel,
    }

    return render(
        request,
        'akun/dashboard_operator.html',
        context
    )


@login_required
def profil_sekolah(request):

    try:

        sekolah = ambil_sekolah_operator(request)

        if sekolah is None:
            return redirect('/')

    except:
        return render(request, 'akun/belum_terhubung.html')

    if request.method == 'POST':
        form = SekolahForm(
            request.POST,
            request.FILES,
            instance=sekolah
        )

        if form.is_valid():
            form.save()
            return redirect('/profil-sekolah/')

    else:
        form = SekolahForm(instance=sekolah)

    context = {
        'form': form,
        'sekolah': sekolah,
        'data_wilayah': json.dumps(DATA_WILAYAH),
    }

    return render(request, 'akun/profil_sekolah.html', context)


@login_required
def data_murid(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # TAHUN AJARAN
    # =====================================================

    tahun_ajaran_aktif = (
        TahunAjaran.objects
        .filter(aktif=True)
        .first()
    )

    if tahun_ajaran_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')


    # =====================================================
    # FILTER TAHUN AJARAN
    # =====================================================

    tahun_id = request.GET.get(
        'tahun'
    )

    if tahun_id:

        tahun_ajaran = (
            TahunAjaran.objects
            .filter(
                id=tahun_id
            )
            .first()
        )

        if tahun_ajaran is None:
            tahun_ajaran = tahun_ajaran_aktif

    else:

        tahun_ajaran = tahun_ajaran_aktif


    # =====================================================
    # ROMBEL SEKOLAH
    # =====================================================

    daftar_rombel = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )
        .order_by(
            'nama_kelas'
        )
    )


    # =====================================================
    # FILTER KELAS
    # =====================================================

    rombel_id = request.GET.get(
        'rombel'
    )


    # =====================================================
    # RIWAYAT MURID
    # =====================================================

    riwayat_murid = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )
        .select_related(
            'murid',
            'rombel',
            'tahun_ajaran'
        )
    )


    if rombel_id:

        riwayat_murid = (
            riwayat_murid
            .filter(
                rombel_id=rombel_id
            )
        )


    riwayat_murid = (
        riwayat_murid
        .order_by(
            'rombel__nama_kelas',
            'murid__nama'
        )
    )


    # =====================================================
    # RINGKASAN
    # =====================================================

    total_siswa = (
        riwayat_murid
        .filter(
            status_akhir='AKTIF'
        )
        .count()
    )


    total_laki_laki = (
        riwayat_murid
        .filter(
            status_akhir='AKTIF',
            murid__jenis_kelamin='L'
        )
        .count()
    )


    total_perempuan = (
        riwayat_murid
        .filter(
            status_akhir='AKTIF',
            murid__jenis_kelamin='P'
        )
        .count()
    )


    total_siswa_baru = (
        riwayat_murid
        .filter(
            status_masuk='SISWA_BARU'
        )
        .count()
    )


    total_pindahan = (
        riwayat_murid
        .filter(
            status_masuk='PINDAHAN'
        )
        .count()
    )


    # =====================================================
    # SEMUA TAHUN AJARAN
    # =====================================================

    daftar_tahun_ajaran = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )


    context = {

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,

        'tahun_ajaran_aktif':
            tahun_ajaran_aktif,

        'daftar_tahun_ajaran':
            daftar_tahun_ajaran,

        'daftar_rombel':
            daftar_rombel,

        'riwayat_murid':
            riwayat_murid,

        'rombel_terpilih':
            rombel_id,

        'total_siswa':
            total_siswa,

        'total_laki_laki':
            total_laki_laki,

        'total_perempuan':
            total_perempuan,

        'total_siswa_baru':
            total_siswa_baru,

        'total_pindahan':
            total_pindahan,
    }


    return render(
        request,
        'akun/data_murid.html',
        context
    )


@login_required
def tambah_murid(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # Tahun ajaran aktif
    tahun_ajaran = (
        TahunAjaran.objects
        .filter(aktif=True)
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_murid'
        )


    if request.method == 'POST':

        form = MuridForm(
            request.POST,
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )

        if form.is_valid():

            with transaction.atomic():

                murid = form.save(
                    commit=False
                )

                murid.sekolah = sekolah

                rombel = (
                    form.cleaned_data[
                        'rombel'
                    ]
                )

                status_masuk = (
                    form.cleaned_data[
                        'status_masuk'
                    ]
                )

                # Field kelas lama kita sinkronkan
                # sementara untuk kompatibilitas
                murid.kelas = (
                    rombel.nama_kelas
                )

                murid.save()


                # Buat riwayat akademik
                riwayat_baru = RiwayatMurid(
                    murid=murid,
                    rombel=rombel,
                    tahun_ajaran=tahun_ajaran,
                    status_masuk=status_masuk,
                    status_akhir='AKTIF'
                )

                riwayat_baru.full_clean()
                riwayat_baru.save()


            messages.success(
                request,
                (
                    f'{murid.nama} berhasil '
                    f'ditambahkan ke '
                    f'{rombel.nama_kelas}.'
                )
            )

            return redirect(
                'data_murid'
            )

    else:

        form = MuridForm(
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )


    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,
    }


    return render(
        request,
        'akun/tambah_murid.html',
        context
    )


@login_required
def data_aset(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data aset sekolah.'
        )

        return redirect('/')

    tab_aktif = request.GET.get(
        'tab',
        'bangunan'
    ).lower()

    if tab_aktif not in {
        'bangunan',
        'sanitasi',
        'perlengkapan',
    }:
        tab_aktif = 'bangunan'

    jenis_aset = (
        JenisAset.objects
        .filter(aktif=True)
        .order_by(
            'kategori',
            'urutan',
            'nama'
        )
    )

    data_aset_sekolah = list(
        AsetSekolah.objects
        .filter(
            sekolah=sekolah,
            sudah_diisi=True
        )
        .select_related(
            'jenis_aset'
        )
        .prefetch_related(
            'foto_dokumentasi'
        )
    )

    data_aset_per_jenis = {
        data.jenis_aset_id: data
        for data in data_aset_sekolah
    }

    bangunan = []
    sanitasi = []
    perlengkapan = []

    for jenis in jenis_aset:

        jenis.data_sekolah_aktif = (
            data_aset_per_jenis.get(
                jenis.id
            )
        )

        if jenis.kategori == 'BANGUNAN':

            bangunan.append(
                jenis
            )

        elif jenis.kategori == 'SANITASI':

            sanitasi.append(
                jenis
            )

        elif (
            jenis.kategori
            == 'PERLENGKAPAN'
        ):

            perlengkapan.append(
                jenis
            )

    jumlah_master = len(
        jenis_aset
    )

    jumlah_terisi = len(
        data_aset_per_jenis
    )

    if jumlah_master > 0:

        persentase_pengisian = round(
            (
                jumlah_terisi
                / jumlah_master
            )
            * 100
        )

    else:

        persentase_pengisian = 0

    jumlah_ada = sum(
        1
        for data in data_aset_sekolah
        if data.ketersediaan == 'ADA'
    )

    jumlah_tidak_ada = sum(
        1
        for data in data_aset_sekolah
        if data.ketersediaan == 'TIDAK_ADA'
    )

    jumlah_rusak_ringan = sum(
        data.jumlah_rusak_ringan
        for data in data_aset_sekolah
    )

    jumlah_rusak_berat = sum(
        data.jumlah_rusak_berat
        for data in data_aset_sekolah
    )

    penilaian = (
        hitung_penilaian_sekolah(
            sekolah
        )
    )

    context = {
        'sekolah': sekolah,

        'bangunan':
            bangunan,

        'sanitasi':
            sanitasi,

        'perlengkapan':
            perlengkapan,

        'jumlah_master':
            jumlah_master,

        'jumlah_terisi':
            jumlah_terisi,

        'persentase_pengisian':
            persentase_pengisian,

        'jumlah_ada':
            jumlah_ada,

        'jumlah_tidak_ada':
            jumlah_tidak_ada,

        'jumlah_rusak_ringan':
            jumlah_rusak_ringan,

        'jumlah_rusak_berat':
            jumlah_rusak_berat,
        
        'penilaian': penilaian,

        'tab_aktif': tab_aktif,
    }

    return render(
        request,
        'akun/data_aset.html',
        context
    )


@login_required
def prioritas_bantuan_sekolah(
    request
):

    profil = ambil_profil_user(
        request
    )

    if user_admin_kabupaten(request):

        daftar_sekolah = (
            Sekolah.objects
            .all()
            .select_related(
                'kategori'
            )
            .order_by('nama')
        )

    elif user_admin_kecamatan(request):

        daftar_sekolah = (
            Sekolah.objects
            .filter(
                kecamatan=(
                    profil.kecamatan
                )
            )
            .select_related(
                'kategori'
            )
            .order_by('nama')
        )

    else:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke halaman prioritas bantuan.'
        )

        return redirect('/')

    hasil_final = []
    hasil_belum_lengkap = []

    for sekolah in daftar_sekolah:

        penilaian = (
            hitung_penilaian_sekolah(
                sekolah
            )
        )

        data = {
            'sekolah': sekolah,
            'penilaian': penilaian,
        }

        if penilaian['penilaian_final']:

            hasil_final.append(
                data
            )

        else:

            hasil_belum_lengkap.append(
                data
            )

    hasil_final.sort(
        key=lambda data: (
            data['penilaian']
            ['skor_prioritas']
            or 0
        ),
        reverse=True
    )

    hasil_belum_lengkap.sort(
        key=lambda data: (
            data['penilaian']
            ['progres']
        ),
        reverse=True
    )

    jumlah_sekolah = (
        len(hasil_final)
        + len(hasil_belum_lengkap)
    )

    jumlah_sangat_mendesak = sum(
        1
        for data in hasil_final
        if (
            data['penilaian']
            ['skor_prioritas']
            >= 80
        )
    )

    jumlah_mendesak = sum(
        1
        for data in hasil_final
        if (
            60
            <=
            data['penilaian']
            ['skor_prioritas']
            < 80
        )
    )

    jumlah_legalitas_belum_aman = sum(
        1
        for data in (
            hasil_final
            + hasil_belum_lengkap
        )
        if not (
            data['penilaian']
            ['legalitas_aman']
        )
    )

    context = {
        'hasil_final':
            hasil_final,

        'hasil_belum_lengkap':
            hasil_belum_lengkap,

        'jumlah_sekolah':
            jumlah_sekolah,

        'jumlah_sudah_dinilai':
            len(hasil_final),

        'jumlah_belum_lengkap':
            len(hasil_belum_lengkap),

        'jumlah_sangat_mendesak':
            jumlah_sangat_mendesak,

        'jumlah_mendesak':
            jumlah_mendesak,

        'jumlah_legalitas_belum_aman':
            jumlah_legalitas_belum_aman,
    }

    return render(
        request,
        'akun/prioritas_bantuan_sekolah.html',
        context
    )


@login_required
def detail_prioritas_sekolah(
    request,
    id
):

    profil = ambil_profil_user(
        request
    )

    if user_admin_kabupaten(request):

        sekolah = get_object_or_404(
            Sekolah.objects.select_related(
                'kategori'
            ),
            id=id
        )

    elif user_admin_kecamatan(request):

        sekolah = get_object_or_404(
            Sekolah.objects.select_related(
                'kategori'
            ),
            id=id,
            kecamatan=profil.kecamatan
        )

    else:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke detail penilaian sekolah.'
        )

        return redirect('/')

    penilaian = (
        hitung_penilaian_sekolah(
            sekolah
        )
    )

    daftar_aset = list(
        AsetSekolah.objects
        .filter(
            sekolah=sekolah,
            sudah_diisi=True
        )
        .select_related(
            'jenis_aset'
        )
        .prefetch_related(
            'foto_dokumentasi'
        )
        .order_by(
            'jenis_aset__kategori',
            'jenis_aset__urutan'
        )
    )

    bangunan = []
    sanitasi = []
    perlengkapan = []

    for data in daftar_aset:

        if (
            data.jenis_aset.kategori
            == 'BANGUNAN'
        ):

            bangunan.append(data)

        elif (
            data.jenis_aset.kategori
            == 'SANITASI'
        ):

            sanitasi.append(data)

        elif (
            data.jenis_aset.kategori
            == 'PERLENGKAPAN'
        ):

            perlengkapan.append(data)

    # =====================================================
    # DATA TANAH & LEGALITAS
    # =====================================================

    daftar_bidang_tanah = list(
        BidangTanahSekolah.objects
        .filter(
            sekolah=sekolah
        )
        .prefetch_related(
            'legalitas',
            'legalitas__lampiran'
        )
        .order_by(
            'nomor_bidang'
        )
    )


    jumlah_bidang_tanah = (
        len(
            daftar_bidang_tanah
        )
    )


    total_luas_tanah = sum(
        (
            bidang.luas_bidang_m2
            or 0
        )
        for bidang in daftar_bidang_tanah
    )


    jumlah_bidang_berlegalitas = sum(
        1
        for bidang in daftar_bidang_tanah
        if bidang.legalitas.all()
    )


    jumlah_bidang_belum_berlegalitas = (
        jumlah_bidang_tanah
        - jumlah_bidang_berlegalitas
    )

    daftar_rombel = list(
        DataRombel.objects
        .filter(sekolah=sekolah)
        .order_by('nama_kelas')
    )

    context = {
        'sekolah':
            sekolah,

        'penilaian':
            penilaian,

        'bangunan':
            bangunan,

        'sanitasi':
            sanitasi,

        'perlengkapan':
            perlengkapan,

        'daftar_bidang_tanah':
            daftar_bidang_tanah,

        'jumlah_bidang_tanah':
            jumlah_bidang_tanah,

        'total_luas_tanah':
            total_luas_tanah,

        'jumlah_bidang_berlegalitas':
            jumlah_bidang_berlegalitas,

        'jumlah_bidang_belum_berlegalitas':
            jumlah_bidang_belum_berlegalitas,

        'daftar_rombel':
            daftar_rombel,
    }

    return render(
        request,
        'akun/detail_prioritas_sekolah.html',
        context
    )


@login_required
def tanah_legalitas_sekolah(
    request
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data tanah sekolah.'
        )

        return redirect('/')


    # =====================================================
    # DAFTAR BIDANG TANAH
    # =====================================================

    daftar_bidang = (
        BidangTanahSekolah.objects
        .filter(
            sekolah=sekolah
        )
        .prefetch_related(
            'legalitas',
            'legalitas__lampiran'
        )
        .order_by(
            'nomor_bidang'
        )
    )


    # =====================================================
    # RINGKASAN
    # =====================================================

    jumlah_bidang = (
        daftar_bidang.count()
    )

    total_luas = sum(
        (
            bidang.luas_bidang_m2
            or 0
        )
        for bidang in daftar_bidang
    )


    jumlah_berlegalitas = sum(
        1
        for bidang in daftar_bidang
        if bidang.legalitas.exists()
    )


    jumlah_belum_berlegalitas = (
        jumlah_bidang
        - jumlah_berlegalitas
    )


    # =====================================================
    # STATUS LEGALITAS KESELURUHAN
    # =====================================================

    if jumlah_bidang == 0:

        status_legalitas = (
            'Data tanah belum diisi'
        )

        status_warna = 'secondary'


    elif jumlah_berlegalitas == 0:

        status_legalitas = (
            'Belum Ada Legalitas'
        )

        status_warna = 'danger'


    elif (
        jumlah_berlegalitas
        < jumlah_bidang
    ):

        status_legalitas = (
            'Legalitas Belum Lengkap'
        )

        status_warna = 'warning'


    else:

        status_legalitas = (
            'Legalitas Lengkap'
        )

        status_warna = 'success'


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'sekolah':
            sekolah,

        'daftar_bidang':
            daftar_bidang,

        'jumlah_bidang':
            jumlah_bidang,

        'total_luas':
            total_luas,

        'jumlah_berlegalitas':
            jumlah_berlegalitas,

        'jumlah_belum_berlegalitas':
            jumlah_belum_berlegalitas,

        'status_legalitas':
            status_legalitas,

        'status_warna':
            status_warna,
    }


    return render(
        request,
        'akun/tanah_legalitas_sekolah.html',
        context
    )


@login_required
def tambah_bidang_tanah(
    request
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data tanah sekolah.'
        )

        return redirect('/')


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = BidangTanahSekolahForm(
            request.POST
        )


        if form.is_valid():

            with transaction.atomic():

                # =========================================
                # NOMOR BIDANG OTOMATIS
                # =========================================

                nomor_terakhir = (
                    BidangTanahSekolah.objects
                    .filter(
                        sekolah=sekolah
                    )
                    .aggregate(
                        terbesar=Max(
                            'nomor_bidang'
                        )
                    )[
                        'terbesar'
                    ]
                    or 0
                )


                bidang = form.save(
                    commit=False
                )

                bidang.sekolah = sekolah

                bidang.nomor_bidang = (
                    nomor_terakhir + 1
                )

                bidang.full_clean()

                bidang.save()


            messages.success(
                request,
                (
                    f'Bidang Tanah '
                    f'{bidang.nomor_bidang} '
                    'berhasil ditambahkan.'
                )
            )


            return redirect(
                'tanah_legalitas_sekolah'
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = (
            BidangTanahSekolahForm()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

    }


    return render(
        request,
        'akun/tambah_bidang_tanah.html',
        context
    )


@login_required
def detail_bidang_tanah(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data tanah sekolah.'
        )

        return redirect('/')


    # =====================================================
    # BIDANG TANAH
    # =====================================================

    bidang = get_object_or_404(
        BidangTanahSekolah.objects
        .prefetch_related(
            'legalitas',
            'legalitas__lampiran'
        ),
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # DAFTAR LEGALITAS
    # =====================================================

    daftar_legalitas = (
        bidang.legalitas
        .all()
        .order_by(
            '-tanggal_dokumen',
            '-id'
        )
    )


    jumlah_legalitas = (
        daftar_legalitas.count()
    )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'sekolah':
            sekolah,

        'bidang':
            bidang,

        'daftar_legalitas':
            daftar_legalitas,

        'jumlah_legalitas':
            jumlah_legalitas,
    }


    return render(
        request,
        'akun/detail_bidang_tanah.html',
        context
    )


@login_required
def edit_bidang_tanah(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data tanah sekolah.'
        )

        return redirect('/')


    # =====================================================
    # BIDANG TANAH
    # =====================================================

    bidang = get_object_or_404(
        BidangTanahSekolah,
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = BidangTanahSekolahForm(
            request.POST,
            instance=bidang
        )


        if form.is_valid():

            bidang = form.save(
                commit=False
            )

            bidang.sekolah = sekolah

            bidang.full_clean()

            bidang.save()


            messages.success(
                request,
                (
                    f'Bidang Tanah '
                    f'{bidang.nomor_bidang} '
                    'berhasil diperbarui.'
                )
            )


            return redirect(
                'detail_bidang_tanah',
                id=bidang.id
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = BidangTanahSekolahForm(
            instance=bidang
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'bidang':
            bidang,
    }


    return render(
        request,
        'akun/edit_bidang_tanah.html',
        context
    )


@login_required
def hapus_bidang_tanah(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data tanah sekolah.'
        )

        return redirect('/')


    # =====================================================
    # BIDANG TANAH
    # =====================================================

    bidang = get_object_or_404(
        BidangTanahSekolah,
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # HANYA IZINKAN POST
    # =====================================================

    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan hapus tidak valid.'
        )

        return redirect(
            'detail_bidang_tanah',
            id=bidang.id
        )


    # =====================================================
    # SIMPAN INFORMASI SEBELUM DIHAPUS
    # =====================================================

    nomor_bidang = (
        bidang.nomor_bidang
    )

    jumlah_legalitas = (
        bidang.legalitas.count()
    )


    # =====================================================
    # HAPUS BIDANG
    # =====================================================

    with transaction.atomic():

        bidang.delete()


    # =====================================================
    # PESAN
    # =====================================================

    if jumlah_legalitas > 0:

        messages.success(
            request,
            (
                f'Bidang Tanah {nomor_bidang} '
                f'beserta {jumlah_legalitas} '
                'data legalitas berhasil dihapus.'
            )
        )

    else:

        messages.success(
            request,
            (
                f'Bidang Tanah {nomor_bidang} '
                'berhasil dihapus.'
            )
        )


    return redirect(
        'tanah_legalitas_sekolah'
    )


@login_required
def tambah_legalitas_tanah(
    request,
    bidang_id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data legalitas tanah.'
        )

        return redirect('/')


    # =====================================================
    # BIDANG TANAH
    # =====================================================

    bidang = get_object_or_404(
        BidangTanahSekolah,
        id=bidang_id,
        sekolah=sekolah
    )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = LegalitasTanahForm(
            request.POST
        )


        if form.is_valid():

            with transaction.atomic():

                legalitas = form.save(
                    commit=False
                )

                legalitas.bidang_tanah = (
                    bidang
                )

                legalitas.full_clean()

                legalitas.save()


            messages.success(
                request,
                (
                    'Data legalitas untuk '
                    f'Bidang Tanah '
                    f'{bidang.nomor_bidang} '
                    'berhasil ditambahkan.'
                )
            )


            return redirect(
                'detail_bidang_tanah',
                id=bidang.id
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = (
            LegalitasTanahForm()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'bidang':
            bidang,

    }


    return render(
        request,
        'akun/tambah_legalitas_tanah.html',
        context
    )


@login_required
def edit_legalitas_tanah(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data legalitas tanah.'
        )

        return redirect('/')


    # =====================================================
    # LEGALITAS TANAH
    # =====================================================

    legalitas = get_object_or_404(
        LegalitasTanah.objects
        .select_related(
            'bidang_tanah',
            'bidang_tanah__sekolah'
        ),
        id=id,
        bidang_tanah__sekolah=sekolah
    )


    bidang = (
        legalitas.bidang_tanah
    )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = LegalitasTanahForm(
            request.POST,
            instance=legalitas
        )


        if form.is_valid():

            legalitas = form.save(
                commit=False
            )

            # Pastikan legalitas tetap berada
            # pada bidang yang sama
            legalitas.bidang_tanah = (
                bidang
            )

            legalitas.full_clean()

            legalitas.save()


            messages.success(
                request,
                (
                    'Data legalitas '
                    f'Bidang Tanah '
                    f'{bidang.nomor_bidang} '
                    'berhasil diperbarui.'
                )
            )


            return redirect(
                'detail_bidang_tanah',
                id=bidang.id
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = LegalitasTanahForm(
            instance=legalitas
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'bidang':
            bidang,

        'legalitas':
            legalitas,

    }


    return render(
        request,
        'akun/edit_legalitas_tanah.html',
        context
    )


@login_required
def hapus_legalitas_tanah(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data legalitas tanah.'
        )

        return redirect('/')


    # =====================================================
    # LEGALITAS TANAH
    # =====================================================

    legalitas = get_object_or_404(
        LegalitasTanah.objects
        .select_related(
            'bidang_tanah',
            'bidang_tanah__sekolah'
        )
        .prefetch_related(
            'lampiran'
        ),
        id=id,
        bidang_tanah__sekolah=sekolah
    )


    bidang = (
        legalitas.bidang_tanah
    )


    # =====================================================
    # HANYA POST
    # =====================================================

    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan hapus tidak valid.'
        )

        return redirect(
            'detail_bidang_tanah',
            id=bidang.id
        )


    # =====================================================
    # SIMPAN INFORMASI
    # =====================================================

    nama_legalitas = (
        legalitas.get_jenis_dokumen_display()
    )


    if legalitas.nomor_dokumen:

        nama_legalitas = (
            f'{nama_legalitas} '
            f'{legalitas.nomor_dokumen}'
        )


    jumlah_lampiran = (
        legalitas.lampiran.count()
    )


    # =====================================================
    # HAPUS FILE FISIK LAMPIRAN
    # =====================================================

    with transaction.atomic():

        for lampiran in (
            legalitas.lampiran.all()
        ):

            if lampiran.file:

                lampiran.file.delete(
                    save=False
                )


        legalitas.delete()


    # =====================================================
    # PESAN
    # =====================================================

    if jumlah_lampiran > 0:

        messages.success(
            request,
            (
                f'{nama_legalitas} '
                f'beserta {jumlah_lampiran} '
                'lampiran berhasil dihapus.'
            )
        )

    else:

        messages.success(
            request,
            (
                f'{nama_legalitas} '
                'berhasil dihapus.'
            )
        )


    return redirect(
        'detail_bidang_tanah',
        id=bidang.id
    )


@login_required
def tambah_lampiran_legalitas(
    request,
    legalitas_id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke dokumen legalitas tanah.'
        )

        return redirect('/')


    # =====================================================
    # LEGALITAS TANAH
    # =====================================================

    legalitas = get_object_or_404(
        LegalitasTanah.objects
        .select_related(
            'bidang_tanah',
            'bidang_tanah__sekolah'
        ),
        id=legalitas_id,
        bidang_tanah__sekolah=sekolah
    )


    bidang = (
        legalitas.bidang_tanah
    )


    # =====================================================
    # CEK LAMPIRAN
    #
    # Satu legalitas hanya boleh memiliki
    # satu lampiran.
    # =====================================================

    if legalitas.lampiran.exists():

        messages.warning(
            request,
            'Dokumen legalitas ini sudah '
            'memiliki lampiran.'
        )

        return redirect(
            'detail_bidang_tanah',
            id=bidang.id
        )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = DokumenTanahForm(
            request.POST,
            request.FILES
        )


        if form.is_valid():

            lampiran = form.save(
                commit=False
            )

            lampiran.legalitas = (
                legalitas
            )


            # =================================================
            # NAMA LAMPIRAN
            #
            # Tidak lagi mengikuti input manual operator.
            # Nama yang tampil di tabel nanti ditentukan
            # dari urutan legalitas.
            # =================================================

            if hasattr(
                lampiran,
                'nama_file'
            ):

                lampiran.nama_file = ''


            lampiran.full_clean()

            lampiran.save()


            messages.success(
                request,
                'Lampiran dokumen berhasil ditambahkan.'
            )


            return redirect(
                'detail_bidang_tanah',
                id=bidang.id
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = (
            DokumenTanahForm()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'bidang':
            bidang,

        'legalitas':
            legalitas,

    }


    return render(
        request,
        'akun/tambah_lampiran_legalitas.html',
        context
    )


@login_required
def hapus_lampiran_legalitas(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke dokumen legalitas tanah.'
        )

        return redirect('/')


    # =====================================================
    # LAMPIRAN
    # =====================================================

    lampiran = get_object_or_404(
        DokumenTanah.objects
        .select_related(
            'legalitas',
            'legalitas__bidang_tanah',
            'legalitas__bidang_tanah__sekolah'
        ),
        id=id,
        legalitas__bidang_tanah__sekolah=sekolah
    )


    legalitas = (
        lampiran.legalitas
    )

    bidang = (
        legalitas.bidang_tanah
    )


    # =====================================================
    # HANYA POST
    # =====================================================

    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan hapus tidak valid.'
        )

        return redirect(
            'detail_bidang_tanah',
            id=bidang.id
        )


    # =====================================================
    # SIMPAN NAMA FILE
    # =====================================================

    nama_lampiran = (
        lampiran.nama_file
        or lampiran.file.name.split('/')[-1]
    )


    # =====================================================
    # HAPUS FILE + RECORD
    # =====================================================

    with transaction.atomic():

        if lampiran.file:

            lampiran.file.delete(
                save=False
            )

        lampiran.delete()


    messages.success(
        request,
        (
            f'Lampiran "{nama_lampiran}" '
            'berhasil dihapus.'
        )
    )


    return redirect(
        'detail_bidang_tanah',
        id=bidang.id
    )



@login_required
def data_rombel_sekolah(
    request
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data rombel sekolah.'
        )

        return redirect('/')

    tahun_ajaran = TahunAjaran.objects.filter(
        aktif=True
    ).first()

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')

    daftar_rombel = list(
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )
        .order_by('nama_kelas')
    )

    total_rombel = sum(
        data.jumlah_rombel
        for data in daftar_rombel
    )

    total_siswa = sum(
        data.jumlah_siswa_aktif
        for data in daftar_rombel
    )

    total_kapasitas = sum(
        data.kapasitas_total
        for data in daftar_rombel
    )

    total_kekurangan = sum(
        data.kekurangan_kapasitas
        for data in daftar_rombel
    )

    context = {
        'sekolah': sekolah,

        'tahun_ajaran': tahun_ajaran,

        'daftar_rombel':
            daftar_rombel,

        'total_rombel':
            total_rombel,

        'total_siswa':
            total_siswa,

        'total_kapasitas':
            total_kapasitas,

        'total_kekurangan':
            total_kekurangan,
    }

    return render(
        request,
        'akun/data_rombel_sekolah.html',
        context
    )


@login_required
def tambah_rombel_sekolah(request):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_rombel_sekolah'
        )

    # =====================================================
    # INSTANCE ROMBEL BARU
    # =====================================================

    rombel_baru = DataRombel(
        sekolah=sekolah,
        tahun_ajaran=tahun_ajaran
    )

    # =====================================================
    # FORM
    # =====================================================

    if request.method == 'POST':

        form = DataRombelForm(
            request.POST,
            instance=rombel_baru,
            sekolah=sekolah
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                f'Data rombel berhasil ditambahkan '
                f'untuk Tahun Ajaran '
                f'{tahun_ajaran.nama}.'
            )

            return redirect(
                'data_rombel_sekolah'
            )

    else:

        form = DataRombelForm(
            instance=rombel_baru,
            sekolah=sekolah
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {
        'form': form,
        'sekolah': sekolah,
        'tahun_ajaran': tahun_ajaran,
        'judul': 'Tambah Data Rombel',
        'tombol': 'Simpan Rombel',
    }

    return render(
        request,
        'akun/form_rombel_sekolah.html',
        context
    )


@login_required
def edit_rombel_sekolah(
    request,
    id
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        return redirect('/')

    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_rombel_sekolah'
        )

    # =====================================================
    # AMBIL ROMBEL
    # =====================================================

    rombel = get_object_or_404(
        DataRombel,
        id=id,
        sekolah=sekolah,
        tahun_ajaran=tahun_ajaran
    )

    # =====================================================
    # FORM
    # =====================================================

    if request.method == 'POST':

        form = DataRombelForm(
            request.POST,
            instance=rombel,
            sekolah=sekolah
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                'Data rombel berhasil diubah.'
            )

            return redirect(
                'data_rombel_sekolah'
            )

    else:

        form = DataRombelForm(
            instance=rombel,
            sekolah=sekolah
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {
        'form': form,
        'sekolah': sekolah,
        'rombel': rombel,
        'tahun_ajaran': tahun_ajaran,
        'judul': 'Edit Data Rombel',
        'tombol': 'Simpan Perubahan',
    }

    return render(
        request,
        'akun/form_rombel_sekolah.html',
        context
    )


@login_required
def hapus_rombel_sekolah(
    request,
    id
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')

    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_rombel_sekolah'
        )

    rombel = get_object_or_404(
        DataRombel,
        id=id,
        sekolah=sekolah,
        tahun_ajaran=tahun_ajaran
    )

    if request.method == 'POST':

        nama_kelas = (
            rombel.nama_kelas
        )

        rombel.delete()

        messages.success(
            request,
            (
                f'Data {nama_kelas} '
                'berhasil dihapus.'
            )
        )

    return redirect(
        'data_rombel_sekolah'
    )


@login_required
def isi_aset_sekolah(
    request,
    jenis_id
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke data aset sekolah.'
        )

        return redirect('/')

    jenis_aset = get_object_or_404(
        JenisAset,
        id=jenis_id,
        aktif=True
    )

    aset_sekolah = (
        AsetSekolah.objects
        .filter(
            sekolah=sekolah,
            jenis_aset=jenis_aset
        )
        .first()
    )

    if aset_sekolah is None:

        aset_sekolah = AsetSekolah(
            sekolah=sekolah,
            jenis_aset=jenis_aset,
            ketersediaan='TIDAK_ADA',
            jumlah_baik=0,
            jumlah_rusak_ringan=0,
            jumlah_rusak_berat=0,
            sudah_diisi=False
        )

    if request.method == 'POST':

        form = AsetSekolahForm(
            request.POST,
            request.FILES,
            instance=aset_sekolah
        )

        if form.is_valid():

            with transaction.atomic():

                aset_sekolah = form.save(
                    commit=False
                )

                aset_sekolah.sekolah = sekolah
                aset_sekolah.jenis_aset = jenis_aset
                aset_sekolah.sudah_diisi = True

                aset_sekolah.save()

                if (
                    aset_sekolah.ketersediaan
                    == 'TIDAK_ADA'
                ):

                    for foto_lama in (
                        aset_sekolah
                        .foto_dokumentasi
                        .all()
                    ):

                        foto_lama.foto.delete(
                            save=False
                        )

                    (
                        aset_sekolah
                        .foto_dokumentasi
                        .all()
                        .delete()
                    )

                else:

                    daftar_foto = [
                        (
                            1,
                            form.cleaned_data.get(
                                'foto_1'
                            )
                        ),
                        (
                            2,
                            form.cleaned_data.get(
                                'foto_2'
                            )
                        ),
                        (
                            3,
                            form.cleaned_data.get(
                                'foto_3'
                            )
                        ),
                    ]

                    for (
                        urutan,
                        file_foto
                    ) in daftar_foto:

                        if not file_foto:
                            continue

                        foto_aset, dibuat_foto = (
                            FotoAset.objects
                            .get_or_create(
                                aset_sekolah=(
                                    aset_sekolah
                                ),
                                urutan=urutan
                            )
                        )

                        if (
                            not dibuat_foto
                            and foto_aset.foto
                        ):

                            foto_aset.foto.delete(
                                save=False
                            )

                        foto_aset.foto = (
                            file_foto
                        )

                        foto_aset.full_clean()
                        foto_aset.save()

            messages.success(
                request,
                (
                    f'Data {jenis_aset.nama} '
                    'berhasil disimpan.'
                )
            )

            tab_per_kategori = {
                'BANGUNAN': 'bangunan',
                'SANITASI': 'sanitasi',
                'PERLENGKAPAN': 'perlengkapan',
            }

            tab_tujuan = tab_per_kategori.get(
                jenis_aset.kategori,
                'bangunan'
            )

            url_tujuan = reverse(
                'data_aset'
            )

            return redirect(
                f'{url_tujuan}'
                f'?tab={tab_tujuan}'
                f'#aset-{jenis_aset.id}'
            )

    else:

        form = AsetSekolahForm(
            instance=aset_sekolah
        )

    if aset_sekolah.pk:

        foto_tersimpan = {
            foto.urutan: foto
            for foto in (
                aset_sekolah
                .foto_dokumentasi
                .all()
            )
        }

    else:

        foto_tersimpan = {}

    context = {
        'form': form,
        'sekolah': sekolah,
        'jenis_aset': jenis_aset,
        'aset_sekolah': aset_sekolah,
        'foto_1': (
            foto_tersimpan.get(1)
        ),
        'foto_2': (
            foto_tersimpan.get(2)
        ),
        'foto_3': (
            foto_tersimpan.get(3)
        ),
    }

    return render(
        request,
        'akun/isi_aset_sekolah.html',
        context
    )


@login_required
def tambah_aset(request):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    if request.method == 'POST':

        form = AsetForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            aset = form.save(commit=False)
            aset.sekolah = sekolah
            aset.save()
            messages.success(request, 'Data aset berhasil ditambahkan')

            return redirect('/data-aset/')

    else:
        form = AsetForm()

    context = {
        'form': form,
    }

    return render(request, 'akun/tambah_aset.html', context)


@login_required
def edit_murid(request, id):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # MURID
    # =====================================================

    murid = get_object_or_404(
        Murid,
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # RIWAYAT MURID TAHUN AJARAN AKTIF
    # =====================================================

    riwayat = (
        RiwayatMurid.objects
        .filter(
            murid=murid,
            tahun_ajaran=tahun_ajaran
        )
        .select_related(
            'rombel'
        )
        .first()
    )

    if riwayat is None:

        messages.error(
            request,
            (
                'Riwayat murid pada Tahun Ajaran '
                f'{tahun_ajaran.nama} tidak ditemukan.'
            )
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = MuridForm(
            request.POST,
            instance=murid,
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran
        )

        form.fields.pop(
            'status_masuk',
            None
        )

        if form.is_valid():

            with transaction.atomic():

                murid = form.save(
                    commit=False
                )

                rombel_baru = (
                    form.cleaned_data[
                        'rombel'
                    ]
                )

                # Sinkronisasi field kelas lama
                murid.kelas = (
                    rombel_baru.nama_kelas
                )

                murid.save()


                # Update rombel pada riwayat
                riwayat.rombel = (
                    rombel_baru
                )

                riwayat.full_clean()
                riwayat.save()


            messages.success(
                request,
                'Data murid berhasil diubah.'
            )

            return redirect(
                'data_murid'
            )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = MuridForm(
            instance=murid,
            sekolah=sekolah,
            tahun_ajaran=tahun_ajaran,
            initial={
                'rombel':
                    riwayat.rombel
            }
        )

        form.fields.pop(
            'status_masuk',
            None
        )


    context = {

        'form':
            form,

        'murid':
            murid,

        'riwayat':
            riwayat,

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,
    }


    return render(
        request,
        'akun/edit_murid.html',
        context
    )


@login_required
def hapus_murid(request, id):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # AMBIL MURID
    # =====================================================

    murid = get_object_or_404(
        Murid,
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # CEK RIWAYAT AKADEMIK
    # =====================================================

    punya_riwayat = (
        RiwayatMurid.objects
        .filter(
            murid=murid
        )
        .exists()
    )


    if punya_riwayat:

        messages.error(
            request,
            (
                f'{murid.nama} tidak dapat dihapus '
                'karena sudah memiliki riwayat akademik. '
                'Gunakan proses Pindah, Keluar, atau Lulus '
                'untuk mengakhiri status siswa.'
            )
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # HAPUS MURID
    # =====================================================

    nama_murid = murid.nama

    murid.delete()


    messages.success(
        request,
        (
            f'Data {nama_murid} berhasil '
            'dihapus permanen.'
        )
    )


    return redirect(
        'data_murid'
    )


@login_required
def ubah_status_murid(request, id):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # HANYA POST
    # =====================================================

    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan perubahan status tidak valid.'
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # MURID
    # =====================================================

    murid = get_object_or_404(
        Murid,
        id=id,
        sekolah=sekolah
    )


    # =====================================================
    # RIWAYAT AKTIF
    # =====================================================

    riwayat = (
        RiwayatMurid.objects
        .filter(
            murid=murid,
            tahun_ajaran=tahun_ajaran,
            status_akhir='AKTIF'
        )
        .select_related(
            'rombel'
        )
        .first()
    )


    if riwayat is None:

        messages.error(
            request,
            (
                f'{murid.nama} tidak memiliki '
                f'status aktif pada Tahun Ajaran '
                f'{tahun_ajaran.nama}.'
            )
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # AKSI
    # =====================================================

    aksi = request.POST.get(
        'aksi'
    )


    aksi_valid = (
        'PINDAH',
        'KELUAR',
    )


    if aksi not in aksi_valid:

        messages.error(
            request,
            'Status siswa tidak valid.'
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # KETERANGAN
    # =====================================================

    keterangan = (
        request.POST.get(
            'keterangan',
            ''
        )
        .strip()
    )


    # =====================================================
    # SIMPAN
    # =====================================================

    try:

        with transaction.atomic():

            riwayat.status_akhir = aksi

            if keterangan:

                riwayat.keterangan = (
                    keterangan
                )

            riwayat.full_clean()

            riwayat.save(
                update_fields=[
                    'status_akhir',
                    'keterangan',
                    'updated_at',
                ]
            )


    except Exception as error:

        messages.error(
            request,
            (
                'Perubahan status siswa gagal: '
                f'{error}'
            )
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # PESAN
    # =====================================================

    if aksi == 'PINDAH':

        pesan_status = (
            'Pindah Sekolah'
        )

    else:

        pesan_status = (
            'Keluar'
        )


    messages.success(
        request,
        (
            f'{murid.nama} berhasil ditetapkan '
            f'dengan status {pesan_status} '
            f'pada Tahun Ajaran '
            f'{tahun_ajaran.nama}.'
        )
    )


    return redirect(
        'data_murid'
    )


@login_required
def pulihkan_status_murid(request, id):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan pemulihan status tidak valid.'
        )

        return redirect(
            'data_murid'
        )


    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_murid'
        )


    murid = get_object_or_404(
        Murid,
        id=id,
        sekolah=sekolah
    )


    riwayat = (
        RiwayatMurid.objects
        .filter(
            murid=murid,
            tahun_ajaran=tahun_ajaran
        )
        .select_related(
            'rombel'
        )
        .first()
    )


    if riwayat is None:

        messages.error(
            request,
            (
                f'Riwayat {murid.nama} pada '
                f'Tahun Ajaran {tahun_ajaran.nama} '
                'tidak ditemukan.'
            )
        )

        return redirect(
            'data_murid'
        )


    if riwayat.status_akhir not in (
        'PINDAH',
        'KELUAR',
    ):

        messages.error(
            request,
            (
                f'Status {murid.nama} tidak dapat '
                'dipulihkan melalui menu ini.'
            )
        )

        return redirect(
            'data_murid'
        )


    try:

        with transaction.atomic():

            # =====================================================
            # SIMPAN STATUS SEBELUM DIPULIHKAN
            # =====================================================

            status_sebelumnya = (
                riwayat.status_akhir
            )


            # =====================================================
            # CATAT WAKTU PEMULIHAN
            # =====================================================

            waktu_pemulihan = (
                timezone.localtime()
                .strftime(
                    '%d-%m-%Y %H:%M'
                )
            )


            # =====================================================
            # TAMBAHKAN JEJAK KE KETERANGAN
            # =====================================================

            catatan_pemulihan = (
                f'[Status {status_sebelumnya} '
                f'dipulihkan menjadi AKTIF '
                f'pada {waktu_pemulihan}]'
            )


            if riwayat.keterangan:

                riwayat.keterangan = (
                    riwayat.keterangan.strip()
                    + '\n'
                    + catatan_pemulihan
                )

            else:

                riwayat.keterangan = (
                    catatan_pemulihan
                )


            # =====================================================
            # PULIHKAN STATUS
            # =====================================================

            riwayat.status_akhir = 'AKTIF'


            riwayat.full_clean()


            riwayat.save(
                update_fields=[
                    'status_akhir',
                    'keterangan',
                    'updated_at',
                ]
            )


            # =====================================================
            # SINKRONISASI FIELD KELAS
            # =====================================================

            murid.kelas = (
                riwayat.rombel.nama_kelas
            )

            murid.save(
                update_fields=[
                    'kelas'
                ]
            )


            murid.kelas = (
                riwayat.rombel.nama_kelas
            )

            murid.save(
                update_fields=[
                    'kelas'
                ]
            )


    except Exception as error:

        messages.error(
            request,
            (
                'Pemulihan status murid gagal: '
                f'{error}'
            )
        )

        return redirect(
            'data_murid'
        )


    messages.success(
        request,
        (
            f'{murid.nama} berhasil dikembalikan '
            f'menjadi siswa aktif pada '
            f'Tahun Ajaran {tahun_ajaran.nama}.'
        )
    )


    return redirect(
        'data_murid'
    )


@login_required
def edit_aset(request, id):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    aset = Aset.objects.get(
        id=id,
        sekolah=sekolah
    )

    if request.method == 'POST':

        form = AsetForm(
            request.POST,
            request.FILES,
            instance=aset
        )

        if form.is_valid():

            form.save()
            messages.success(request, 'Data aset berhasil diubah')

            return redirect('/data-aset/')

    else:
        form = AsetForm(instance=aset)

    context = {
        'form': form,
        'aset': aset,
    }

    return render(
        request,
        'akun/edit_aset.html',
        context
    )


@login_required
def hapus_aset(request, id):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    aset = Aset.objects.get(
        id=id,
        sekolah=sekolah
    )

    messages.success(request, 'Data aset berhasil dihapus')
    aset.delete()

    return redirect('/data-aset/')


@login_required
def data_prestasi(request):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    prestasi = Prestasi.objects.filter(sekolah=sekolah)

    context = {
        'prestasi': prestasi,
        'sekolah': sekolah,
    }

    return render(request, 'akun/data_prestasi.html', context)


@login_required
def tambah_prestasi(request):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    if request.method == 'POST':
        form = PrestasiForm(request.POST, request.FILES)

        form.fields['murid'].queryset = Murid.objects.filter(sekolah=sekolah)

        if form.is_valid():
            prestasi = form.save(commit=False)
            prestasi.sekolah = sekolah
            prestasi.save()
            messages.success(request, 'Data prestasi berhasil ditambahkan')

            return redirect('/data-prestasi/')

    else:
        form = PrestasiForm()
        form.fields['murid'].queryset = Murid.objects.filter(sekolah=sekolah)

    context = {
        'form': form,
    }

    return render(request, 'akun/tambah_prestasi.html', context)


@login_required
def edit_prestasi(request, id):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    prestasi = Prestasi.objects.get(
        id=id,
        sekolah=sekolah
    )

    if request.method == 'POST':
        form = PrestasiForm(
            request.POST,
            request.FILES,
            instance=prestasi
        )

        form.fields['murid'].queryset = Murid.objects.filter(sekolah=sekolah)

        if form.is_valid():
            form.save()
            messages.success(request, 'Data prestasi berhasil diubah')
            return redirect('/data-prestasi/')

    else:
        form = PrestasiForm(instance=prestasi)
        form.fields['murid'].queryset = Murid.objects.filter(sekolah=sekolah)

    context = {
        'form': form,
        'prestasi': prestasi,
    }

    return render(request, 'akun/edit_prestasi.html', context)


@login_required
def hapus_prestasi(request, id):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    prestasi = Prestasi.objects.get(
        id=id,
        sekolah=sekolah
    )

    messages.success(request, 'Data prestasi berhasil dihapus')
    prestasi.delete()

    return redirect('/data-prestasi/')


def ambil_sekolah_operator(request):
    if not request.user.is_authenticated:
        return None

    if request.user.is_superuser:
        return None

    try:
        return request.user.sekolah_operator
    except:
        return None



def hitung_kesiapan_tahun_ajaran(
    sekolah,
    tahun_asal,
    tahun_tujuan
):

    # =========================================
    # RIWAYAT SISWA TAHUN ASAL
    # =========================================

    semua_riwayat = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            tahun_ajaran=tahun_asal
        )
    )

    total_siswa = (
        semua_riwayat.count()
    )


    # =========================================
    # ROMBEL TAHUN ASAL
    # =========================================

    jumlah_rombel_asal = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_asal
        )
        .count()
    )

    rombel_tahun_asal = (
        jumlah_rombel_asal > 0
    )


    # =========================================
    # SISWA BELUM DIPROSES
    # =========================================

    belum_diproses = (
        semua_riwayat
        .filter(
            status_akhir='AKTIF'
        )
        .count()
    )


    # =========================================
    # SISWA SUDAH DIPROSES
    # =========================================

    sudah_diproses = (
        total_siswa
        - belum_diproses
    )


    # =========================================
    # PROGRES
    # =========================================

    if total_siswa > 0:

        persentase = round(
            (
                sudah_diproses
                / total_siswa
            )
            * 100
        )

    elif rombel_tahun_asal:

        # Sekolah sudah memiliki struktur
        # akademik tetapi memang tidak
        # mempunyai siswa.
        persentase = 100

    else:

        persentase = 0


    # =========================================
    # ROMBEL TAHUN TUJUAN
    # =========================================

    jumlah_rombel_tujuan = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_tujuan
        )
        .count()
    )

    rombel_tahun_tujuan = (
        jumlah_rombel_tujuan > 0
    )


    # =========================================
    # STATUS KESIAPAN
    # =========================================

    # Sekolah normal yang memiliki siswa
    siap_dengan_siswa = (
        total_siswa > 0
        and
        belum_diproses == 0
        and
        rombel_tahun_tujuan
    )


    # Sekolah yang sudah mempunyai struktur
    # akademik tetapi memang tidak mempunyai
    # siswa pada tahun asal
    siap_tanpa_siswa = (
        total_siswa == 0
        and
        rombel_tahun_asal
        and
        rombel_tahun_tujuan
    )


    siap = (
        siap_dengan_siswa
        or
        siap_tanpa_siswa
    )


    # =========================================
    # PENENTUAN STATUS
    # =========================================

    if siap_dengan_siswa:

        status = 'SIAP'
        warna = 'success'


    elif siap_tanpa_siswa:

        status = 'SIAP - TANPA SISWA'
        warna = 'success'


    elif (
        total_siswa == 0
        and
        not rombel_tahun_asal
    ):

        status = 'BELUM ADA DATA'
        warna = 'secondary'


    elif (
        total_siswa == 0
        and
        rombel_tahun_asal
        and
        not rombel_tahun_tujuan
    ):

        status = 'ROMBEL BELUM SIAP'
        warna = 'warning'


    elif belum_diproses > 0:

        status = 'BELUM SIAP'
        warna = 'danger'


    else:

        status = 'ROMBEL BELUM SIAP'
        warna = 'warning'


    # =========================================
    # HASIL
    # =========================================

    return {

        'total_siswa':
            total_siswa,

        'sudah_diproses':
            sudah_diproses,

        'belum_diproses':
            belum_diproses,

        'persentase':
            persentase,

        'jumlah_rombel_asal':
            jumlah_rombel_asal,

        'rombel_tahun_asal':
            rombel_tahun_asal,

        'jumlah_rombel_tujuan':
            jumlah_rombel_tujuan,

        'rombel_tahun_tujuan':
            rombel_tahun_tujuan,

        'siap':
            siap,

        'siap_tanpa_siswa':
            siap_tanpa_siswa,

        'status':
            status,

        'warna':
            warna,
    }


@login_required
def monitoring_kesiapan_tahun_ajaran(
    request
):

    # =====================================================
    # HAK AKSES
    # =====================================================

    profil = ambil_profil_user(
        request
    )

    admin_kabupaten = (
        user_admin_kabupaten(
            request
        )
    )

    admin_kecamatan = (
        user_admin_kecamatan(
            request
        )
    )


    if not (
        admin_kabupaten
        or
        admin_kecamatan
    ):

        messages.error(
            request,
            (
                'Anda tidak memiliki akses '
                'ke monitoring kesiapan '
                'Tahun Ajaran.'
            )
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF = TAHUN ASAL
    # =====================================================

    tahun_asal = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_asal is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('kelola_tahun_ajaran')


    # =====================================================
    # TAHUN TUJUAN
    #
    # Harus merupakan penerus langsung
    # dari Tahun Ajaran aktif.
    # =====================================================

    tahun_tujuan = (
        TahunAjaran.objects
        .filter(
            aktif=False,
            tahun_sebelumnya=tahun_asal
        )
        .first()
    )


    if tahun_tujuan is None:

        messages.info(
            request,
            (
                f'Belum tersedia Tahun Ajaran berikutnya '
                f'setelah {tahun_asal.nama}. '
                f'Silakan buat Tahun Ajaran berikutnya '
                f'terlebih dahulu.'
            )
        )

        return redirect(
            'kelola_tahun_ajaran'
        )


    # =====================================================
    # DAFTAR SEKOLAH SESUAI HAK AKSES
    # =====================================================

    daftar_sekolah = (
        Sekolah.objects
        .select_related(
            'kategori'
        )
        .order_by(
            'kecamatan',
            'nama'
        )
    )


    # =====================================================
    # ADMIN KECAMATAN
    # =====================================================

    if admin_kecamatan:

        daftar_sekolah = (
            daftar_sekolah
            .filter(
                kecamatan=profil.kecamatan
            )
        )


    # =====================================================
    # FILTER
    # =====================================================

    kecamatan = (
        request.GET
        .get(
            'kecamatan',
            ''
        )
        .strip()
    )

    kategori = (
        request.GET
        .get(
            'kategori',
            ''
        )
        .strip()
    )

    status = (
        request.GET
        .get(
            'status',
            ''
        )
        .strip()
        .upper()
    )

    cari = (
        request.GET
        .get(
            'cari',
            ''
        )
        .strip()
    )


    # =====================================================
    # FILTER KECAMATAN
    # HANYA ADMIN KABUPATEN
    # =====================================================

    if (
        kecamatan
        and
        admin_kabupaten
    ):

        daftar_sekolah = (
            daftar_sekolah
            .filter(
                kecamatan=kecamatan
            )
        )


    # =====================================================
    # FILTER KATEGORI
    # =====================================================

    if kategori:

        daftar_sekolah = (
            daftar_sekolah
            .filter(
                kategori_id=kategori
            )
        )


    # =====================================================
    # PENCARIAN SEKOLAH
    # =====================================================

    if cari:

        daftar_sekolah = (
            daftar_sekolah
            .filter(
                nama__icontains=cari
            )
        )


    # =====================================================
    # HITUNG KESIAPAN SEMUA SEKOLAH
    # =====================================================

    hasil_semua = []


    for sekolah in daftar_sekolah:

        kesiapan = (
            hitung_kesiapan_tahun_ajaran(
                sekolah,
                tahun_asal,
                tahun_tujuan
            )
        )


        # =============================================
        # ALASAN / KETERANGAN
        # =============================================

        if kesiapan['status'] == 'SIAP':

            alasan = (
                'Seluruh siswa telah diproses '
                'dan rombel Tahun Ajaran tujuan '
                'telah tersedia.'
            )


        elif (
            kesiapan['status']
            == 'SIAP - TANPA SISWA'
        ):

            alasan = (
                'Sekolah tidak memiliki siswa '
                'pada Tahun Ajaran asal, '
                'namun rombel Tahun Ajaran '
                'tujuan telah tersedia.'
            )


        elif (
            kesiapan['status']
            == 'BELUM SIAP'
        ):

            alasan = (
                f"{kesiapan['belum_diproses']} "
                'siswa masih belum diproses.'
            )


        elif (
            kesiapan['status']
            == 'ROMBEL BELUM SIAP'
        ):

            alasan = (
                'Proses siswa telah selesai, '
                'tetapi rombel Tahun Ajaran '
                'tujuan belum tersedia.'
            )


        elif (
            kesiapan['status']
            == 'BELUM ADA DATA'
        ):

            alasan = (
                'Belum terdapat data siswa '
                'maupun struktur rombel pada '
                'Tahun Ajaran asal.'
            )


        else:

            alasan = (
                'Status kesiapan belum '
                'dapat ditentukan.'
            )


        hasil_semua.append({
            'sekolah':
                sekolah,

            'kesiapan':
                kesiapan,

            'alasan':
                alasan,
        })


    # =====================================================
    # RINGKASAN
    #
    # Dihitung SEBELUM filter status,
    # supaya kartu ringkasan tetap menggambarkan
    # keseluruhan sekolah dalam cakupan/filter wilayah.
    # =====================================================

    total_sekolah = (
        len(
            hasil_semua
        )
    )


    total_siap_normal = sum(
        1
        for data in hasil_semua
        if (
            data['kesiapan']['status']
            == 'SIAP'
        )
    )


    total_siap_tanpa_siswa = sum(
        1
        for data in hasil_semua
        if (
            data['kesiapan']['status']
            == 'SIAP - TANPA SISWA'
        )
    )


    total_siap = (
        total_siap_normal
        +
        total_siap_tanpa_siswa
    )


    total_belum_siap = sum(
        1
        for data in hasil_semua
        if (
            data['kesiapan']['status']
            == 'BELUM SIAP'
        )
    )


    total_rombel_belum_siap = sum(
        1
        for data in hasil_semua
        if (
            data['kesiapan']['status']
            == 'ROMBEL BELUM SIAP'
        )
    )


    total_belum_ada_data = sum(
        1
        for data in hasil_semua
        if (
            data['kesiapan']['status']
            == 'BELUM ADA DATA'
        )
    )


    # =====================================================
    # PERSENTASE KESIAPAN
    # =====================================================

    if total_sekolah > 0:

        persentase_siap = round(
            (
                total_siap
                / total_sekolah
            )
            * 100,
            1
        )

    else:

        persentase_siap = 0


    # =====================================================
    # STATUS KESIAPAN KABUPATEN / KECAMATAN
    # =====================================================

    semua_siap = (
        total_sekolah > 0
        and
        total_siap == total_sekolah
    )


    # =====================================================
    # FILTER STATUS
    #
    # Filter dilakukan SETELAH ringkasan dihitung.
    # =====================================================

    if status:

        if status == 'SIAP':

            hasil_monitoring = [
                data
                for data in hasil_semua
                if data['kesiapan']['siap']
            ]

        else:

            hasil_monitoring = [
                data
                for data in hasil_semua
                if (
                    data['kesiapan']['status']
                    == status
                )
            ]

    else:

        hasil_monitoring = (
            hasil_semua
        )


    # =====================================================
    # RINGKASAN TAMBAHAN
    # =====================================================

    total_perlu_proses_siswa = (
        total_belum_siap
    )

    total_perlu_rombel = (
        total_rombel_belum_siap
    )

    total_tanpa_data = (
        total_belum_ada_data
    )


    # =====================================================
    # PILIHAN FILTER
    # =====================================================

    daftar_kecamatan = (
        Sekolah.objects
        .exclude(
            kecamatan__isnull=True
        )
        .exclude(
            kecamatan=''
        )
        .values_list(
            'kecamatan',
            flat=True
        )
        .distinct()
        .order_by(
            'kecamatan'
        )
    )


    daftar_kategori = (
        KategoriSekolah.objects
        .all()
        .order_by(
            'nama'
        )
    )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'tahun_asal':
            tahun_asal,

        'tahun_tujuan':
            tahun_tujuan,

        'hasil_monitoring':
            hasil_monitoring,


        # =========================
        # RINGKASAN
        # =========================

        'total_sekolah':
            total_sekolah,

        'total_siap':
            total_siap,

        'total_siap_normal':
            total_siap_normal,

        'total_siap_tanpa_siswa':
            total_siap_tanpa_siswa,

        'total_belum_siap':
            total_belum_siap,

        'total_rombel_belum_siap':
            total_rombel_belum_siap,

        'total_belum_ada_data':
            total_belum_ada_data,

        'persentase_siap':
            persentase_siap,

        'semua_siap':
            semua_siap,


        # =========================
        # TINDAK LANJUT
        # =========================

        'total_perlu_proses_siswa':
            total_perlu_proses_siswa,

        'total_perlu_rombel':
            total_perlu_rombel,

        'total_tanpa_data':
            total_tanpa_data,


        # =========================
        # FILTER
        # =========================

        'daftar_kecamatan':
            daftar_kecamatan,

        'daftar_kategori':
            daftar_kategori,

        'filter_kecamatan':
            kecamatan,

        'filter_kategori':
            kategori,

        'filter_status':
            status,

        'filter_cari':
            cari,


        # =========================
        # AKSES
        # =========================

        'admin_kabupaten':
            admin_kabupaten,

        'admin_kecamatan':
            admin_kecamatan,
    }


    return render(
        request,
        'akun/monitoring_kesiapan_tahun_ajaran.html',
        context
    )


@login_required
def kelola_tahun_ajaran(request):

    # =====================================================
    # HAK AKSES
    # =====================================================

    if not user_admin_kabupaten(request):

        messages.error(
            request,
            (
                'Pengelolaan Tahun Ajaran hanya '
                'dapat dilakukan oleh Admin Kabupaten.'
            )
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    # =====================================================
    # RINGKASAN KESIAPAN TAHUN BERIKUTNYA
    # =====================================================

    ringkasan_kesiapan = None
    tahun_tujuan_kesiapan = None


    if tahun_aktif:

        tahun_tujuan_kesiapan = (
            TahunAjaran.objects
            .filter(
                aktif=False,
                tahun_sebelumnya=tahun_aktif
            )
            .first()
        )


        if tahun_tujuan_kesiapan:

            daftar_sekolah = (
                Sekolah.objects
                .all()
                .order_by('nama')
            )


            # =============================================
            # COUNTER
            # =============================================

            total_sekolah_kesiapan = 0

            total_siap_kesiapan = 0

            total_siap_normal_kesiapan = 0

            total_siap_tanpa_siswa_kesiapan = 0

            total_belum_siap_kesiapan = 0

            total_rombel_belum_siap_kesiapan = 0

            total_belum_ada_data_kesiapan = 0


            # =============================================
            # HITUNG SETIAP SEKOLAH
            # =============================================

            for sekolah in daftar_sekolah:

                hasil = (
                    hitung_kesiapan_tahun_ajaran(
                        sekolah,
                        tahun_aktif,
                        tahun_tujuan_kesiapan
                    )
                )


                total_sekolah_kesiapan += 1


                # =========================================
                # SEMUA STATUS DENGAN siap=True
                # DIANGGAP SIAP
                # =========================================

                if hasil['siap']:

                    total_siap_kesiapan += 1


                    if hasil['status'] == 'SIAP':

                        total_siap_normal_kesiapan += 1


                    elif (
                        hasil['status']
                        == 'SIAP - TANPA SISWA'
                    ):

                        total_siap_tanpa_siswa_kesiapan += 1


                # =========================================
                # BELUM SIAP
                # =========================================

                elif hasil['status'] == 'BELUM SIAP':

                    total_belum_siap_kesiapan += 1


                # =========================================
                # ROMBEL BELUM SIAP
                # =========================================

                elif (
                    hasil['status']
                    == 'ROMBEL BELUM SIAP'
                ):

                    total_rombel_belum_siap_kesiapan += 1


                # =========================================
                # BELUM ADA DATA
                # =========================================

                elif (
                    hasil['status']
                    == 'BELUM ADA DATA'
                ):

                    total_belum_ada_data_kesiapan += 1


            # =============================================
            # PERSENTASE KESIAPAN
            # =============================================

            if total_sekolah_kesiapan > 0:

                persentase_kesiapan = round(
                    (
                        total_siap_kesiapan
                        /
                        total_sekolah_kesiapan
                    )
                    * 100,
                    1
                )

            else:

                persentase_kesiapan = 0


            # =============================================
            # BOLEH DIAKTIFKAN?
            # =============================================

            siap_diaktifkan = (
                total_sekolah_kesiapan > 0
                and
                total_siap_kesiapan
                == total_sekolah_kesiapan
            )


            # =============================================
            # RINGKASAN
            # =============================================

            ringkasan_kesiapan = {

                'total_sekolah':
                    total_sekolah_kesiapan,

                'total_siap':
                    total_siap_kesiapan,

                'total_siap_normal':
                    total_siap_normal_kesiapan,

                'total_siap_tanpa_siswa':
                    total_siap_tanpa_siswa_kesiapan,

                'total_belum_siap':
                    total_belum_siap_kesiapan,

                'total_rombel_belum_siap':
                    total_rombel_belum_siap_kesiapan,

                'total_belum_ada_data':
                    total_belum_ada_data_kesiapan,

                'persentase':
                    persentase_kesiapan,

                'siap_diaktifkan':
                    siap_diaktifkan,
            }


    # =====================================================
    # FORM TAMBAH TAHUN AJARAN
    # =====================================================

    if request.method == 'POST':

        form = TahunAjaranForm(
            request.POST
        )


        if form.is_valid():

            # =============================================
            # HARUS ADA TAHUN AKTIF
            # =============================================

            if tahun_aktif is None:

                messages.error(
                    request,
                    (
                        'Tidak ditemukan Tahun Ajaran aktif '
                        'sebagai periode sebelumnya.'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            tahun_baru = (
                form.save(
                    commit=False
                )
            )


            # =============================================
            # VALIDASI FORMAT / URUTAN
            # =============================================

            try:

                bagian_baru = (
                    tahun_baru.nama
                    .split('/')
                )

                awal_baru = int(
                    bagian_baru[0]
                )

                akhir_baru = int(
                    bagian_baru[1]
                )


                bagian_aktif = (
                    tahun_aktif.nama
                    .split('/')
                )

                awal_aktif = int(
                    bagian_aktif[0]
                )

                akhir_aktif = int(
                    bagian_aktif[1]
                )


            except (
                ValueError,
                IndexError,
                TypeError
            ):

                messages.error(
                    request,
                    (
                        'Format Tahun Ajaran '
                        'tidak valid.'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            # =============================================
            # HARUS SATU PERIODE SETELAH TAHUN AKTIF
            # =============================================

            if (
                awal_baru != awal_aktif + 1
                or
                akhir_baru != akhir_aktif + 1
            ):

                messages.error(
                    request,
                    (
                        f'Tahun Ajaran baru harus '
                        f'satu periode setelah '
                        f'{tahun_aktif.nama}. '
                        f'Tahun berikutnya seharusnya '
                        f'{awal_aktif + 1}/'
                        f'{akhir_aktif + 1}.'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            # =============================================
            # CEK DUPLIKAT NAMA
            # =============================================

            if (
                TahunAjaran.objects
                .filter(
                    nama=tahun_baru.nama
                )
                .exists()
            ):

                messages.error(
                    request,
                    (
                        f'Tahun Ajaran '
                        f'{tahun_baru.nama} '
                        f'sudah tersedia.'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            # =============================================
            # PENGAMAN:
            # TAHUN AKTIF SUDAH PUNYA PENERUS
            # =============================================

            if (
                TahunAjaran.objects
                .filter(
                    tahun_sebelumnya=tahun_aktif
                )
                .exists()
            ):

                messages.error(
                    request,
                    (
                        f'Tahun Ajaran '
                        f'{tahun_aktif.nama} '
                        f'sudah memiliki Tahun Ajaran '
                        f'berikutnya.'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            # =============================================
            # RELASI OTOMATIS
            # =============================================

            tahun_baru.tahun_sebelumnya = (
                tahun_aktif
            )


            # Tahun baru selalu dibuat
            # dalam kondisi belum aktif.
            tahun_baru.aktif = False


            # =============================================
            # SIMPAN
            # =============================================

            try:

                tahun_baru.full_clean()

                tahun_baru.save()


            except Exception as error:

                messages.error(
                    request,
                    (
                        'Tahun Ajaran gagal dibuat: '
                        f'{error}'
                    )
                )

                return redirect(
                    'kelola_tahun_ajaran'
                )


            messages.success(
                request,
                (
                    f'Tahun Ajaran '
                    f'{tahun_baru.nama} '
                    f'berhasil ditambahkan. '
                    f'Tahun sebelumnya otomatis '
                    f'{tahun_aktif.nama}.'
                )
            )


            return redirect(
                'kelola_tahun_ajaran'
            )


    else:

        form = TahunAjaranForm()


    # =====================================================
    # DAFTAR TAHUN AJARAN
    # =====================================================

    daftar_tahun = (
        TahunAjaran.objects
        .select_related(
            'tahun_sebelumnya'
        )
        .all()
        .order_by(
            '-nama'
        )
    )


    # =====================================================
    # TAHUN BERIKUTNYA YANG DIHARAPKAN
    # =====================================================

    tahun_berikutnya = None


    if tahun_aktif:

        try:

            bagian = (
                tahun_aktif.nama
                .split('/')
            )

            awal = int(
                bagian[0]
            )

            akhir = int(
                bagian[1]
            )

            tahun_berikutnya = (
                f'{awal + 1}/'
                f'{akhir + 1}'
            )


        except (
            ValueError,
            IndexError,
            TypeError
        ):

            tahun_berikutnya = None


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'daftar_tahun':
            daftar_tahun,

        'tahun_aktif':
            tahun_aktif,

        'tahun_berikutnya':
            tahun_berikutnya,

        'tahun_tujuan_kesiapan':
            tahun_tujuan_kesiapan,

        'ringkasan_kesiapan':
            ringkasan_kesiapan,
    }


    # =====================================================
    # RENDER
    # =====================================================

    return render(
        request,
        'akun/kelola_tahun_ajaran.html',
        context
    )


def parse_tanggal_excel(nilai):

    if nilai is None:
        return None

    if isinstance(nilai, datetime):
        return nilai.date()

    if isinstance(nilai, date):
        return nilai

    nilai = str(nilai).strip()

    format_tanggal = [
        '%Y-%m-%d',
        '%d-%m-%Y',
        '%d/%m/%Y',
        '%Y/%m/%d',
    ]

    for format_item in format_tanggal:

        try:

            return datetime.strptime(
                nilai,
                format_item
            ).date()

        except ValueError:
            continue

    raise ValueError(
        f'Format tanggal tidak dikenali: {nilai}'
    )


@login_required
def import_murid_excel(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke import data murid.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_murid'
        )


    # =====================================================
    # HASIL IMPORT
    # =====================================================

    hasil_import = None
    detail_gagal = []


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        file_excel = (
            request.FILES.get(
                'file_excel'
            )
        )


        # =================================================
        # FILE WAJIB
        # =================================================

        if not file_excel:

            messages.error(
                request,
                'Silakan pilih file Excel.'
            )

            return redirect(
                'import_murid_excel'
            )


        # =================================================
        # VALIDASI EKSTENSI
        # =================================================

        nama_file = (
            file_excel.name.lower()
        )


        if not (
            nama_file.endswith('.xlsx')
            or
            nama_file.endswith('.xlsm')
        ):

            messages.error(
                request,
                'File harus berformat XLSX atau XLSM.'
            )

            return redirect(
                'import_murid_excel'
            )


        # =================================================
        # BACA WORKBOOK
        # =================================================

        try:

            workbook = load_workbook(
                file_excel,
                data_only=True
            )

            sheet = (
                workbook.active
            )


        except Exception as error:

            messages.error(
                request,
                (
                    'File Excel tidak dapat dibaca: '
                    f'{error}'
                )
            )

            return redirect(
                'import_murid_excel'
            )


        # =================================================
        # VALIDASI HEADER
        # =================================================

        header_wajib = [
            'Nama',
            'NISN',
            'NIK',
            'Jenis Kelamin',
            'Tempat Lahir',
            'Tanggal Lahir',
            'Alamat',
            'Kelas',
            'Status Masuk',
        ]


        header_excel = [
            (
                str(cell.value).strip()
                if cell.value is not None
                else ''
            )
            for cell in sheet[1]
        ]


        if header_excel[:9] != header_wajib:

            messages.error(
                request,
                (
                    'Format Excel tidak sesuai '
                    'template SIPETAKU.'
                )
            )

            return redirect(
                'import_murid_excel'
            )


        # =================================================
        # COUNTER
        # =================================================

        total_baru = 0
        total_update = 0
        total_gagal = 0


        # =================================================
        # PROSES SETIAP BARIS
        # =================================================

        for nomor_baris, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            # Lewati baris yang benar-benar kosong.
            if not any(row):
                continue


            try:

                # =========================================
                # AMBIL DATA
                # =========================================

                nama = (
                    str(row[0]).strip()
                    if row[0]
                    else ''
                )

                nisn = (
                    str(row[1]).strip()
                    if row[1]
                    else None
                )

                nik = (
                    str(row[2]).strip()
                    if row[2]
                    else None
                )

                jenis_kelamin = (
                    str(row[3]).strip()
                    if row[3]
                    else ''
                )

                tempat_lahir = (
                    str(row[4]).strip()
                    if row[4]
                    else ''
                )

                tanggal_lahir = (
                    parse_tanggal_excel(
                        row[5]
                    )
                )

                alamat = (
                    str(row[6]).strip()
                    if row[6]
                    else ''
                )

                nama_kelas = (
                    str(row[7]).strip()
                    if row[7]
                    else ''
                )

                status_masuk = (
                    str(row[8])
                    .strip()
                    .upper()
                    if row[8]
                    else 'SISWA_BARU'
                )


                # =========================================
                # VALIDASI DASAR
                # =========================================

                if not nama:

                    raise ValueError(
                        'Nama siswa wajib diisi.'
                    )


                if not nisn and not nik:

                    raise ValueError(
                        'Minimal NISN atau NIK '
                        'harus diisi.'
                    )


                if not tempat_lahir:

                    raise ValueError(
                        'Tempat lahir wajib diisi.'
                    )


                if tanggal_lahir is None:

                    raise ValueError(
                        'Tanggal lahir wajib diisi.'
                    )


                if not nama_kelas:

                    raise ValueError(
                        'Kelas wajib diisi.'
                    )


                # =========================================
                # NORMALISASI JENIS KELAMIN
                # =========================================

                jk_upper = (
                    jenis_kelamin.upper()
                )


                if jk_upper in (
                    'L',
                    'LAKI-LAKI',
                    'LAKI LAKI',
                ):

                    jenis_kelamin = 'L'


                elif jk_upper in (
                    'P',
                    'PEREMPUAN',
                ):

                    jenis_kelamin = 'P'


                else:

                    raise ValueError(
                        'Jenis kelamin harus '
                        'L atau P.'
                    )


                # =========================================
                # NORMALISASI STATUS MASUK
                # =========================================

                status_normal = (
                    status_masuk
                    .strip()
                    .upper()
                    .replace('-', ' ')
                    .replace('_', ' ')
                )

                status_normal = (
                    ' '.join(
                        status_normal.split()
                    )
                )


                if status_normal in (
                    'SISWA BARU',
                    'BARU',
                ):

                    status_masuk = (
                        'SISWA_BARU'
                    )


                elif status_normal in (
                    'PINDAHAN',
                    'SISWA PINDAHAN',
                    'PINDAH',
                ):

                    status_masuk = (
                        'PINDAHAN'
                    )


                elif status_normal in (
                    'LANJUT',
                    'LANJUTAN',
                    'SISWA LANJUTAN',
                ):

                    status_masuk = (
                        'LANJUT'
                    )


                else:

                    raise ValueError(
                        (
                            'Status Masuk tidak dikenali. '
                            'Gunakan Siswa Baru, '
                            'Pindahan, atau Lanjutan.'
                        )
                    )


                # =========================================
                # CARI ROMBEL
                # =========================================

                rombel = (
                    DataRombel.objects
                    .filter(
                        sekolah=sekolah,
                        tahun_ajaran=tahun_ajaran,
                        nama_kelas__iexact=nama_kelas
                    )
                    .first()
                )


                if rombel is None:

                    raise ValueError(
                        (
                            f'Rombel "{nama_kelas}" '
                            'belum tersedia untuk '
                            f'Tahun Ajaran '
                            f'{tahun_ajaran.nama}.'
                        )
                    )


                # =========================================
                # CARI MURID BERDASARKAN NISN DAN NIK
                # =========================================

                murid_nisn = None
                murid_nik = None


                if nisn:

                    murid_nisn = (
                        Murid.objects
                        .filter(
                            nisn=nisn
                        )
                        .first()
                    )


                if nik:

                    murid_nik = (
                        Murid.objects
                        .filter(
                            nik=nik
                        )
                        .first()
                    )


                # =========================================
                # VALIDASI KONFLIK IDENTITAS
                # =========================================

                if (
                    murid_nisn
                    and
                    murid_nik
                    and
                    murid_nisn.id != murid_nik.id
                ):

                    raise ValueError(
                        (
                            'NISN dan NIK pada baris ini '
                            'terdaftar pada dua murid yang berbeda. '
                            'Periksa kembali data identitas siswa.'
                        )
                    )


                # =========================================
                # TENTUKAN MURID
                # =========================================

                murid = (
                    murid_nisn
                    or
                    murid_nik
                )


                # =========================================
                # TANDAI APAKAH BARU / UPDATE
                # =========================================

                murid_baru = (
                    murid is None
                )


                # =========================================
                # PROSES DATABASE
                # =========================================

                with transaction.atomic():

                    # =====================================
                    # UPDATE MURID
                    # =====================================

                    if murid:

                        # Tidak boleh mengambil
                        # siswa sekolah lain.
                        if (
                            murid.sekolah_id
                            != sekolah.id
                        ):

                            raise ValueError(
                                (
                                    'NISN/NIK sudah '
                                    'terdaftar pada '
                                    'sekolah lain.'
                                )
                            )


                        murid.nama = nama


                        if nisn:
                            murid.nisn = nisn


                        if nik:
                            murid.nik = nik


                        murid.jenis_kelamin = (
                            jenis_kelamin
                        )

                        murid.tempat_lahir = (
                            tempat_lahir
                        )

                        murid.tanggal_lahir = (
                            tanggal_lahir
                        )

                        murid.alamat = (
                            alamat
                        )

                        # Snapshot kelas terakhir /
                        # kelas terkini.
                        murid.kelas = (
                            rombel.nama_kelas
                        )

                        murid.save()


                    # =====================================
                    # MURID BARU
                    # =====================================

                    else:

                        murid = (
                            Murid.objects.create(

                                sekolah=sekolah,

                                nama=nama,

                                nisn=nisn,

                                nik=nik,

                                jenis_kelamin=(
                                    jenis_kelamin
                                ),

                                tempat_lahir=(
                                    tempat_lahir
                                ),

                                tanggal_lahir=(
                                    tanggal_lahir
                                ),

                                kelas=(
                                    rombel.nama_kelas
                                ),

                                alamat=alamat,
                            )
                        )


                    # =====================================
                    # RIWAYAT TAHUN AJARAN
                    # =====================================

                    riwayat = (
                        RiwayatMurid.objects
                        .filter(
                            murid=murid,
                            tahun_ajaran=tahun_ajaran
                        )
                        .first()
                    )


                    # =====================================
                    # RIWAYAT SUDAH ADA
                    # =====================================

                    if riwayat:

                        # Status Masuk lama tidak ditimpa.
                        # Import hanya boleh memperbarui
                        # rombel selama siswa masih AKTIF.

                        if (
                            riwayat.status_akhir
                            != 'AKTIF'
                        ):

                            raise ValueError(
                                (
                                    'Siswa sudah memiliki riwayat '
                                    f'dengan status '
                                    f'{riwayat.get_status_akhir_display()} '
                                    f'pada Tahun Ajaran '
                                    f'{tahun_ajaran.nama}.'
                                )
                            )


                        riwayat.rombel = (
                            rombel
                        )

                        riwayat.full_clean()

                        riwayat.save(
                            update_fields=[
                                'rombel'
                            ]
                        )


                    # =====================================
                    # BUAT RIWAYAT BARU
                    # =====================================

                    else:

                        riwayat_baru = (
                            RiwayatMurid(

                                murid=murid,

                                tahun_ajaran=(
                                    tahun_ajaran
                                ),

                                rombel=rombel,

                                status_masuk=(
                                    status_masuk
                                ),

                                status_akhir='AKTIF'
                            )
                        )


                        riwayat_baru.full_clean()

                        riwayat_baru.save()


                # =========================================
                # COUNTER HANYA SETELAH TRANSAKSI SUKSES
                # =========================================

                if murid_baru:

                    total_baru += 1

                else:

                    total_update += 1


            # =================================================
            # GAGAL SATU BARIS
            # =================================================

            except Exception as error:

                total_gagal += 1

                detail_gagal.append({

                    'baris':
                        nomor_baris,

                    'nama':
                        (
                            str(row[0])
                            if row[0]
                            else '-'
                        ),

                    'error':
                        str(error),
                })


        # =====================================================
        # HASIL IMPORT
        # =====================================================

        hasil_import = {

            'baru':
                total_baru,

            'update':
                total_update,

            'gagal':
                total_gagal,
        }


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,

        'hasil_import':
            hasil_import,

        'detail_gagal':
            detail_gagal,
    }


    # =====================================================
    # RENDER
    # =====================================================

    return render(
        request,
        'akun/import_murid_excel.html',
        context
    )


@login_required
def download_template_murid(request):

    sekolah = ambil_sekolah_operator(request)

    if sekolah is None:
        return redirect('/')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Template Murid'

    # Header
    header = [
        'Nama',
        'NISN',
        'NIK',
        'Jenis Kelamin',
        'Tempat Lahir',
        'Tanggal Lahir',
        'Alamat',
        'Kelas',
        'Status Masuk',
    ]

    sheet.append(header)

    # Contoh data
    sheet.append([
        'Ahmad Fauzan',
        '0123456789',
        '6400000000000001',
        'L',
        'Sangatta',
        '2013-05-12',
        'Jl. Pendidikan',
        'Kelas 7A',
        'SISWA_BARU',
    ])

    # Lebar kolom sederhana
    lebar_kolom = {
        'A': 30,
        'B': 18,
        'C': 22,
        'D': 18,
        'E': 20,
        'F': 18,
        'G': 35,
        'H': 18,
        'I': 20,
    }

    for kolom, lebar in lebar_kolom.items():
        sheet.column_dimensions[kolom].width = lebar

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response[
        'Content-Disposition'
    ] = (
        'attachment; '
        'filename="template_import_murid_sipetaku.xlsx"'
    )

    workbook.save(response)

    return response


@login_required
def persiapan_tahun_ajaran(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN YANG BELUM AKTIF
    # =====================================================

    daftar_tahun_tujuan = (
        TahunAjaran.objects
        .filter(
            aktif=False,
            tahun_sebelumnya=tahun_aktif
        )
        .order_by('nama')
    )


    # =====================================================
    # PILIH TAHUN TUJUAN
    # =====================================================

    tahun_id = request.GET.get(
        'tahun'
    )

    tahun_tujuan = None


    if tahun_id:

        tahun_tujuan = (
            TahunAjaran.objects
            .filter(
                id=tahun_id,
                aktif=False,
                tahun_sebelumnya=tahun_aktif
            )
            .first()
        )


    if tahun_tujuan is None:

        tahun_tujuan = (
            daftar_tahun_tujuan
            .first()
        )


    # =====================================================
    # ROMBEL TAHUN AKTIF
    # =====================================================

    rombel_tahun_aktif = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_aktif
        )
        .order_by(
            'nama_kelas'
        )
    )


    # =====================================================
    # ROMBEL TAHUN TUJUAN
    # =====================================================

    if tahun_tujuan:

        rombel_tahun_tujuan = (
            DataRombel.objects
            .filter(
                sekolah=sekolah,
                tahun_ajaran=tahun_tujuan
            )
            .order_by(
                'nama_kelas'
            )
        )

    else:

        rombel_tahun_tujuan = (
            DataRombel.objects
            .none()
        )


    context = {

        'sekolah':
            sekolah,

        'tahun_aktif':
            tahun_aktif,

        'tahun_tujuan':
            tahun_tujuan,

        'daftar_tahun_tujuan':
            daftar_tahun_tujuan,

        'rombel_tahun_aktif':
            rombel_tahun_aktif,

        'rombel_tahun_tujuan':
            rombel_tahun_tujuan,
    }


    return render(
        request,
        'akun/persiapan_tahun_ajaran.html',
        context
    )


@login_required
def tambah_rombel_tahun_berikutnya(
    request,
    tahun_id
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        return redirect('/')

    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'persiapan_tahun_ajaran'
        )

    # =====================================================
    # TAHUN AJARAN TUJUAN
    # =====================================================

    tahun_tujuan = get_object_or_404(
        TahunAjaran,
        id=tahun_id,
        aktif=False,
        tahun_sebelumnya=tahun_aktif
    )

    # =====================================================
    # INSTANCE ROMBEL BARU
    # =====================================================

    rombel_baru = DataRombel(
        sekolah=sekolah,
        tahun_ajaran=tahun_tujuan
    )

    # =====================================================
    # FORM
    # =====================================================

    if request.method == 'POST':

        form = DataRombelForm(
            request.POST,
            instance=rombel_baru,
            sekolah=sekolah
        )

        if form.is_valid():

            data = form.save(
                commit=False
            )

            data.sekolah = sekolah

            data.tahun_ajaran = (
                tahun_tujuan
            )

            data.save()

            messages.success(
                request,
                (
                    f'Rombel {data.nama_kelas} '
                    f'berhasil dibuat untuk '
                    f'Tahun Ajaran '
                    f'{tahun_tujuan.nama}.'
                )
            )

            return redirect(
                (
                    '/persiapan-tahun-ajaran/'
                    f'?tahun={tahun_tujuan.id}'
                )
            )

    else:

        form = DataRombelForm(
            instance=rombel_baru,
            sekolah=sekolah
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'tahun_tujuan':
            tahun_tujuan,

        'judul':
            'Tambah Rombel Tahun Berikutnya',

        'tombol':
            'Simpan Rombel',
    }

    return render(
        request,
        'akun/tambah_rombel_tahun_berikutnya.html',
        context
    )


@login_required
def edit_rombel_tahun_berikutnya(
    request,
    id
):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses ke data sekolah.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'persiapan_tahun_ajaran'
        )


    # =====================================================
    # AMBIL ROMBEL TAHUN PERSIAPAN
    #
    # Syarat:
    # - milik sekolah operator
    # - tahun rombel belum aktif
    # - tahun rombel merupakan kelanjutan langsung
    #   dari tahun yang sedang aktif
    # =====================================================

    rombel = get_object_or_404(
        DataRombel,
        id=id,
        sekolah=sekolah,
        tahun_ajaran__aktif=False,
        tahun_ajaran__tahun_sebelumnya=tahun_aktif
    )


    tahun_tujuan = (
        rombel.tahun_ajaran
    )


    # =====================================================
    # CEK APAKAH ROMBEL SUDAH MEMILIKI SISWA
    # =====================================================

    punya_siswa = (
        rombel.riwayat_murid
        .exists()
    )


    # =====================================================
    # FORM
    # =====================================================

    if request.method == 'POST':

        form = DataRombelForm(
            request.POST,
            instance=rombel,
            sekolah=sekolah
        )

    else:

        form = DataRombelForm(
            instance=rombel,
            sekolah=sekolah
        )


    # =====================================================
    # KUNCI IDENTITAS ROMBEL JIKA SUDAH MEMILIKI SISWA
    #
    # Tingkat dan nama kelas tidak boleh diubah.
    # Kapasitas, jumlah rombel dan keterangan masih
    # boleh diedit.
    # =====================================================

    if punya_siswa:

        form.fields[
            'tingkat'
        ].disabled = True

        form.fields[
            'tingkat'
        ].help_text = (
            'Tingkat kelas tidak dapat diubah karena '
            'rombel ini sudah memiliki data siswa.'
        )


        form.fields[
            'nama_kelas'
        ].disabled = True

        form.fields[
            'nama_kelas'
        ].help_text = (
            'Nama kelas tidak dapat diubah karena '
            'rombel ini sudah memiliki data siswa.'
        )


    # =====================================================
    # PROSES SIMPAN
    # =====================================================

    if request.method == 'POST':

        if form.is_valid():

            try:

                with transaction.atomic():

                    data = form.save(
                        commit=False
                    )


                    # =========================================
                    # KUNCI SEKOLAH DAN TAHUN AJARAN
                    #
                    # Jangan percaya nilai dari browser.
                    # =========================================

                    data.sekolah = (
                        sekolah
                    )

                    data.tahun_ajaran = (
                        tahun_tujuan
                    )


                    # =========================================
                    # VALIDASI MODEL
                    # =========================================

                    data.full_clean()

                    data.save()


            except Exception as error:

                messages.error(
                    request,
                    (
                        'Perubahan rombel gagal: '
                        f'{error}'
                    )
                )

            else:

                messages.success(
                    request,
                    (
                        f'Rombel {data.nama_kelas} '
                        f'Tahun Ajaran '
                        f'{tahun_tujuan.nama} '
                        f'berhasil diperbarui.'
                    )
                )

                return redirect(
                    (
                        '/persiapan-tahun-ajaran/'
                        f'?tahun={tahun_tujuan.id}'
                    )
                )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'sekolah':
            sekolah,

        'tahun_aktif':
            tahun_aktif,

        'tahun_tujuan':
            tahun_tujuan,

        'rombel':
            rombel,

        'punya_siswa':
            punya_siswa,

        'judul':
            'Edit Rombel Tahun Persiapan',

        'tombol':
            'Simpan Perubahan',
    }


    return render(
        request,
        'akun/tambah_rombel_tahun_berikutnya.html',
        context
    )


@login_required
@require_POST
def hapus_rombel_tahun_berikutnya(
    request,
    id
):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'persiapan_tahun_ajaran'
        )


    # =====================================================
    # ROMBEL TAHUN PERSIAPAN
    # =====================================================

    rombel = get_object_or_404(
        DataRombel,
        id=id,
        sekolah=sekolah,
        tahun_ajaran__aktif=False,
        tahun_ajaran__tahun_sebelumnya=tahun_aktif
    )

    tahun_tujuan = (
        rombel.tahun_ajaran
    )


    # =====================================================
    # CEK RIWAYAT SISWA
    # =====================================================

    if rombel.riwayat_murid.exists():

        messages.error(
            request,
            (
                f'Rombel {rombel.nama_kelas} '
                f'tidak dapat dihapus karena '
                f'sudah memiliki data siswa.'
            )
        )

        return redirect(
            (
                '/persiapan-tahun-ajaran/'
                f'?tahun={tahun_tujuan.id}'
            )
        )


    nama_rombel = (
        rombel.nama_kelas
    )

    rombel.delete()


    messages.success(
        request,
        (
            f'Rombel {nama_rombel} '
            f'Tahun Ajaran '
            f'{tahun_tujuan.nama} '
            f'berhasil dihapus.'
        )
    )


    return redirect(
        (
            '/persiapan-tahun-ajaran/'
            f'?tahun={tahun_tujuan.id}'
        )
    )


@login_required
def proses_akhir_tahun(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # TAHUN AKTIF
    # =====================================================

    tahun_aktif = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_aktif is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN TUJUAN
    # =====================================================

    tahun_tujuan = (
        TahunAjaran.objects
        .filter(
            aktif=False,
            tahun_sebelumnya=tahun_aktif
        )
        .first()
    )


    if tahun_tujuan is None:

        messages.error(
            request,
            (
                f'Belum tersedia Tahun Ajaran '
                f'berikutnya setelah '
                f'{tahun_aktif.nama}.'
            )
        )

        return redirect(
            'persiapan_tahun_ajaran'
        )


    # =====================================================
    # ROMBEL
    # =====================================================

    rombel_asal_id = request.GET.get(
        'rombel'
    )

    daftar_rombel_asal = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_aktif
        )
        .order_by('nama_kelas')
    )

    daftar_rombel_tujuan = (
        DataRombel.objects
        .filter(
            sekolah=sekolah,
            tahun_ajaran=tahun_tujuan
        )
        .order_by('nama_kelas')
    )


    # =====================================================
    # SISWA AKTIF YANG BELUM DIPROSES
    # =====================================================

    riwayat_siswa = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            tahun_ajaran=tahun_aktif,
            status_akhir='AKTIF'
        )
        .select_related(
            'murid',
            'rombel'
        )
        .order_by(
            'rombel__nama_kelas',
            'murid__nama'
        )
    )

    if rombel_asal_id:

        riwayat_siswa = (
            riwayat_siswa
            .filter(
                rombel_id=rombel_asal_id
            )
        )


    # =====================================================
    # PROSES MASSAL
    # =====================================================

    if request.method == 'POST':

        daftar_id = request.POST.getlist(
            'riwayat_ids'
        )

        aksi = request.POST.get(
            'aksi_massal'
        )

        rombel_tujuan_id = (
            request.POST.get(
                'rombel_tujuan'
            )
        )


        if not daftar_id:

            messages.error(
                request,
                'Pilih minimal satu siswa.'
            )

            return redirect(
                request.get_full_path()
            )


        aksi_valid = (
            'NAIK_KELAS',
            'TINGGAL_KELAS',
            'LULUS',
        )

        if aksi not in aksi_valid:

            messages.error(
                request,
                'Aksi tidak valid.'
            )

            return redirect(
                request.get_full_path()
            )


        rombel_tujuan = None


        if aksi in (
            'NAIK_KELAS',
            'TINGGAL_KELAS',
        ):

            if not rombel_tujuan_id:

                messages.error(
                    request,
                    'Rombel tujuan wajib dipilih.'
                )

                return redirect(
                    request.get_full_path()
                )


            rombel_tujuan = get_object_or_404(
                DataRombel,
                id=rombel_tujuan_id,
                sekolah=sekolah,
                tahun_ajaran=tahun_tujuan
            )


        daftar_riwayat = (
            RiwayatMurid.objects
            .filter(
                id__in=daftar_id,
                murid__sekolah=sekolah,
                tahun_ajaran=tahun_aktif,
                status_akhir='AKTIF'
            )
            .select_related(
                'murid',
                'rombel'
            )
        )


        # =====================================================
        # VALIDASI AKSI BERDASARKAN TINGKAT KELAS
        # =====================================================

        kategori = (
            sekolah.kategori.nama
            .strip()
            .upper()
        )

        siswa_aksi_tidak_valid = []


        for riwayat in daftar_riwayat:

            # =============================================
            # PASTIKAN SISWA MEMILIKI ROMBEL DAN TINGKAT
            # =============================================

            if (
                riwayat.rombel is None
                or
                riwayat.rombel.tingkat is None
            ):

                siswa_aksi_tidak_valid.append(
                    riwayat.murid.nama
                )

                continue


            tingkat_asal = (
                riwayat.rombel.tingkat
            )


            # =============================================
            # SD
            # =============================================

            if kategori == 'SD':

                # Kelas 1 - 5
                if tingkat_asal in (
                    1, 2, 3, 4, 5
                ):

                    if aksi not in (
                        'NAIK_KELAS',
                        'TINGGAL_KELAS',
                    ):

                        siswa_aksi_tidak_valid.append(
                            riwayat.murid.nama
                        )


                # Kelas 6
                elif tingkat_asal == 6:

                    if aksi not in (
                        'LULUS',
                        'TINGGAL_KELAS',
                    ):

                        siswa_aksi_tidak_valid.append(
                            riwayat.murid.nama
                        )


                else:

                    siswa_aksi_tidak_valid.append(
                        riwayat.murid.nama
                    )


            # =============================================
            # SMP
            # =============================================

            elif kategori == 'SMP':

                # Kelas 7 - 8
                if tingkat_asal in (
                    7, 8
                ):

                    if aksi not in (
                        'NAIK_KELAS',
                        'TINGGAL_KELAS',
                    ):

                        siswa_aksi_tidak_valid.append(
                            riwayat.murid.nama
                        )


                # Kelas 9
                elif tingkat_asal == 9:

                    if aksi not in (
                        'LULUS',
                        'TINGGAL_KELAS',
                    ):

                        siswa_aksi_tidak_valid.append(
                            riwayat.murid.nama
                        )


            else:

                siswa_aksi_tidak_valid.append(
                    riwayat.murid.nama
                )


        # =====================================================
        # JIKA AKSI TIDAK SESUAI TINGKAT
        # =====================================================

        if siswa_aksi_tidak_valid:

            messages.error(
                request,
                (
                    'Proses dibatalkan. '
                    'Aksi akademik yang dipilih tidak sesuai '
                    'dengan tingkat kelas siswa.'
                )
            )

            return redirect(
                request.get_full_path()
            )

        # =====================================================
        # VALIDASI KELULUSAN
        # BERDASARKAN TINGKAT ROMBEL
        # =====================================================

        if aksi == 'LULUS':

            kategori = (
                sekolah.kategori.nama
                .strip()
                .upper()
            )

            # =================================================
            # TENTUKAN TINGKAT AKHIR
            # =================================================

            tingkat_akhir = None

            if kategori == 'SD':

                tingkat_akhir = 6

            elif kategori == 'SMP':

                tingkat_akhir = 9


            # =================================================
            # VALIDASI KHUSUS SD DAN SMP
            # =================================================

            if tingkat_akhir is not None:

                siswa_tidak_valid = []


                for riwayat in daftar_riwayat:

                    # -----------------------------------------
                    # SISWA BELUM MEMILIKI ROMBEL
                    # -----------------------------------------

                    if riwayat.rombel is None:

                        siswa_tidak_valid.append(
                            riwayat.murid.nama
                        )

                        continue


                    # -----------------------------------------
                    # TINGKAT ROMBEL BELUM DIISI
                    # -----------------------------------------

                    if riwayat.rombel.tingkat is None:

                        siswa_tidak_valid.append(
                            riwayat.murid.nama
                        )

                        continue


                    # -----------------------------------------
                    # BUKAN TINGKAT AKHIR
                    # -----------------------------------------

                    if (
                        riwayat.rombel.tingkat
                        != tingkat_akhir
                    ):

                        siswa_tidak_valid.append(
                            riwayat.murid.nama
                        )


                # =============================================
                # JIKA ADA SISWA YANG BUKAN KELAS AKHIR
                # =============================================

                if siswa_tidak_valid:

                    messages.error(
                        request,
                        (
                            'Proses kelulusan dibatalkan. '
                            f'Untuk jenjang {kategori}, '
                            f'kelulusan hanya dapat dilakukan '
                            f'pada Kelas {tingkat_akhir}. '
                            'Terdapat siswa yang bukan '
                            'kelas akhir.'
                        )
                    )

                    return redirect(
                        request.get_full_path()
                    )


        # =====================================================
        # VALIDASI NAIK KELAS / TINGGAL KELAS
        # BERDASARKAN TINGKAT ROMBEL
        # =====================================================

        if aksi in (
            'NAIK_KELAS',
            'TINGGAL_KELAS',
        ):

            # =============================================
            # PASTIKAN ROMBEL TUJUAN MEMILIKI TINGKAT
            # =============================================

            tingkat_tujuan = (
                rombel_tujuan.tingkat
            )

            if tingkat_tujuan is None:

                messages.error(
                    request,
                    (
                        'Rombel tujuan belum memiliki '
                        'data tingkat kelas.'
                    )
                )

                return redirect(
                    request.get_full_path()
                )


            # =============================================
            # VALIDASI SETIAP SISWA
            # =============================================

            siswa_tidak_valid = []


            for riwayat in daftar_riwayat:

                # -----------------------------------------
                # Pastikan siswa memiliki rombel
                # -----------------------------------------

                if riwayat.rombel is None:

                    siswa_tidak_valid.append(
                        (
                            f'{riwayat.murid.nama} '
                            '- belum memiliki rombel asal'
                        )
                    )

                    continue


                # -----------------------------------------
                # Ambil tingkat langsung dari DataRombel
                # -----------------------------------------

                tingkat_asal = (
                    riwayat.rombel.tingkat
                )


                if tingkat_asal is None:

                    siswa_tidak_valid.append(
                        (
                            f'{riwayat.murid.nama} '
                            '- rombel asal belum memiliki '
                            'tingkat kelas'
                        )
                    )

                    continue


                # =========================================
                # NAIK KELAS
                # =========================================

                if aksi == 'NAIK_KELAS':

                    tingkat_seharusnya = (
                        tingkat_asal + 1
                    )

                    if (
                        tingkat_tujuan
                        != tingkat_seharusnya
                    ):

                        siswa_tidak_valid.append(
                            (
                                f'{riwayat.murid.nama}: '
                                f'{riwayat.rombel.nama_kelas} '
                                f'→ '
                                f'{rombel_tujuan.nama_kelas}'
                            )
                        )


                # =========================================
                # TINGGAL KELAS
                # =========================================

                elif aksi == 'TINGGAL_KELAS':

                    if (
                        tingkat_tujuan
                        != tingkat_asal
                    ):

                        siswa_tidak_valid.append(
                            (
                                f'{riwayat.murid.nama}: '
                                f'{riwayat.rombel.nama_kelas} '
                                f'→ '
                                f'{rombel_tujuan.nama_kelas}'
                            )
                        )


            # =============================================
            # JIKA ADA SISWA TIDAK VALID
            # =============================================

            if siswa_tidak_valid:

                if aksi == 'NAIK_KELAS':

                    pesan = (
                        'Proses naik kelas dibatalkan. '
                        'Rombel tujuan harus berada '
                        'satu tingkat di atas '
                        'rombel asal siswa.'
                    )

                else:

                    pesan = (
                        'Proses tinggal kelas dibatalkan. '
                        'Rombel tujuan harus memiliki '
                        'tingkat yang sama dengan '
                        'rombel asal siswa.'
                    )


                messages.error(
                    request,
                    pesan
                )

                return redirect(
                    request.get_full_path()
                )

        # =====================================================
        # VALIDASI DUPLIKASI RIWAYAT TAHUN TUJUAN
        # =====================================================

        if aksi in (
            'NAIK_KELAS',
            'TINGGAL_KELAS',
        ):

            siswa_sudah_punya_riwayat = []

            for riwayat in daftar_riwayat:

                sudah_ada = (
                    RiwayatMurid.objects
                    .filter(
                        murid=riwayat.murid,
                        tahun_ajaran=tahun_tujuan
                    )
                    .exists()
                )

                if sudah_ada:

                    siswa_sudah_punya_riwayat.append(
                        riwayat.murid.nama
                    )


            if siswa_sudah_punya_riwayat:

                messages.error(
                    request,
                    (
                        'Proses dibatalkan karena terdapat '
                        'siswa yang sudah memiliki riwayat '
                        f'pada Tahun Ajaran {tahun_tujuan.nama}. '
                        'Data riwayat tidak dibuat ulang.'
                    )
                )

                return redirect(
                    request.get_full_path()
                )


        # =====================================================
        # PERINGATAN KAPASITAS ROMBEL TUJUAN
        # =====================================================

        peringatan_kapasitas = None

        if aksi in (
            'NAIK_KELAS',
            'TINGGAL_KELAS',
        ):

            jumlah_siswa_saat_ini = (
                rombel_tujuan.jumlah_siswa_aktif
            )

            jumlah_siswa_masuk = (
                daftar_riwayat.count()
            )

            jumlah_setelah_proses = (
                jumlah_siswa_saat_ini
                + jumlah_siswa_masuk
            )

            kapasitas_total = (
                rombel_tujuan.kapasitas_total
            )


            if (
                kapasitas_total > 0
                and
                jumlah_setelah_proses
                > kapasitas_total
            ):

                kelebihan = (
                    jumlah_setelah_proses
                    - kapasitas_total
                )

                lanjut_kapasitas = (
                    request.POST.get(
                        'lanjut_kapasitas'
                    )
                )


                if lanjut_kapasitas != 'YA':

                    messages.warning(
                        request,
                        (
                            f'Rombel {rombel_tujuan.nama_kelas} '
                            f'memiliki kapasitas ideal '
                            f'{kapasitas_total} siswa. '
                            f'Saat ini terdapat '
                            f'{jumlah_siswa_saat_ini} siswa dan '
                            f'{jumlah_siswa_masuk} siswa akan '
                            f'diproses masuk. '
                            f'Setelah proses jumlah siswa menjadi '
                            f'{jumlah_setelah_proses}, sehingga '
                            f'melebihi kapasitas sebanyak '
                            f'{kelebihan} siswa.'
                        )
                    )

                    return render(
                        request,
                        'akun/konfirmasi_kapasitas_rombel.html',
                        {
                            'sekolah':
                                sekolah,

                            'tahun_aktif':
                                tahun_aktif,

                            'tahun_tujuan':
                                tahun_tujuan,

                            'rombel_tujuan':
                                rombel_tujuan,

                            'jumlah_siswa_saat_ini':
                                jumlah_siswa_saat_ini,

                            'jumlah_siswa_masuk':
                                jumlah_siswa_masuk,

                            'jumlah_setelah_proses':
                                jumlah_setelah_proses,

                            'kapasitas_total':
                                kapasitas_total,

                            'kelebihan':
                                kelebihan,

                            'daftar_id':
                                daftar_id,

                            'aksi':
                                aksi,

                            'rombel_asal_id':
                                rombel_asal_id,
                        }
                    )


        total_diproses = 0


        try:

            with transaction.atomic():

                for riwayat in daftar_riwayat:

                    # =====================================
                    # NAIK KELAS / TINGGAL KELAS
                    # =====================================

                    if aksi in (
                        'NAIK_KELAS',
                        'TINGGAL_KELAS',
                    ):

                        riwayat.status_akhir = (
                            aksi
                        )

                        riwayat.save(
                            update_fields=[
                                'status_akhir'
                            ]
                        )


                        riwayat_baru = RiwayatMurid(
                            murid=riwayat.murid,
                            rombel=rombel_tujuan,
                            tahun_ajaran=tahun_tujuan,
                            status_masuk=aksi,
                            status_akhir='AKTIF'
                        )

                        riwayat_baru.full_clean()
                        riwayat_baru.save()


                        riwayat.murid.kelas = (
                            rombel_tujuan.nama_kelas
                        )

                        riwayat.murid.save(
                            update_fields=[
                                'kelas'
                            ]
                        )



                    # =====================================
                    # LULUS
                    # =====================================

                    elif aksi == 'LULUS':

                        tahun_lulus = int(
                            tahun_aktif
                            .nama
                            .split('/')[-1]
                        )

                        riwayat.status_akhir = (
                            'LULUS'
                        )

                        riwayat.tahun_lulus = (
                            tahun_lulus
                        )

                        riwayat.save(
                            update_fields=[
                                'status_akhir',
                                'tahun_lulus'
                            ]
                        )


                    total_diproses += 1


        except Exception as error:

            messages.error(
                request,
                (
                    'Proses dibatalkan karena terjadi '
                    f'kesalahan: {error}'
                )
            )

            return redirect(
                request.get_full_path()
            )


        messages.success(
            request,
            (
                f'{total_diproses} siswa '
                'berhasil diproses.'
            )
        )


        url_kembali = (
            '/proses-akhir-tahun/'
        )

        if rombel_asal_id:

            url_kembali += (
                f'?rombel={rombel_asal_id}'
            )

        return redirect(
            url_kembali
        )


    context = {

        'sekolah':
            sekolah,

        'tahun_aktif':
            tahun_aktif,

        'tahun_tujuan':
            tahun_tujuan,

        'daftar_rombel_asal':
            daftar_rombel_asal,

        'daftar_rombel_tujuan':
            daftar_rombel_tujuan,

        'riwayat_siswa':
            riwayat_siswa,

        'rombel_asal_terpilih':
            rombel_asal_id,
    }


    return render(
        request,
        'akun/proses_akhir_tahun.html',
        context
    )


@login_required
def aktifkan_tahun_ajaran(
    request,
    tahun_id
):

    # =====================================================
    # HAK AKSES
    # =====================================================

    if not user_admin_kabupaten(request):

        messages.error(
            request,
            (
                'Aktivasi Tahun Ajaran hanya dapat '
                'dilakukan oleh Admin Kabupaten.'
            )
        )

        return redirect('/')


    # =====================================================
    # HANYA POST
    # =====================================================

    if request.method != 'POST':

        messages.error(
            request,
            'Permintaan aktivasi tidak valid.'
        )

        return redirect(
            'monitoring_kesiapan_tahun_ajaran'
        )


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_lama = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_lama is None:

        messages.error(
            request,
            'Tidak ditemukan Tahun Ajaran aktif.'
        )

        return redirect(
            'kelola_tahun_ajaran'
        )


    # =====================================================
    # TAHUN TUJUAN
    #
    # Harus:
    # - belum aktif
    # - penerus langsung Tahun Ajaran aktif
    # =====================================================

    tahun_baru = get_object_or_404(
        TahunAjaran,
        id=tahun_id,
        aktif=False,
        tahun_sebelumnya=tahun_lama
    )


    # =====================================================
    # DAFTAR SELURUH SEKOLAH
    # =====================================================

    daftar_sekolah = (
        Sekolah.objects
        .all()
        .order_by(
            'nama'
        )
    )


    total_sekolah = (
        daftar_sekolah.count()
    )


    # =====================================================
    # PENGAMAN:
    # JANGAN AKTIFKAN JIKA TIDAK ADA SEKOLAH
    # =====================================================

    if total_sekolah == 0:

        messages.error(
            request,
            (
                'Tahun Ajaran tidak dapat diaktifkan '
                'karena belum terdapat data sekolah '
                'di dalam sistem.'
            )
        )

        return redirect(
            'monitoring_kesiapan_tahun_ajaran'
        )


    # =====================================================
    # CEK KESIAPAN SELURUH SEKOLAH
    # =====================================================

    total_siap = 0

    sekolah_belum_siap = []


    for sekolah in daftar_sekolah:

        kesiapan = (
            hitung_kesiapan_tahun_ajaran(
                sekolah,
                tahun_lama,
                tahun_baru
            )
        )


        if kesiapan['siap']:

            total_siap += 1

        else:

            sekolah_belum_siap.append({

                'sekolah':
                    sekolah,

                'status':
                    kesiapan['status'],

                'belum_diproses':
                    kesiapan[
                        'belum_diproses'
                    ],

                'rombel':
                    kesiapan[
                        'jumlah_rombel_tujuan'
                    ],
            })


    # =====================================================
    # TOLAK JIKA MASIH ADA SEKOLAH BELUM SIAP
    # =====================================================

    if sekolah_belum_siap:

        jumlah_belum_siap = (
            len(
                sekolah_belum_siap
            )
        )


        messages.error(
            request,
            (
                f'Tahun Ajaran {tahun_baru.nama} '
                f'belum dapat diaktifkan. '
                f'{total_siap} dari '
                f'{total_sekolah} sekolah '
                f'telah dinyatakan siap. '
                f'Masih terdapat '
                f'{jumlah_belum_siap} sekolah '
                f'yang belum memenuhi '
                f'kriteria kesiapan.'
            )
        )


        return redirect(
            'monitoring_kesiapan_tahun_ajaran'
        )


    # =====================================================
    # PEGAWAI AKTIF TAHUN LAMA
    # =====================================================

    pegawai_aktif_tahun_lama = (
        RiwayatPegawai.objects
        .filter(
            tahun_ajaran=tahun_lama,
            status_akhir='AKTIF'
        )
        .select_related(
            'pegawai',
            'sekolah'
        )
        .order_by(
            'pegawai__nama'
        )
    )


    total_pegawai_lanjutan = (
        pegawai_aktif_tahun_lama.count()
    )


    # =====================================================
    # AKTIVASI GLOBAL + LANJUTKAN PEGAWAI
    # =====================================================

    try:

        with transaction.atomic():

            # =================================================
            # LANJUTKAN PEGAWAI AKTIF KE TAHUN BARU
            # =================================================

            total_pegawai_dibuat = 0
            total_pegawai_sudah_ada = 0


            for riwayat_lama in (
                pegawai_aktif_tahun_lama
            ):

                # ---------------------------------------------
                # PENGAMAN DUPLIKASI
                # ---------------------------------------------

                sudah_ada = (
                    RiwayatPegawai.objects
                    .filter(
                        pegawai=(
                            riwayat_lama.pegawai
                        ),
                        tahun_ajaran=tahun_baru,
                        status_akhir='AKTIF'
                    )
                    .exists()
                )


                if sudah_ada:

                    total_pegawai_sudah_ada += 1

                    continue


                # ---------------------------------------------
                # BUAT RIWAYAT TAHUN BARU
                # ---------------------------------------------

                riwayat_baru = (
                    RiwayatPegawai(
                        pegawai=(
                            riwayat_lama.pegawai
                        ),

                        sekolah=(
                            riwayat_lama.sekolah
                        ),

                        tahun_ajaran=(
                            tahun_baru
                        ),

                        jenis_pegawai=(
                            riwayat_lama
                            .jenis_pegawai
                        ),

                        jabatan=(
                            riwayat_lama
                            .jabatan
                        ),

                        mata_pelajaran=(
                            riwayat_lama
                            .mata_pelajaran
                        ),

                        status_pegawai=(
                            riwayat_lama
                            .status_pegawai
                        ),

                        status_akhir='AKTIF',

                        # Tanggal mulai penempatan
                        # tetap dipertahankan.
                        tanggal_mulai=(
                            riwayat_lama
                            .tanggal_mulai
                        ),

                        tanggal_selesai=None,

                        tanggal_pensiun=(
                            riwayat_lama
                            .tanggal_pensiun
                        ),

                        keterangan=(
                            riwayat_lama
                            .keterangan
                        ),
                    )
                )


                riwayat_baru.full_clean()

                riwayat_baru.save()

                total_pegawai_dibuat += 1


            # =================================================
            # AKTIFKAN TAHUN AJARAN BARU
            # =================================================

            tahun_baru.aktif = True

            tahun_baru.save(
                update_fields=[
                    'aktif'
                ]
            )


    except Exception as error:

        messages.error(
            request,
            (
                'Aktivasi Tahun Ajaran gagal: '
                f'{error}'
            )
        )

        return redirect(
            'monitoring_kesiapan_tahun_ajaran'
        )


    # =====================================================
    # BERHASIL
    # =====================================================

    pesan_pegawai = (
        f'{total_pegawai_dibuat} pegawai aktif '
        f'berhasil dilanjutkan ke Tahun Ajaran '
        f'{tahun_baru.nama}.'
    )


    if total_pegawai_sudah_ada > 0:

        pesan_pegawai += (
            f' {total_pegawai_sudah_ada} pegawai '
            f'telah memiliki riwayat aktif pada '
            f'Tahun Ajaran tujuan sehingga tidak '
            f'dibuat ulang.'
        )


    messages.success(
        request,
        (
            f'Tahun Ajaran {tahun_baru.nama} '
            f'berhasil diaktifkan secara global. '
            f'Seluruh {total_sekolah} sekolah '
            f'telah memenuhi kriteria kesiapan. '
            f'{pesan_pegawai}'
        )
    )


    return redirect(
        'kelola_tahun_ajaran'
    )


@login_required
def riwayat_murid(request, id):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    murid = get_object_or_404(
        Murid,
        id=id,
        sekolah=sekolah
    )


    daftar_riwayat = (
        RiwayatMurid.objects
        .filter(
            murid=murid
        )
        .select_related(
            'tahun_ajaran',
            'rombel'
        )
        .order_by(
            'tahun_ajaran__nama'
        )
    )


    context = {

        'sekolah':
            sekolah,

        'murid':
            murid,

        'daftar_riwayat':
            daftar_riwayat,
    }


    return render(
        request,
        'akun/riwayat_murid.html',
        context
    )


@login_required
def arsip_lulusan(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # FILTER TAHUN LULUS
    # =====================================================

    tahun_lulus = request.GET.get(
        'tahun_lulus'
    )


    daftar_lulusan = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            status_akhir='LULUS'
        )
        .select_related(
            'murid',
            'rombel',
            'tahun_ajaran'
        )
        .order_by(
            '-tahun_lulus',
            'murid__nama'
        )
    )


    if tahun_lulus:

        daftar_lulusan = (
            daftar_lulusan
            .filter(
                tahun_lulus=tahun_lulus
            )
        )


    # =====================================================
    # PILIHAN TAHUN
    # =====================================================

    daftar_tahun_lulus = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            status_akhir='LULUS',
            tahun_lulus__isnull=False
        )
        .values_list(
            'tahun_lulus',
            flat=True
        )
        .distinct()
        .order_by(
            '-tahun_lulus'
        )
    )


    total_lulusan = (
        daftar_lulusan.count()
    )


    total_laki_laki = (
        daftar_lulusan
        .filter(
            murid__jenis_kelamin='L'
        )
        .count()
    )


    total_perempuan = (
        daftar_lulusan
        .filter(
            murid__jenis_kelamin='P'
        )
        .count()
    )


    context = {

        'sekolah':
            sekolah,

        'daftar_lulusan':
            daftar_lulusan,

        'daftar_tahun_lulus':
            daftar_tahun_lulus,

        'tahun_lulus_terpilih':
            tahun_lulus,

        'total_lulusan':
            total_lulusan,

        'total_laki_laki':
            total_laki_laki,

        'total_perempuan':
            total_perempuan,
    }


    return render(
        request,
        'akun/arsip_lulusan.html',
        context
    )


@login_required
def arsip_siswa_keluar(request):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    # =====================================================
    # FILTER STATUS
    # =====================================================

    status = request.GET.get(
        'status'
    )


    daftar_siswa = (
        RiwayatMurid.objects
        .filter(
            murid__sekolah=sekolah,
            status_akhir__in=[
                'PINDAH',
                'KELUAR',
            ]
        )
        .select_related(
            'murid',
            'rombel',
            'tahun_ajaran'
        )
        .order_by(
            '-tahun_ajaran__nama',
            'murid__nama'
        )
    )


    if status in (
        'PINDAH',
        'KELUAR',
    ):

        daftar_siswa = (
            daftar_siswa
            .filter(
                status_akhir=status
            )
        )


    # =====================================================
    # RINGKASAN
    # =====================================================

    total_siswa = (
        daftar_siswa.count()
    )

    total_pindah = (
        daftar_siswa
        .filter(
            status_akhir='PINDAH'
        )
        .count()
    )

    total_keluar = (
        daftar_siswa
        .filter(
            status_akhir='KELUAR'
        )
        .count()
    )


    context = {

        'sekolah':
            sekolah,

        'daftar_siswa':
            daftar_siswa,

        'status_terpilih':
            status,

        'total_siswa':
            total_siswa,

        'total_pindah':
            total_pindah,

        'total_keluar':
            total_keluar,
    }


    return render(
        request,
        'akun/arsip_siswa_keluar.html',
        context
    )


@login_required
def data_pegawai(request):

    # =====================================================
    # PROFIL + HAK AKSES
    # =====================================================

    profil = ambil_profil_user(request)

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(aktif=True)
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect('/')


    # =====================================================
    # QUERY DASAR
    # =====================================================

    riwayat_pegawai = (
        RiwayatPegawai.objects
        .filter(
            tahun_ajaran=tahun_ajaran,
            status_akhir='AKTIF'
        )
        .select_related(
            'pegawai',
            'sekolah',
            'tahun_ajaran'
        )
    )


    # =====================================================
    # BATASI BERDASARKAN ROLE
    # =====================================================

    sekolah_operator = None


    if request.user.is_superuser:

        pass


    elif (
        profil
        and
        profil.role == 'admin_kabupaten'
    ):

        pass


    elif (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                sekolah__kecamatan=profil.kecamatan
            )
        )


    elif (
        profil
        and
        profil.role == 'operator_sekolah'
    ):

        sekolah_operator = (
            ambil_sekolah_operator(
                request
            )
        )

        if sekolah_operator is None:

            messages.error(
                request,
                'Operator belum terhubung '
                'dengan sekolah.'
            )

            return redirect('/')


        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                sekolah=sekolah_operator
            )
        )


    else:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke Data Pegawai.'
        )

        return redirect('/')


    # =====================================================
    # FILTER
    # =====================================================

    jenis_pegawai = (
        request.GET.get(
            'jenis',
            ''
        )
        .strip()
        .upper()
    )

    status_pegawai = (
        request.GET.get(
            'status',
            ''
        )
        .strip()
        .upper()
    )

    kecamatan = (
        request.GET.get(
            'kecamatan',
            ''
        )
        .strip()
    )

    sekolah_id = (
        request.GET.get(
            'sekolah',
            ''
        )
        .strip()
    )

    cari = (
        request.GET.get(
            'cari',
            ''
        )
        .strip()
    )


    # =====================================================
    # FILTER QUERY
    # =====================================================

    if jenis_pegawai in (
        'PENDIDIK',
        'TENAGA_KEPENDIDIKAN',
    ):

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                jenis_pegawai=jenis_pegawai
            )
        )


    if status_pegawai:

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                status_pegawai=status_pegawai
            )
        )


    if (
        kecamatan
        and (
            request.user.is_superuser
            or (
                profil
                and
                profil.role == 'admin_kabupaten'
            )
        )
    ):

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                sekolah__kecamatan=kecamatan
            )
        )


    if sekolah_id:

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                sekolah_id=sekolah_id
            )
        )


    if cari:

        riwayat_pegawai = (
            riwayat_pegawai
            .filter(
                Q(
                    pegawai__nama__icontains=cari
                )
                |
                Q(
                    pegawai__nip__icontains=cari
                )
                |
                Q(
                    pegawai__nik__icontains=cari
                )
                |
                Q(
                    pegawai__nuptk__icontains=cari
                )
                |
                Q(
                    jabatan__icontains=cari
                )
                |
                Q(
                    mata_pelajaran__icontains=cari
                )
                |
                Q(
                    sekolah__nama__icontains=cari
                )
            )
        )


    riwayat_pegawai = (
        riwayat_pegawai
        .order_by(
            'sekolah__kecamatan',
            'sekolah__nama',
            'pegawai__nama'
        )
    )


    # =====================================================
    # RINGKASAN BERDASARKAN DATA YANG BOLEH DIAKSES
    # =====================================================

    ringkasan_qs = (
        RiwayatPegawai.objects
        .filter(
            tahun_ajaran=tahun_ajaran,
            status_akhir='AKTIF'
        )
    )


    if (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        ringkasan_qs = (
            ringkasan_qs
            .filter(
                sekolah__kecamatan=profil.kecamatan
            )
        )


    elif (
        profil
        and
        profil.role == 'operator_sekolah'
        and sekolah_operator
    ):

        ringkasan_qs = (
            ringkasan_qs
            .filter(
                sekolah=sekolah_operator
            )
        )


    total_pegawai = (
        ringkasan_qs.count()
    )

    total_pendidik = (
        ringkasan_qs
        .filter(
            jenis_pegawai='PENDIDIK'
        )
        .count()
    )

    total_tenaga_kependidikan = (
        ringkasan_qs
        .filter(
            jenis_pegawai='TENAGA_KEPENDIDIKAN'
        )
        .count()
    )

    total_laki_laki = (
        ringkasan_qs
        .filter(
            pegawai__jenis_kelamin='L'
        )
        .count()
    )

    total_perempuan = (
        ringkasan_qs
        .filter(
            pegawai__jenis_kelamin='P'
        )
        .count()
    )


    # =====================================================
    # PILIHAN FILTER
    # =====================================================

    if (
        request.user.is_superuser
        or (
            profil
            and
            profil.role == 'admin_kabupaten'
        )
    ):

        daftar_kecamatan = (
            Sekolah.objects
            .exclude(
                kecamatan=''
            )
            .values_list(
                'kecamatan',
                flat=True
            )
            .distinct()
            .order_by(
                'kecamatan'
            )
        )

        daftar_sekolah = (
            Sekolah.objects
            .all()
            .order_by(
                'kecamatan',
                'nama'
            )
        )


    elif (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        daftar_kecamatan = (
            Sekolah.objects
            .filter(
                kecamatan=profil.kecamatan
            )
            .values_list(
                'kecamatan',
                flat=True
            )
            .distinct()
        )

        daftar_sekolah = (
            Sekolah.objects
            .filter(
                kecamatan=profil.kecamatan
            )
            .order_by(
                'nama'
            )
        )


    else:

        daftar_kecamatan = []

        daftar_sekolah = (
            Sekolah.objects
            .filter(
                id=sekolah_operator.id
            )
            if sekolah_operator
            else Sekolah.objects.none()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'tahun_ajaran':
            tahun_ajaran,

        'riwayat_pegawai':
            riwayat_pegawai,

        'total_pegawai':
            total_pegawai,

        'total_pendidik':
            total_pendidik,

        'total_tenaga_kependidikan':
            total_tenaga_kependidikan,

        'total_laki_laki':
            total_laki_laki,

        'total_perempuan':
            total_perempuan,

        'pilihan_status':
            RiwayatPegawai.STATUS_PEGAWAI,

        'daftar_kecamatan':
            daftar_kecamatan,

        'daftar_sekolah':
            daftar_sekolah,

        'filter_jenis':
            jenis_pegawai,

        'filter_status':
            status_pegawai,

        'filter_kecamatan':
            kecamatan,

        'filter_sekolah':
            sekolah_id,

        'filter_cari':
            cari,

        'admin_kabupaten':
            (
                request.user.is_superuser
                or (
                    profil
                    and
                    profil.role == 'admin_kabupaten'
                )
            ),

        'admin_kecamatan':
            (
                profil
                and
                profil.role == 'admin_kecamatan'
            ),

        'operator_sekolah':
            (
                profil
                and
                profil.role == 'operator_sekolah'
            ),
    }


    return render(
        request,
        'akun/data_pegawai.html',
        context
    )


@login_required
def tambah_pegawai(request):

    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'untuk menambah Data Pegawai.'
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_pegawai'
        )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form_pegawai = (
            PegawaiForm(
                request.POST
            )
        )

        form_riwayat = (
            RiwayatPegawaiForm(
                request.POST
            )
        )


        if (
            form_pegawai.is_valid()
            and
            form_riwayat.is_valid()
        ):

            try:

                with transaction.atomic():

                    # =====================================
                    # SIMPAN IDENTITAS PEGAWAI
                    # =====================================

                    pegawai = (
                        form_pegawai.save(
                            commit=False
                        )
                    )

                    pegawai.full_clean()

                    pegawai.save()


                    # =====================================
                    # SIMPAN RIWAYAT PENEMPATAN
                    # =====================================

                    riwayat = (
                        form_riwayat.save(
                            commit=False
                        )
                    )

                    riwayat.pegawai = (
                        pegawai
                    )

                    riwayat.sekolah = (
                        sekolah
                    )

                    riwayat.tahun_ajaran = (
                        tahun_ajaran
                    )

                    riwayat.status_akhir = (
                        'AKTIF'
                    )

                    riwayat.full_clean()

                    riwayat.save()


                messages.success(
                    request,
                    (
                        f'{pegawai.nama} berhasil '
                        f'ditambahkan sebagai pegawai '
                        f'{sekolah.nama} pada '
                        f'Tahun Ajaran '
                        f'{tahun_ajaran.nama}.'
                    )
                )


                return redirect(
                    'data_pegawai'
                )


            except Exception as error:

                messages.error(
                    request,
                    (
                        'Data pegawai gagal disimpan. '
                        f'{error}'
                    )
                )


    # =====================================================
    # GET
    # =====================================================

    else:

        form_pegawai = (
            PegawaiForm()
        )

        form_riwayat = (
            RiwayatPegawaiForm()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form_pegawai':
            form_pegawai,

        'form_riwayat':
            form_riwayat,

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,
    }


    return render(
        request,
        'akun/tambah_pegawai.html',
        context
    )


@login_required
def edit_pegawai(request, id):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_pegawai'
        )


    riwayat = get_object_or_404(
        RiwayatPegawai.objects.select_related(
            'pegawai',
            'sekolah',
            'tahun_ajaran'
        ),
        id=id,
        sekolah=sekolah,
        tahun_ajaran=tahun_ajaran,
        status_akhir='AKTIF'
    )

    pegawai = riwayat.pegawai


    if request.method == 'POST':

        form_pegawai = PegawaiForm(
            request.POST,
            instance=pegawai
        )

        form_riwayat = RiwayatPegawaiForm(
            request.POST,
            instance=riwayat
        )


        if (
            form_pegawai.is_valid()
            and
            form_riwayat.is_valid()
        ):

            try:

                with transaction.atomic():

                    pegawai = form_pegawai.save()

                    riwayat = (
                        form_riwayat.save(
                            commit=False
                        )
                    )

                    # Tetap dikunci ke penempatan aktif.
                    riwayat.pegawai = pegawai
                    riwayat.sekolah = sekolah
                    riwayat.tahun_ajaran = (
                        tahun_ajaran
                    )
                    riwayat.status_akhir = (
                        'AKTIF'
                    )

                    riwayat.full_clean()
                    riwayat.save()


                messages.success(
                    request,
                    (
                        f'Data pegawai '
                        f'{pegawai.nama} '
                        f'berhasil diperbarui.'
                    )
                )

                return redirect(
                    'data_pegawai'
                )


            except Exception as error:

                messages.error(
                    request,
                    (
                        'Data pegawai gagal '
                        f'diperbarui. {error}'
                    )
                )


    else:

        form_pegawai = PegawaiForm(
            instance=pegawai
        )

        form_riwayat = RiwayatPegawaiForm(
            instance=riwayat
        )


    context = {

        'form_pegawai':
            form_pegawai,

        'form_riwayat':
            form_riwayat,

        'pegawai':
            pegawai,

        'riwayat':
            riwayat,

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,
    }


    return render(
        request,
        'akun/edit_pegawai.html',
        context
    )


@login_required
def riwayat_pegawai(request, id):

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:
        return redirect('/')


    pegawai = get_object_or_404(
        Pegawai,
        id=id
    )


    # Pastikan pegawai ini memang pernah / sedang
    # terkait dengan sekolah operator.
    punya_akses = (
        RiwayatPegawai.objects
        .filter(
            pegawai=pegawai,
            sekolah=sekolah
        )
        .exists()
    )


    if not punya_akses:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke riwayat pegawai ini.'
        )

        return redirect(
            'data_pegawai'
        )


    daftar_riwayat = (
        RiwayatPegawai.objects
        .filter(
            pegawai=pegawai
        )
        .select_related(
            'sekolah',
            'tahun_ajaran'
        )
        .order_by(
            '-tahun_ajaran__nama',
            '-tanggal_mulai',
            '-id'
        )
    )


    context = {

        'pegawai':
            pegawai,

        'sekolah':
            sekolah,

        'daftar_riwayat':
            daftar_riwayat,
    }


    return render(
        request,
        'akun/riwayat_pegawai.html',
        context
    )


@login_required
def mutasi_pegawai(request, id):

    # =====================================================
    # HAK AKSES MUTASI
    # =====================================================

    profil = getattr(
        request.user,
        'profil',
        None
    )

    boleh_mutasi = (
        request.user.is_superuser
        or (
            profil
            and
            profil.role == 'admin_kabupaten'
        )
    )

    if not boleh_mutasi:

        messages.error(
            request,
            (
                'Anda tidak memiliki kewenangan '
                'untuk melakukan mutasi pegawai.'
            )
        )

        return redirect(
            'data_pegawai'
        )


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_pegawai'
        )


    # =====================================================
    # RIWAYAT AKTIF
    # =====================================================

    riwayat_lama = get_object_or_404(
        RiwayatPegawai.objects
        .select_related(
            'pegawai',
            'sekolah',
            'tahun_ajaran'
        ),
        id=id,
        tahun_ajaran=tahun_ajaran,
        status_akhir='AKTIF'
    )

    pegawai = riwayat_lama.pegawai

    sekolah = riwayat_lama.sekolah


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = MutasiPegawaiForm(
            request.POST,
            sekolah_asal=sekolah
        )


        if form.is_valid():

            sekolah_tujuan = (
                form.cleaned_data[
                    'sekolah_tujuan'
                ]
            )

            tanggal_mutasi = (
                form.cleaned_data[
                    'tanggal_mutasi'
                ]
            )


            # =============================================
            # VALIDASI TANGGAL MUTASI
            # =============================================

            if (
                riwayat_lama.tanggal_mulai
                and
                tanggal_mutasi
                < riwayat_lama.tanggal_mulai
            ):

                form.add_error(
                    'tanggal_mutasi',
                    (
                        'Tanggal mutasi tidak boleh '
                        'lebih awal dari tanggal mulai '
                        'penempatan saat ini.'
                    )
                )


            else:

                try:

                    with transaction.atomic():

                        # =================================
                        # KUNCI RIWAYAT LAMA
                        # =================================

                        riwayat_lama = (
                            RiwayatPegawai.objects
                            .select_for_update()
                            .get(
                                id=riwayat_lama.id
                            )
                        )


                        # =================================
                        # PASTIKAN MASIH AKTIF
                        # =================================

                        if (
                            riwayat_lama.status_akhir
                            != 'AKTIF'
                        ):

                            raise ValueError(
                                (
                                    'Riwayat pegawai ini '
                                    'sudah tidak aktif.'
                                )
                            )


                        # =================================
                        # TUTUP RIWAYAT LAMA
                        # =================================

                        riwayat_lama.status_akhir = (
                            'MUTASI'
                        )

                        riwayat_lama.tanggal_selesai = (
                            tanggal_mutasi
                        )

                        riwayat_lama.full_clean()

                        riwayat_lama.save(
                            update_fields=[
                                'status_akhir',
                                'tanggal_selesai',
                                'updated_at',
                            ]
                        )


                        # =================================
                        # BUAT RIWAYAT BARU
                        # =================================

                        riwayat_baru = (
                            RiwayatPegawai(
                                pegawai=pegawai,
                                sekolah=sekolah_tujuan,
                                tahun_ajaran=tahun_ajaran,

                                jenis_pegawai=(
                                    form.cleaned_data[
                                        'jenis_pegawai'
                                    ]
                                ),

                                jabatan=(
                                    form.cleaned_data[
                                        'jabatan'
                                    ]
                                ),

                                mata_pelajaran=(
                                    form.cleaned_data[
                                        'mata_pelajaran'
                                    ]
                                ),

                                status_pegawai=(
                                    form.cleaned_data[
                                        'status_pegawai'
                                    ]
                                ),

                                status_akhir='AKTIF',

                                tanggal_mulai=(
                                    tanggal_mutasi
                                ),

                                tanggal_selesai=None,

                                tanggal_pensiun=(
                                    riwayat_lama
                                    .tanggal_pensiun
                                ),

                                keterangan=(
                                    form.cleaned_data[
                                        'keterangan'
                                    ]
                                ),
                            )
                        )

                        riwayat_baru.full_clean()

                        riwayat_baru.save()


                    messages.success(
                        request,
                        (
                            f'{pegawai.nama} berhasil '
                            f'dimutasi dari '
                            f'{sekolah.nama} ke '
                            f'{sekolah_tujuan.nama}.'
                        )
                    )

                    return redirect(
                        'data_pegawai'
                    )


                except Exception as error:

                    messages.error(
                        request,
                        (
                            'Mutasi pegawai gagal. '
                            f'{error}'
                        )
                    )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = MutasiPegawaiForm(
            sekolah_asal=sekolah,
            initial={

                'jenis_pegawai':
                    riwayat_lama.jenis_pegawai,

                'jabatan':
                    riwayat_lama.jabatan,

                'mata_pelajaran':
                    riwayat_lama.mata_pelajaran,

                'status_pegawai':
                    riwayat_lama.status_pegawai,
            }
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'form':
            form,

        'pegawai':
            pegawai,

        'riwayat_lama':
            riwayat_lama,

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,
    }


    return render(
        request,
        'akun/mutasi_pegawai.html',
        context
    )


@login_required
def akhiri_pegawai(request, id):

    # =====================================================
    # HAK AKSES
    # =====================================================

    profil = ambil_profil_user(
        request
    )

    boleh_proses = (
        request.user.is_superuser
        or (
            profil
            and
            profil.role == 'admin_kabupaten'
        )
    )

    if not boleh_proses:

        messages.error(
            request,
            (
                'Anda tidak memiliki kewenangan '
                'untuk mengakhiri status pegawai.'
            )
        )

        return redirect(
            'data_pegawai'
        )


    # =====================================================
    # RIWAYAT AKTIF
    # =====================================================

    riwayat = get_object_or_404(
        RiwayatPegawai.objects
        .select_related(
            'pegawai',
            'sekolah',
            'tahun_ajaran'
        ),
        id=id,
        status_akhir='AKTIF'
    )

    pegawai = (
        riwayat.pegawai
    )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        form = AkhiriPegawaiForm(
            request.POST
        )


        if form.is_valid():

            status_akhir = (
                form.cleaned_data[
                    'status_akhir'
                ]
            )

            tanggal_selesai = (
                form.cleaned_data[
                    'tanggal_selesai'
                ]
            )

            keterangan = (
                form.cleaned_data[
                    'keterangan'
                ]
            )


            # =============================================
            # VALIDASI TANGGAL
            # =============================================

            if (
                riwayat.tanggal_mulai
                and
                tanggal_selesai
                < riwayat.tanggal_mulai
            ):

                form.add_error(
                    'tanggal_selesai',
                    (
                        'Tanggal efektif tidak boleh '
                        'lebih awal dari tanggal mulai '
                        'penempatan.'
                    )
                )


            else:

                try:

                    with transaction.atomic():

                        riwayat = (
                            RiwayatPegawai.objects
                            .select_for_update()
                            .get(
                                id=riwayat.id
                            )
                        )


                        if (
                            riwayat.status_akhir
                            != 'AKTIF'
                        ):

                            raise ValueError(
                                (
                                    'Pegawai ini sudah '
                                    'tidak berstatus aktif.'
                                )
                            )


                        riwayat.status_akhir = (
                            status_akhir
                        )

                        riwayat.tanggal_selesai = (
                            tanggal_selesai
                        )


                        if keterangan:

                            riwayat.keterangan = (
                                keterangan
                            )


                        # =================================
                        # KHUSUS PENSIUN
                        # =================================

                        if (
                            status_akhir
                            == 'PENSIUN'
                        ):

                            riwayat.tanggal_pensiun = (
                                tanggal_selesai
                            )


                        riwayat.full_clean()

                        riwayat.save()


                    messages.success(
                        request,
                        (
                            f'Status {pegawai.nama} '
                            f'berhasil diubah menjadi '
                            f'{riwayat.get_status_akhir_display()}.'
                        )
                    )

                    return redirect(
                        'data_pegawai'
                    )


                except Exception as error:

                    messages.error(
                        request,
                        (
                            'Perubahan status pegawai '
                            f'gagal. {error}'
                        )
                    )


    # =====================================================
    # GET
    # =====================================================

    else:

        form = AkhiriPegawaiForm()


    context = {

        'form':
            form,

        'pegawai':
            pegawai,

        'riwayat':
            riwayat,
    }


    return render(
        request,
        'akun/akhiri_pegawai.html',
        context
    )


@login_required
def download_template_pegawai(request):

    # =====================================================
    # HAK AKSES
    # =====================================================

    profil = ambil_profil_user(
        request
    )

    boleh_download = (
        request.user.is_superuser
        or (
            profil
            and
            profil.role
            in (
                'admin_kabupaten',
                'admin_kecamatan',
                'operator_sekolah',
            )
        )
    )

    if not boleh_download:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'untuk mengunduh template pegawai.'
        )

        return redirect('/')


    # =====================================================
    # BUAT WORKBOOK
    # =====================================================

    wb = Workbook()

    ws = wb.active

    ws.title = (
        'Template Pegawai'
    )


    # =====================================================
    # HEADER
    # =====================================================

    header = [
        'Nama',
        'NIP',
        'NIK',
        'NUPTK',
        'Jenis Kelamin',
        'Tempat Lahir',
        'Tanggal Lahir',
        'Alamat',
        'No HP',
        'Jenis Pegawai',
        'Jabatan',
        'Mata Pelajaran',
        'Status Pegawai',
        'Tanggal Mulai',
        'Tanggal Pensiun',
        'Keterangan',
    ]

    ws.append(
        header
    )


    # =====================================================
    # CONTOH DATA
    # =====================================================

    ws.append([
        'Ahmad Saputra',
        '198901012020011001',
        '6402000000000001',
        '',
        'L',
        'Sangatta',
        '1989-01-01',
        'Sangatta',
        '081234567890',
        'PENDIDIK',
        'Guru Matematika',
        'Matematika',
        'PNS',
        '2026-01-01',
        '2049-01-01',
        '',
    ])


    ws.append([
        'Siti Rahma',
        '',
        '6402000000000002',
        '1234567890123456',
        'P',
        'Bengalon',
        '1993-05-10',
        'Bengalon',
        '081234567891',
        'PENDIDIK',
        'Guru Kelas',
        '',
        'PPPK',
        '2026-01-01',
        '',
        '',
    ])


    ws.append([
        'Budi Santoso',
        '',
        '6402000000000003',
        '',
        'L',
        'Sangatta',
        '1990-08-15',
        'Sangatta',
        '081234567892',
        'TENAGA_KEPENDIDIKAN',
        'Operator Sekolah',
        '',
        'HONORER',
        '2026-01-01',
        '',
        '',
    ])


    # =====================================================
    # SHEET PETUNJUK
    # =====================================================

    petunjuk = (
        wb.create_sheet(
            'Petunjuk'
        )
    )


    petunjuk.append([
        'Kolom',
        'Ketentuan',
    ])


    daftar_petunjuk = [

        (
            'Nama',
            'Wajib diisi.'
        ),

        (
            'NIP',
            (
                'Opsional. Jika diisi, '
                'harus unik.'
            )
        ),

        (
            'NIK',
            (
                'Disarankan diisi. '
                'Jika diisi, harus unik.'
            )
        ),

        (
            'NUPTK',
            (
                'Opsional. Jika diisi, '
                'harus unik.'
            )
        ),

        (
            'Jenis Kelamin',
            'Gunakan L atau P.'
        ),

        (
            'Tanggal Lahir',
            'Format YYYY-MM-DD.'
        ),

        (
            'Jenis Pegawai',
            (
                'Gunakan PENDIDIK atau '
                'TENAGA_KEPENDIDIKAN.'
            )
        ),

        (
            'Jabatan',
            (
                'Contoh: Kepala Sekolah, '
                'Guru Kelas, Guru Matematika, '
                'Kepala TU, Operator Sekolah.'
            )
        ),

        (
            'Mata Pelajaran',
            (
                'Diisi untuk PENDIDIK jika relevan. '
                'Kosongkan untuk Tenaga Kependidikan.'
            )
        ),

        (
            'Status Pegawai',
            (
                'Gunakan PNS, PPPK, HONORER, '
                'GTY, atau GTT.'
            )
        ),

        (
            'Tanggal Mulai',
            'Format YYYY-MM-DD.'
        ),

        (
            'Tanggal Pensiun',
            (
                'Opsional. '
                'Format YYYY-MM-DD.'
            )
        ),

        (
            'Sekolah',
            (
                'Tidak perlu ditulis di Excel. '
                'Sekolah ditentukan otomatis '
                'oleh akun operator.'
            )
        ),

        (
            'Tahun Ajaran',
            (
                'Tidak perlu ditulis di Excel. '
                'Sistem otomatis menggunakan '
                'Tahun Ajaran aktif.'
            )
        ),
    ]


    for data in daftar_petunjuk:

        petunjuk.append(
            list(data)
        )


    # =====================================================
    # LEBAR KOLOM TEMPLATE
    # =====================================================

    lebar_kolom = {

        'A': 28,
        'B': 24,
        'C': 22,
        'D': 22,
        'E': 18,
        'F': 20,
        'G': 18,
        'H': 30,
        'I': 18,
        'J': 26,
        'K': 24,
        'L': 22,
        'M': 20,
        'N': 18,
        'O': 18,
        'P': 30,
    }


    for kolom, lebar in (
        lebar_kolom.items()
    ):

        ws.column_dimensions[
            kolom
        ].width = lebar


    petunjuk.column_dimensions[
        'A'
    ].width = 24

    petunjuk.column_dimensions[
        'B'
    ].width = 75


    # =====================================================
    # RESPONSE
    # =====================================================

    response = HttpResponse(
        content_type=(
            'application/'
            'vnd.openxmlformats-'
            'officedocument.'
            'spreadsheetml.sheet'
        )
    )


    response[
        'Content-Disposition'
    ] = (
        'attachment; '
        'filename='
        '"template_import_pegawai.xlsx"'
    )


    wb.save(
        response
    )


    return response


@login_required
def import_pegawai_excel(request):

    # =====================================================
    # HAK AKSES
    # Hanya Operator Sekolah
    # =====================================================

    profil = ambil_profil_user(
        request
    )

    if not (
        profil
        and
        profil.role == 'operator_sekolah'
    ):

        messages.error(
            request,
            (
                'Import Data Pegawai hanya dapat '
                'dilakukan oleh Operator Sekolah.'
            )
        )

        return redirect('/')


    # =====================================================
    # SEKOLAH OPERATOR
    # =====================================================

    sekolah = ambil_sekolah_operator(
        request
    )

    if sekolah is None:

        messages.error(
            request,
            (
                'Akun operator belum terhubung '
                'dengan sekolah.'
            )
        )

        return redirect('/')


    # =====================================================
    # TAHUN AJARAN AKTIF
    # =====================================================

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )

    if tahun_ajaran is None:

        messages.error(
            request,
            'Belum ada Tahun Ajaran aktif.'
        )

        return redirect(
            'data_pegawai'
        )


    # =====================================================
    # GET
    # =====================================================

    if request.method != 'POST':

        return render(
            request,
            'akun/import_pegawai.html',
            {
                'sekolah':
                    sekolah,

                'tahun_ajaran':
                    tahun_ajaran,
            }
        )


    # =====================================================
    # FILE
    # =====================================================

    file_excel = (
        request.FILES.get(
            'file_excel'
        )
    )

    if not file_excel:

        messages.error(
            request,
            'Silakan pilih file Excel.'
        )

        return redirect(
            'import_pegawai_excel'
        )


    # =====================================================
    # VALIDASI EXTENSION
    # =====================================================

    if not (
        file_excel.name
        .lower()
        .endswith('.xlsx')
    ):

        messages.error(
            request,
            (
                'Format file harus '
                'Microsoft Excel (.xlsx).'
            )
        )

        return redirect(
            'import_pegawai_excel'
        )


    # =====================================================
    # BACA WORKBOOK
    # =====================================================

    try:

        workbook = load_workbook(
            file_excel,
            data_only=True
        )

        worksheet = (
            workbook[
                'Template Pegawai'
            ]
            if 'Template Pegawai'
            in workbook.sheetnames
            else workbook.active
        )


    except Exception as error:

        messages.error(
            request,
            (
                'File Excel tidak dapat dibaca. '
                f'{error}'
            )
        )

        return redirect(
            'import_pegawai_excel'
        )


    # =====================================================
    # HEADER WAJIB
    # =====================================================

    header_wajib = [
        'Nama',
        'NIP',
        'NIK',
        'NUPTK',
        'Jenis Kelamin',
        'Tempat Lahir',
        'Tanggal Lahir',
        'Alamat',
        'No HP',
        'Jenis Pegawai',
        'Jabatan',
        'Mata Pelajaran',
        'Status Pegawai',
        'Tanggal Mulai',
        'Tanggal Pensiun',
        'Keterangan',
    ]


    header_excel = [
        (
            str(cell.value).strip()
            if cell.value is not None
            else ''
        )
        for cell
        in worksheet[1]
    ]


    if (
        header_excel[:len(header_wajib)]
        != header_wajib
    ):

        messages.error(
            request,
            (
                'Format kolom Excel tidak sesuai '
                'dengan template resmi SIPETAKU.'
            )
        )

        return redirect(
            'import_pegawai_excel'
        )


    # =====================================================
    # HELPER NILAI TEKS
    # =====================================================

    def nilai_teks(value):

        if value is None:
            return None

        # Excel kadang membaca nomor identitas
        # sebagai angka / float.
        if isinstance(value, float):

            if value.is_integer():

                return str(
                    int(value)
                )

        if isinstance(value, int):

            return str(value)

        value = str(value).strip()

        if not value:
            return None

        return value


    # =====================================================
    # HELPER TANGGAL
    # =====================================================

    def nilai_tanggal(
        value,
        nama_kolom
    ):

        if value is None:
            return None


        if isinstance(
            value,
            datetime
        ):

            return value.date()


        if isinstance(
            value,
            date
        ):

            return value


        value = str(value).strip()

        if not value:
            return None


        daftar_format = (
            '%Y-%m-%d',
            '%d/%m/%Y',
            '%d-%m-%Y',
        )


        for format_tanggal in (
            daftar_format
        ):

            try:

                return datetime.strptime(
                    value,
                    format_tanggal
                ).date()

            except ValueError:

                continue


        raise ValueError(
            (
                f'{nama_kolom} tidak valid. '
                'Gunakan format YYYY-MM-DD.'
            )
        )


    # =====================================================
    # COUNTER
    # =====================================================

    total_baru = 0
    total_update = 0
    total_gagal = 0

    daftar_gagal = []


    # =====================================================
    # PROSES BARIS
    # =====================================================

    for nomor_baris, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        # ---------------------------------------------
        # Lewati baris benar-benar kosong
        # ---------------------------------------------

        if not any(
            value is not None
            and str(value).strip()
            for value in row
        ):

            continue


        try:

            with transaction.atomic():

                # =====================================
                # AMBIL DATA
                # =====================================

                nama = nilai_teks(
                    row[0]
                )

                nip = nilai_teks(
                    row[1]
                )

                nik = nilai_teks(
                    row[2]
                )

                nuptk = nilai_teks(
                    row[3]
                )

                jenis_kelamin = nilai_teks(
                    row[4]
                )

                tempat_lahir = nilai_teks(
                    row[5]
                )

                tanggal_lahir = nilai_tanggal(
                    row[6],
                    'Tanggal Lahir'
                )

                alamat = nilai_teks(
                    row[7]
                )

                no_hp = nilai_teks(
                    row[8]
                )

                jenis_pegawai = nilai_teks(
                    row[9]
                )

                jabatan = nilai_teks(
                    row[10]
                )

                mata_pelajaran = nilai_teks(
                    row[11]
                )

                status_pegawai = nilai_teks(
                    row[12]
                )

                tanggal_mulai = nilai_tanggal(
                    row[13],
                    'Tanggal Mulai'
                )

                tanggal_pensiun = nilai_tanggal(
                    row[14],
                    'Tanggal Pensiun'
                )

                keterangan = nilai_teks(
                    row[15]
                )


                # =====================================
                # VALIDASI DATA WAJIB
                # =====================================

                if not nama:

                    raise ValueError(
                        'Nama wajib diisi.'
                    )


                if not jenis_kelamin:

                    raise ValueError(
                        (
                            'Jenis Kelamin '
                            'wajib diisi.'
                        )
                    )


                jenis_kelamin = (
                    jenis_kelamin
                    .strip()
                    .upper()
                )


                if jenis_kelamin not in (
                    'L',
                    'P',
                ):

                    raise ValueError(
                        (
                            'Jenis Kelamin harus '
                            'L atau P.'
                        )
                    )


                if not tempat_lahir:

                    raise ValueError(
                        (
                            'Tempat Lahir '
                            'wajib diisi.'
                        )
                    )


                if tanggal_lahir is None:

                    raise ValueError(
                        (
                            'Tanggal Lahir '
                            'wajib diisi.'
                        )
                    )


                if not jenis_pegawai:

                    raise ValueError(
                        (
                            'Jenis Pegawai '
                            'wajib diisi.'
                        )
                    )


                jenis_pegawai = (
                    jenis_pegawai
                    .strip()
                    .upper()
                    .replace(' ', '_')
                )


                # Normalisasi beberapa penulisan
                if jenis_pegawai in (
                    'GURU',
                    'GURU/PENDIDIK',
                    'GURU_/_PENDIDIK',
                ):

                    jenis_pegawai = (
                        'PENDIDIK'
                    )


                if jenis_pegawai in (
                    'TENAGA_KEPENDIDIKAN',
                    'TENDIK',
                ):

                    jenis_pegawai = (
                        'TENAGA_KEPENDIDIKAN'
                    )


                if jenis_pegawai not in (
                    'PENDIDIK',
                    'TENAGA_KEPENDIDIKAN',
                ):

                    raise ValueError(
                        (
                            'Jenis Pegawai harus '
                            'PENDIDIK atau '
                            'TENAGA_KEPENDIDIKAN.'
                        )
                    )


                if not jabatan:

                    raise ValueError(
                        'Jabatan wajib diisi.'
                    )


                if not status_pegawai:

                    raise ValueError(
                        (
                            'Status Pegawai '
                            'wajib diisi.'
                        )
                    )


                status_pegawai = (
                    status_pegawai
                    .strip()
                    .upper()
                )


                if status_pegawai not in (
                    'PNS',
                    'PPPK',
                    'HONORER',
                    'GTY',
                    'GTT',
                ):

                    raise ValueError(
                        (
                            'Status Pegawai harus '
                            'PNS, PPPK, HONORER, '
                            'GTY, atau GTT.'
                        )
                    )


                # Tenaga kependidikan
                # tidak menggunakan mapel.
                if (
                    jenis_pegawai
                    == 'TENAGA_KEPENDIDIKAN'
                ):

                    mata_pelajaran = None


                # =====================================
                # IDENTIFIKASI PEGAWAI
                # =====================================

                pegawai_nik = None
                pegawai_nip = None
                pegawai_nuptk = None


                if nik:

                    pegawai_nik = (
                        Pegawai.objects
                        .filter(
                            nik=nik
                        )
                        .first()
                    )


                if nip:

                    pegawai_nip = (
                        Pegawai.objects
                        .filter(
                            nip=nip
                        )
                        .first()
                    )


                if nuptk:

                    pegawai_nuptk = (
                        Pegawai.objects
                        .filter(
                            nuptk=nuptk
                        )
                        .first()
                    )


                # =====================================
                # CEK KONFLIK IDENTITAS
                # =====================================

                kandidat = [
                    p
                    for p in (
                        pegawai_nik,
                        pegawai_nip,
                        pegawai_nuptk,
                    )
                    if p is not None
                ]


                kandidat_id = {
                    p.id
                    for p in kandidat
                }


                if len(kandidat_id) > 1:

                    raise ValueError(
                        (
                            'NIK, NIP, atau NUPTK '
                            'pada baris ini terdaftar '
                            'pada pegawai yang berbeda. '
                            'Periksa kembali identitas.'
                        )
                    )


                # =====================================
                # PEGAWAI LAMA / BARU
                # =====================================

                pegawai = (
                    pegawai_nik
                    or pegawai_nip
                    or pegawai_nuptk
                )


                pegawai_baru = (
                    pegawai is None
                )


                # =====================================
                # BUAT PEGAWAI BARU
                # =====================================

                if pegawai_baru:

                    pegawai = Pegawai(
                        nama=nama,
                        nip=nip,
                        nik=nik,
                        nuptk=nuptk,
                        jenis_kelamin=(
                            jenis_kelamin
                        ),
                        tempat_lahir=(
                            tempat_lahir
                        ),
                        tanggal_lahir=(
                            tanggal_lahir
                        ),
                        alamat=alamat,
                        no_hp=no_hp,
                    )


                # =====================================
                # UPDATE IDENTITAS PEGAWAI LAMA
                # =====================================

                else:

                    pegawai.nama = nama

                    pegawai.jenis_kelamin = (
                        jenis_kelamin
                    )

                    pegawai.tempat_lahir = (
                        tempat_lahir
                    )

                    pegawai.tanggal_lahir = (
                        tanggal_lahir
                    )

                    pegawai.alamat = alamat

                    pegawai.no_hp = no_hp


                    # Identitas unik tidak kita
                    # kosongkan bila Excel kosong.
                    if nip:
                        pegawai.nip = nip

                    if nik:
                        pegawai.nik = nik

                    if nuptk:
                        pegawai.nuptk = nuptk


                pegawai.full_clean()

                pegawai.save()


                # =====================================
                # CEK RIWAYAT AKTIF TAHUN INI
                # =====================================

                riwayat_aktif = (
                    RiwayatPegawai.objects
                    .filter(
                        pegawai=pegawai,
                        tahun_ajaran=(
                            tahun_ajaran
                        ),
                        status_akhir='AKTIF'
                    )
                    .select_related(
                        'sekolah'
                    )
                    .first()
                )


                # =====================================
                # JIKA AKTIF DI SEKOLAH LAIN
                # =====================================

                if (
                    riwayat_aktif
                    and
                    riwayat_aktif.sekolah_id
                    != sekolah.id
                ):

                    raise ValueError(
                        (
                            f'Pegawai sudah aktif di '
                            f'{riwayat_aktif.sekolah.nama} '
                            f'pada Tahun Ajaran '
                            f'{tahun_ajaran.nama}. '
                            f'Perpindahan sekolah harus '
                            f'dilakukan melalui fitur Mutasi.'
                        )
                    )


                # =====================================
                # UPDATE RIWAYAT DI SEKOLAH YANG SAMA
                # =====================================

                if riwayat_aktif:

                    riwayat_aktif.jenis_pegawai = (
                        jenis_pegawai
                    )

                    riwayat_aktif.jabatan = (
                        jabatan
                    )

                    riwayat_aktif.mata_pelajaran = (
                        mata_pelajaran
                    )

                    riwayat_aktif.status_pegawai = (
                        status_pegawai
                    )

                    riwayat_aktif.tanggal_mulai = (
                        tanggal_mulai
                    )

                    riwayat_aktif.tanggal_pensiun = (
                        tanggal_pensiun
                    )

                    riwayat_aktif.keterangan = (
                        keterangan or ''
                    )

                    riwayat_aktif.full_clean()

                    riwayat_aktif.save()

                    total_update += 1


                # =====================================
                # BUAT RIWAYAT BARU
                # =====================================

                else:

                    riwayat_baru = (
                        RiwayatPegawai(
                            pegawai=pegawai,
                            sekolah=sekolah,
                            tahun_ajaran=(
                                tahun_ajaran
                            ),
                            jenis_pegawai=(
                                jenis_pegawai
                            ),
                            jabatan=jabatan,
                            mata_pelajaran=(
                                mata_pelajaran
                            ),
                            status_pegawai=(
                                status_pegawai
                            ),
                            status_akhir='AKTIF',
                            tanggal_mulai=(
                                tanggal_mulai
                            ),
                            tanggal_selesai=None,
                            tanggal_pensiun=(
                                tanggal_pensiun
                            ),
                            keterangan=(
                                keterangan or ''
                            ),
                        )
                    )

                    riwayat_baru.full_clean()

                    riwayat_baru.save()


                    if pegawai_baru:

                        total_baru += 1

                    else:

                        # Pegawai sudah pernah ada,
                        # tetapi belum punya riwayat
                        # aktif tahun ini.
                        total_update += 1


        except Exception as error:

            total_gagal += 1

            nama_error = (
                nilai_teks(row[0])
                or
                f'Baris {nomor_baris}'
            )

            daftar_gagal.append({
                'baris':
                    nomor_baris,

                'nama':
                    nama_error,

                'pesan':
                    str(error),
            })


    # =====================================================
    # HASIL
    # =====================================================

    context = {

        'sekolah':
            sekolah,

        'tahun_ajaran':
            tahun_ajaran,

        'total_baru':
            total_baru,

        'total_update':
            total_update,

        'total_gagal':
            total_gagal,

        'daftar_gagal':
            daftar_gagal,

        'import_selesai':
            True,
    }


    return render(
        request,
        'akun/import_pegawai.html',
        context
    )


@login_required
def arsip_pegawai(request):

    # =====================================================
    # PROFIL + TAHUN AJARAN
    # =====================================================

    profil = ambil_profil_user(request)

    tahun_ajaran = (
        TahunAjaran.objects
        .filter(
            aktif=True
        )
        .first()
    )


    # =====================================================
    # QUERY DASAR
    # =====================================================

    arsip = (
        RiwayatPegawai.objects
        .filter(
            status_akhir__in=[
                'MUTASI',
                'PENSIUN',
                'BERHENTI',
                'MENINGGAL',
            ]
        )
        .select_related(
            'pegawai',
            'sekolah',
            'tahun_ajaran'
        )
    )


    # =====================================================
    # HAK AKSES
    # =====================================================

    sekolah_operator = None


    if request.user.is_superuser:

        pass


    elif (
        profil
        and
        profil.role == 'admin_kabupaten'
    ):

        pass


    elif (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        arsip = (
            arsip
            .filter(
                sekolah__kecamatan=profil.kecamatan
            )
        )


    elif (
        profil
        and
        profil.role == 'operator_sekolah'
    ):

        sekolah_operator = (
            ambil_sekolah_operator(
                request
            )
        )


        if sekolah_operator is None:

            messages.error(
                request,
                'Operator belum terhubung '
                'dengan sekolah.'
            )

            return redirect('/')


        arsip = (
            arsip
            .filter(
                sekolah=sekolah_operator
            )
        )


    else:

        messages.error(
            request,
            'Anda tidak memiliki akses '
            'ke Arsip Pegawai.'
        )

        return redirect('/')


    # =====================================================
    # FILTER
    # =====================================================

    status_akhir = (
        request.GET.get(
            'status',
            ''
        )
        .strip()
        .upper()
    )

    kecamatan = (
        request.GET.get(
            'kecamatan',
            ''
        )
        .strip()
    )

    sekolah_id = (
        request.GET.get(
            'sekolah',
            ''
        )
        .strip()
    )

    cari = (
        request.GET.get(
            'cari',
            ''
        )
        .strip()
    )


    # =====================================================
    # FILTER STATUS
    # =====================================================

    if status_akhir in (
        'MUTASI',
        'PENSIUN',
        'BERHENTI',
        'MENINGGAL',
    ):

        arsip = (
            arsip
            .filter(
                status_akhir=status_akhir
            )
        )


    # =====================================================
    # FILTER KECAMATAN
    # =====================================================

    if (
        kecamatan
        and (
            request.user.is_superuser
            or (
                profil
                and
                profil.role == 'admin_kabupaten'
            )
        )
    ):

        arsip = (
            arsip
            .filter(
                sekolah__kecamatan=kecamatan
            )
        )


    # =====================================================
    # FILTER SEKOLAH
    # =====================================================

    if sekolah_id:

        arsip = (
            arsip
            .filter(
                sekolah_id=sekolah_id
            )
        )


    # =====================================================
    # PENCARIAN
    # =====================================================

    if cari:

        arsip = (
            arsip
            .filter(
                Q(
                    pegawai__nama__icontains=cari
                )
                |
                Q(
                    pegawai__nip__icontains=cari
                )
                |
                Q(
                    pegawai__nik__icontains=cari
                )
                |
                Q(
                    pegawai__nuptk__icontains=cari
                )
                |
                Q(
                    jabatan__icontains=cari
                )
                |
                Q(
                    sekolah__nama__icontains=cari
                )
            )
        )


    arsip = (
        arsip
        .order_by(
            '-tanggal_selesai',
            '-tahun_ajaran__nama',
            'pegawai__nama'
        )
    )


    # =====================================================
    # QUERY RINGKASAN SESUAI HAK AKSES
    # =====================================================

    ringkasan_qs = (
        RiwayatPegawai.objects
        .filter(
            status_akhir__in=[
                'MUTASI',
                'PENSIUN',
                'BERHENTI',
                'MENINGGAL',
            ]
        )
    )


    if (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        ringkasan_qs = (
            ringkasan_qs
            .filter(
                sekolah__kecamatan=profil.kecamatan
            )
        )


    elif (
        profil
        and
        profil.role == 'operator_sekolah'
        and sekolah_operator
    ):

        ringkasan_qs = (
            ringkasan_qs
            .filter(
                sekolah=sekolah_operator
            )
        )


    total_arsip = (
        ringkasan_qs.count()
    )

    total_mutasi = (
        ringkasan_qs
        .filter(
            status_akhir='MUTASI'
        )
        .count()
    )

    total_pensiun = (
        ringkasan_qs
        .filter(
            status_akhir='PENSIUN'
        )
        .count()
    )

    total_berhenti = (
        ringkasan_qs
        .filter(
            status_akhir='BERHENTI'
        )
        .count()
    )

    total_meninggal = (
        ringkasan_qs
        .filter(
            status_akhir='MENINGGAL'
        )
        .count()
    )


    # =====================================================
    # PILIHAN FILTER
    # =====================================================

    if (
        request.user.is_superuser
        or (
            profil
            and
            profil.role == 'admin_kabupaten'
        )
    ):

        daftar_kecamatan = (
            Sekolah.objects
            .exclude(
                kecamatan=''
            )
            .values_list(
                'kecamatan',
                flat=True
            )
            .distinct()
            .order_by(
                'kecamatan'
            )
        )

        daftar_sekolah = (
            Sekolah.objects
            .all()
            .order_by(
                'kecamatan',
                'nama'
            )
        )


    elif (
        profil
        and
        profil.role == 'admin_kecamatan'
    ):

        daftar_kecamatan = (
            Sekolah.objects
            .filter(
                kecamatan=profil.kecamatan
            )
            .values_list(
                'kecamatan',
                flat=True
            )
            .distinct()
        )

        daftar_sekolah = (
            Sekolah.objects
            .filter(
                kecamatan=profil.kecamatan
            )
            .order_by(
                'nama'
            )
        )


    else:

        daftar_kecamatan = []

        daftar_sekolah = (
            Sekolah.objects
            .filter(
                id=sekolah_operator.id
            )
            if sekolah_operator
            else Sekolah.objects.none()
        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'arsip':
            arsip,

        'tahun_ajaran':
            tahun_ajaran,

        'total_arsip':
            total_arsip,

        'total_mutasi':
            total_mutasi,

        'total_pensiun':
            total_pensiun,

        'total_berhenti':
            total_berhenti,

        'total_meninggal':
            total_meninggal,

        'daftar_kecamatan':
            daftar_kecamatan,

        'daftar_sekolah':
            daftar_sekolah,

        'filter_status':
            status_akhir,

        'filter_kecamatan':
            kecamatan,

        'filter_sekolah':
            sekolah_id,

        'filter_cari':
            cari,

        'admin_kabupaten':
            (
                request.user.is_superuser
                or (
                    profil
                    and
                    profil.role == 'admin_kabupaten'
                )
            ),

        'admin_kecamatan':
            (
                profil
                and
                profil.role == 'admin_kecamatan'
            ),

        'operator_sekolah':
            (
                profil
                and
                profil.role == 'operator_sekolah'
            ),
    }


    return render(
        request,
        'akun/arsip_pegawai.html',
        context
    )


@login_required
def ubah_password(request):

    # =====================================================
    # USER LOGIN
    # =====================================================

    user = request.user


    # =====================================================
    # PROSES POST
    # =====================================================

    if request.method == 'POST':

        password_lama = (
            request.POST
            .get(
                'password_lama',
                ''
            )
        )

        password_baru = (
            request.POST
            .get(
                'password_baru',
                ''
            )
        )

        konfirmasi_password = (
            request.POST
            .get(
                'konfirmasi_password',
                ''
            )
        )


        # =================================================
        # VALIDASI PASSWORD LAMA
        # =================================================

        if not user.check_password(
            password_lama
        ):

            messages.error(
                request,
                'Password lama tidak sesuai.'
            )

            return redirect(
                'ubah_password'
            )


        # =================================================
        # PASSWORD BARU WAJIB DIISI
        # =================================================

        if not password_baru:

            messages.error(
                request,
                'Password baru wajib diisi.'
            )

            return redirect(
                'ubah_password'
            )


        # =================================================
        # KONFIRMASI PASSWORD
        # =================================================

        if (
            password_baru
            != konfirmasi_password
        ):

            messages.error(
                request,
                (
                    'Konfirmasi password baru '
                    'tidak sesuai.'
                )
            )

            return redirect(
                'ubah_password'
            )


        # =================================================
        # PASSWORD BARU TIDAK BOLEH SAMA
        # =================================================

        if user.check_password(
            password_baru
        ):

            messages.error(
                request,
                (
                    'Password baru tidak boleh '
                    'sama dengan password lama.'
                )
            )

            return redirect(
                'ubah_password'
            )


        # =================================================
        # VALIDATOR PASSWORD DJANGO
        # =================================================

        try:

            validate_password(
                password_baru,
                user=user
            )

        except ValidationError as error:

            for pesan in error.messages:

                messages.error(
                    request,
                    pesan
                )

            return redirect(
                'ubah_password'
            )


        # =================================================
        # SIMPAN PASSWORD
        # =================================================

        user.set_password(
            password_baru
        )

        user.save(
            update_fields=[
                'password'
            ]
        )


        # =================================================
        # PERTAHANKAN SESSION LOGIN
        # =================================================

        update_session_auth_hash(
            request,
            user
        )


        # =================================================
        # BERHASIL
        # =================================================

        messages.success(
            request,
            (
                'Password berhasil diubah. '
                'Gunakan password baru untuk '
                'login berikutnya.'
            )
        )

        return redirect(
            'ubah_password'
        )


    # =====================================================
    # TAMPILKAN HALAMAN
    # =====================================================

    context = {
        'user_login': user,
    }


    return render(
        request,
        'akun/ubah_password.html',
        context
    )